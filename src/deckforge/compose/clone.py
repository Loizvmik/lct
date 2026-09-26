"""Сборка слайда клоном слайда-примера шаблона: CLONE → BIND → ADAPT.

Старый путь (`builder.place_slide`) строит слайд с нуля: пустой слайд на
лейауте, декор перерисовывается по слепку `DecorShape`, текст льётся в
новые текстовые рамки. По дороге теряется всё, чего слепок не описывает:
группы, градиенты, фото-контейнеры, нестандартные фигуры, стрелки,
форматирование текста. Пользователь видит «белые листы не в дизайне
шаблона».

Здесь слайд-пример копируется целиком (дерево фигур, фон, связи с
картинками), а потом меняется только то, что обязано поменяться: текст в
слотах, лишние единицы повтора, фото пользователя. Решение, какой путь
выбрать и принят ли результат, остаётся за `builder._place_best_candidate`:
модуль только умеет клонировать, сопоставлять и переписывать.

Модель здесь не вызывается (граница слоя `compose/`), python-pptx трогается
только ради частей пакета и связей: сами фигуры правятся прямо по lxml.
"""
from __future__ import annotations
import copy
import io
import re
from dataclasses import dataclass, replace
from typing import Callable, Iterable

from lxml import etree
from PIL import Image
from pptx.opc.packuri import PackURI

from deckforge.ooxml.geometry import Box, Canvas
from deckforge.ooxml.ns import NS, qn
from deckforge.ooxml.walk import ShapeRef, walk_shapes
from deckforge.template.patterns import DecorShape, PatternSlot

# Порог совпадения коробки слота из профиля с фигурой клона (IoU). Коробки
# слотов сняты с ТОГО ЖЕ слайда-примера, но проходят притяжку к полям сетки
# (`patterns._snap_to_margins`), поэтому ровно 1.0 не выходит. Замер на
# VK Education: у 150 слотов из 160 IoU выше 0.97, притянутые к полям дают
# 0.85-0.92. Ниже 0.85 начинаются уже соседние фигуры той же строки
# (подпись месяца на диаграмме Ганта: 0.77 с соседней ячейкой), и
# совпадение по ним было бы подменой, а не узнаванием.
SLOT_MATCH_IOU = 0.85

# Запасной критерий для вырожденных коробок (линия, стрелка: высота или
# ширина около нуля, IoU для них не определён): все четыре края в пределах
# половины процента холста, это меньше толщины любой видимой линии шаблона.
_EDGE_TOLERANCE = 0.005

# Связи слайда, которые клону не нужны: лейаут у клона свой (add_slide уже
# связал), заметки пишутся заново (`builder._write_speaker_notes`), чужая
# страница заметок на клоне показывала бы текст примера.
_SKIPPED_RELTYPES = ("/notesSlide", "/slideLayout")

_R_NS = "{%s}" % NS["r"]

# Роли слотов, под которые в клоне стоит не текстовая фигура.
_PICTURE_ROLES = frozenset({"image", "icon"})
_FRAME_ROLES = frozenset({"table", "chart"})

# Плейсхолдеры, которые PowerPoint заполняет сам (номер слайда, дата,
# колонтитулы): пустые в разметке, но осмысленные, как и в
# `builder._AUTO_FILLED_PLACEHOLDER_TYPES`.
_AUTO_PH_TYPES = frozenset({"sldNum", "dt", "ftr", "hdr"})

# Висячий отступ маркера списка (EMU, 0.3125″): без него символ маркера
# прилипает к первой букве пункта. Ставится только там, где шаблон свой
# отступ не задал.
_BULLET_INDENT_EMU = 285750

_SLIDE_NUMBER_RE = re.compile(r"slide(\d+)\.xml$")
_PARTNAME_RE = re.compile(r"^(.*?)(\d*)(\.[^./]+)$")

# Порядок детей a:pPr по схеме OOXML: маркер обязан стоять перед tabLst/
# defRPr/extLst, иначе PowerPoint считает файл повреждённым.
_PPR_TAIL = (qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst"))


@dataclass(frozen=True)
class TextStyle:
    """Кегль/гарнитура/интерлиньяж/поля рамки, которыми текст фигуры клона
    реально будет отрисован. Нужен замеру в `builder`: мерить надо тем же,
    чем рисует PowerPoint и чем потом мерит аудит (L03)."""
    size_pt: float | None
    family: str | None
    line_spacing: float | None
    insets_in: tuple[float, float, float, float]  # слева, сверху, справа, снизу


# ---------------------------------------------------------------------------
# CLONE
# ---------------------------------------------------------------------------


def sample_slides_by_number(prs) -> dict[int, object]:
    """Слайды-примеры по номеру части (`ppt/slides/slideN.xml` → N), тот же
    номер, что `Pattern.source_slide_index`. Берётся ДО
    `builder._clear_sample_slides`: после `drop_rel` слайды пропадают из
    `prs.slides`, но сами объекты частей живут в памяти, и клон может с
    ними работать."""
    result = {}
    for slide in prs.slides:
        match = _SLIDE_NUMBER_RE.search(str(slide.part.partname))
        if match:
            result[int(match.group(1))] = slide
    return result


def clone_example_slide(prs, source_slide, layout):
    """Новый слайд на `layout` с копией дерева фигур, фона и связей
    `source_slide`. Плейсхолдеры, которые `add_slide` положил сам, убираются
    до копирования: у примера свои, с его текстом и форматированием, и
    двойной комплект дал бы пустые «Щёлкните, чтобы добавить текст» под
    настоящими фигурами."""
    slide = prs.slides.add_slide(layout)
    sp_tree = slide.shapes._spTree  # noqa: SLF001: python-pptx не даёт публичного доступа к p:spTree
    for el in list(sp_tree):
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            sp_tree.remove(el)

    rid_map = _copy_relationships(source_slide.part, slide.part)
    copied = []
    for el in source_slide.shapes._spTree:  # noqa: SLF001
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        dup = copy.deepcopy(el)
        sp_tree.append(dup)
        copied.append(dup)

    # Свой фон есть у 20 из 55 примеров VK Education: без него клон встаёт
    # на фон лейаута, и тёмный слайд-разделитель выходит белым.
    src_root, dst_root = source_slide._element, slide._element  # noqa: SLF001
    src_csld, dst_csld = src_root.find(qn("p:cSld")), dst_root.find(qn("p:cSld"))
    src_bg = src_csld.find(qn("p:bg"))
    if src_bg is not None:
        old_bg = dst_csld.find(qn("p:bg"))
        if old_bg is not None:
            dst_csld.remove(old_bg)
        bg = copy.deepcopy(src_bg)
        dst_csld.insert(0, bg)
        copied.append(bg)
    src_map = src_root.find(qn("p:clrMapOvr"))
    if src_map is not None:
        old_map = dst_root.find(qn("p:clrMapOvr"))
        if old_map is not None:
            dst_root.replace(old_map, copy.deepcopy(src_map))
        else:
            dst_csld.addnext(copy.deepcopy(src_map))
    for attr in ("show", "showMasterSp", "showMasterPhAnim"):
        if src_root.get(attr) is not None:
            dst_root.set(attr, src_root.get(attr))

    for el in copied:
        _remap_rids(el, rid_map)
    return slide


def _copy_relationships(source_part, target_part) -> dict[str, str]:
    """Связи примера переносятся в клон, каждая получает СВОЙ rId клона:
    номера у двух частей независимы, и совпадение старого rId с занятым
    номером клона (rId1 у клона уже занят лейаутом) подменило бы картинку
    лейаутом. Возвращает старый rId → новый."""
    mapping: dict[str, str] = {}
    for rid, rel in source_part.rels.items():
        if rel.reltype.endswith(_SKIPPED_RELTYPES):
            continue
        if rel.is_external:
            mapping[rid] = target_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            mapping[rid] = target_part.relate_to(rel.target_part, rel.reltype)
    return mapping


def _remap_rids(root, mapping: dict[str, str]) -> None:
    """Переписывает ВСЕ атрибуты пространства имён связей (r:embed, r:link,
    r:id, r:pict…) по таблице `mapping`. Перечислять имена атрибутов не
    стали: svg-картинка во вложенном extLst, гиперссылка, видео ссылаются
    каждый своим атрибутом, и пропущенный оставил бы битую ссылку."""
    for el in root.iter():
        for name, value in el.attrib.items():
            if name.startswith(_R_NS) and value in mapping:
                el.set(name, mapping[value])


# Метка клона в имени слайда: префикс и `pattern_id` примера. Отсюда её
# читают сборка (`builder._try_clone` ставит), `scripts/inspect_deck.py` и
# аудит (`clone_pattern_id`: D05 сравнивает заполненность с примером).
CLONE_MARK_PREFIX = "deckforge:clone:"


def clone_pattern_id(slide_root) -> str | None:
    """`pattern_id` примера, клоном которого собран слайд (`p:sld`), либо
    `None`, если метки нет: слайд собран с нуля или файл чужой."""
    c_sld = slide_root.find(qn("p:cSld"))
    name = (c_sld.get("name") or "") if c_sld is not None else ""
    if not name.startswith(CLONE_MARK_PREFIX):
        return None
    return name[len(CLONE_MARK_PREFIX):] or None


def mark_slide(slide, label: str) -> None:
    """Метка пути сборки в имени слайда (`p:cSld/@name`): поле штатное,
    PowerPoint его не показывает, а `scripts/inspect_deck.py` читает, каким
    путём собран каждый слайд готового файла."""
    slide._element.find(qn("p:cSld")).set("name", label)  # noqa: SLF001


# ---------------------------------------------------------------------------
# Сопоставление слотов и фигур
# ---------------------------------------------------------------------------


def box_iou(a: Box, b: Box) -> float:
    inter = a.intersect(b)
    if inter is None:
        return 0.0
    union = a.area + b.area - inter.area
    return inter.area / union if union > 0 else 0.0


def same_box(a: Box, b: Box) -> bool:
    if box_iou(a, b) >= SLOT_MATCH_IOU:
        return True
    return (
        abs(a.left - b.left) <= _EDGE_TOLERANCE and abs(a.top - b.top) <= _EDGE_TOLERANCE
        and abs(a.right - b.right) <= _EDGE_TOLERANCE and abs(a.bottom - b.bottom) <= _EDGE_TOLERANCE
    )


def _box_score(a: Box, b: Box) -> float:
    return max(box_iou(a, b), 1.0 if same_box(a, b) else 0.0)


def slide_refs(slide, canvas: Canvas) -> list[ShapeRef]:
    """Листья дерева фигур слайда с коробками в долях холста. Плейсхолдер
    без своего `a:xfrm` получает коробку плейсхолдера лейаута (по idx, затем
    по типу): тот же порядок наследования, по которому майнинг снимал
    коробки слотов (`patterns._resolve_slide_boxes`), иначе заголовок почти
    любого примера не нашёлся бы вовсе."""
    refs = list(walk_shapes(slide._element, canvas))  # noqa: SLF001
    if not any(r.box is None and r.is_placeholder for r in refs):
        return refs
    layout_refs = list(walk_shapes(slide.slide_layout._element, canvas))  # noqa: SLF001
    by_idx = {r.ph_idx: r.box for r in layout_refs if r.is_placeholder and r.box is not None and r.ph_idx is not None}
    by_type: dict[str | None, Box] = {}
    for r in layout_refs:
        if r.is_placeholder and r.box is not None and r.ph_type not in by_type:
            by_type[r.ph_type] = r.box
    out = []
    for r in refs:
        if r.box is None and r.is_placeholder:
            box = by_idx.get(r.ph_idx) or by_type.get(r.ph_type)
            if box is not None:
                r = replace(r, box=box)
        out.append(r)
    return out


def _has_text_body(element) -> bool:
    return element.find(qn("p:txBody")) is not None


def _kind_fits(role: str, ref: ShapeRef) -> bool:
    if role in _PICTURE_ROLES:
        return ref.kind == "picture"
    if role in _FRAME_ROLES:
        return ref.kind == "graphic_frame"
    return ref.kind == "shape" and _has_text_body(ref.element)


def clone_map(refs: Iterable[ShapeRef]) -> dict[str, ShapeRef]:
    """Id фигуры (`p:cNvPr/@id`) → лист дерева клона. `deepcopy` в
    `clone_example_slide` id не меняет, а python-pptx новым фигурам даёт
    id больше занятых, поэтому id из профиля (`PatternSlot.source_shape_id`)
    указывает в клоне ровно на ту фигуру, с которой снят слот. Id, который
    встречается на слайде дважды (так бывает у файлов после ручной
    склейки), в карту не попадает: по нему фигуру не узнать, и вызывающий
    идёт по коробке."""
    seen: dict[str, ShapeRef] = {}
    doubled: set[str] = set()
    for ref in refs:
        if not ref.shape_id:
            continue
        if ref.shape_id in seen:
            doubled.add(ref.shape_id)
        seen[ref.shape_id] = ref
    return {k: v for k, v in seen.items() if k not in doubled}


def match_slots(slide, slots: list[PatternSlot], canvas: Canvas) -> dict[int, ShapeRef | None]:
    """Номер слота в `slots` → фигура клона на его месте или `None`, если
    такой не нашлось.

    Сначала по id исходной фигуры (`clone_map`): слот снят с этого же
    слайда, и id узнаёт фигуру, даже когда притяжка к полям сдвинула
    коробку слота ниже порога. Фигура по id принимается, только если вид
    подходит слоту (`_kind_fits`): id из профиля другого разбора мог
    указать на картинку там, где нужна надпись. Остальные слоты идут по
    IoU коробок (не ниже `SLOT_MATCH_IOU`), жадно по убыванию и каждая
    фигура одному слоту: у двух слотов одной строки таблицы-таймлайна
    коробки близки, и без этого оба указали бы на одну и ту же фигуру."""
    refs = [r for r in slide_refs(slide, canvas) if r.box is not None]
    by_id = clone_map(refs)
    result: dict[int, ShapeRef | None] = {i: None for i in range(len(slots))}
    used: set[int] = set()
    for i, slot in enumerate(slots):
        ref = by_id.get(slot.source_shape_id) if slot.source_shape_id else None
        if ref is None or id(ref.element) in used or not _kind_fits(slot.role, ref):
            continue
        result[i] = ref
        used.add(id(ref.element))
    pairs = []
    for i, slot in enumerate(slots):
        if result[i] is not None:
            continue
        for j, ref in enumerate(refs):
            if id(ref.element) in used or not _kind_fits(slot.role, ref):
                continue
            score = _box_score(slot.box, ref.box)
            if score >= SLOT_MATCH_IOU:
                pairs.append((score, i, j))
    pairs.sort(key=lambda t: -t[0])
    for _score, i, j in pairs:
        if result[i] is not None or id(refs[j].element) in used:
            continue
        result[i] = refs[j]
        used.add(id(refs[j].element))
    return result


# ---------------------------------------------------------------------------
# BIND
# ---------------------------------------------------------------------------


def bind_text(shape_element, paragraphs, *, bullet_char: str | None = None) -> None:
    """Заменяет текст фигуры, сохраняя форматирование примера. Образец
    абзаца берётся с первого абзаца с текстом (его `a:pPr`), образец run с
    первого run (его `a:rPr`: кегль, цвет, начертание, гарнитура). Каждый
    новый абзац делается копией образца; остальные абзацы, run, переносы и
    поля примера удаляются. `a:bodyPr`/`a:lstStyle` не трогаются: там поля рамки,
    автоподбор и стили уровней, которые шаблон задал этой фигуре.

    `paragraphs`: объекты с `.text` и `.bullet` (`blocks.Paragraph`).
    Маркер ставится, только если пунктов больше одного и `bullet_char`
    передан: одиночный «список» в карточке с маркером выглядит опечаткой,
    а у примера, чей абзац уже несёт свой маркер, он остаётся шаблонным."""
    tx_body = shape_element.find(qn("p:txBody"))
    if tx_body is None:
        raise ValueError("bind_text: у фигуры нет p:txBody")
    _bind_tx_body(tx_body, paragraphs, bullet_char)


def _bind_tx_body(tx_body, paragraphs, bullet_char: str | None) -> None:
    """Общая часть `bind_text` и ячейки таблицы (`a:tc/a:txBody`): у фигуры
    и у ячейки разметка абзацев одна и та же, различается только обёртка."""
    old_paragraphs = tx_body.findall(qn("a:p"))
    proto_p = next((p for p in old_paragraphs if p.find(qn("a:r")) is not None), None)
    if proto_p is None:
        proto_p = old_paragraphs[0] if old_paragraphs else etree.Element(qn("a:p"))
    proto_r = proto_p.find(qn("a:r"))
    proto_p = copy.deepcopy(proto_p)
    proto_r = copy.deepcopy(proto_r) if proto_r is not None else _run_from_end_props(proto_p)
    for child in list(proto_p):
        if child.tag not in (qn("a:pPr"), qn("a:endParaRPr")):
            proto_p.remove(child)
    r_pr = proto_r.find(qn("a:rPr"))
    if r_pr is not None:
        # Флаг «орфографическая ошибка» относится к тексту примера, не к нашему.
        r_pr.attrib.pop("err", None)
    for p in old_paragraphs:
        tx_body.remove(p)

    items = list(paragraphs) or [_EmptyParagraph()]
    with_bullets = bullet_char is not None and len(items) > 1
    for para in items:
        p = copy.deepcopy(proto_p)
        end = p.find(qn("a:endParaRPr"))
        lines = (para.text or "").split("\n")
        for k, line in enumerate(lines):
            if k:
                br = etree.Element(qn("a:br"))
                if r_pr is not None:
                    br.append(copy.deepcopy(r_pr))
                _insert_before(p, br, end)
            run = copy.deepcopy(proto_r)
            if getattr(para, "bold", False):
                # Заголовок карточки, склеенный с телом (`blocks._assign_cards`):
                # начертание примера остаётся, добавляется только жирность.
                run_pr = run.find(qn("a:rPr"))
                if run_pr is None:
                    run_pr = etree.Element(qn("a:rPr"))
                    run.insert(0, run_pr)
                run_pr.set("b", "1")
            t = run.find(qn("a:t"))
            if t is None:
                t = etree.SubElement(run, qn("a:t"))
            t.text = line
            _insert_before(p, run, end)
        if with_bullets and getattr(para, "bullet", False):
            _ensure_bullet(p, bullet_char)
        tx_body.append(p)


class _EmptyParagraph:
    text = ""
    bullet = False


def _insert_before(parent, el, anchor) -> None:
    if anchor is not None:
        anchor.addprevious(el)
    else:
        parent.append(el)


def _run_from_end_props(p) -> etree._Element:
    """Образец run для пустой фигуры примера (пустой плейсхолдер): свойства
    берутся из `a:endParaRPr`: это то, чем PowerPoint напечатал бы первую
    букву, если бы её ввели руками."""
    run = etree.Element(qn("a:r"))
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        r_pr = copy.deepcopy(end)
        r_pr.tag = qn("a:rPr")
        # Кегль из `endParaRPr` не переносится: у пустого плейсхолдера это
        # заглушка экспорта (VK Education, обложка раздела: 16 pt при
        # заголовке макета 48 pt), и записанный явно он перекрывал
        # унаследованный от макета, тот самый, что измерил разбор шаблона.
        # Без `sz` текст наследует кегль макета, как и напечатанный руками.
        r_pr.attrib.pop("sz", None)
        run.append(r_pr)
    etree.SubElement(run, qn("a:t"))
    return run


def _ensure_bullet(p, bullet_char: str) -> None:
    p_pr = p.find(qn("a:pPr"))
    if p_pr is None:
        p_pr = etree.Element(qn("a:pPr"))
        p.insert(0, p_pr)
    if p_pr.find(qn("a:buChar")) is not None or p_pr.find(qn("a:buAutoNum")) is not None:
        return
    for tag in ("a:buNone", "a:buBlip"):
        el = p_pr.find(qn(tag))
        if el is not None:
            p_pr.remove(el)
    bu = etree.Element(qn("a:buChar"))
    bu.set("char", bullet_char)
    tail = next((c for c in p_pr if c.tag in _PPR_TAIL), None)
    _insert_before(p_pr, bu, tail)
    # Пример без своего отступа (или с явным нулевым, как у абзаца-подписи)
    # получает висячий: у пункта списка маркер не должен касаться буквы.
    if int(p_pr.get("marL") or 0) == 0 and int(p_pr.get("indent") or 0) >= 0:
        p_pr.set("marL", str(_BULLET_INDENT_EMU))
        p_pr.set("indent", str(-_BULLET_INDENT_EMU))


def text_style(shape_element) -> TextStyle:
    """Чем будет нарисован текст фигуры после `bind_text`: кегль и
    гарнитура первого run, интерлиньяж первого абзаца, поля рамки из
    `a:bodyPr` (по умолчанию OOXML 0.1″ по бокам и 0.05″ сверху и снизу).
    `None` у кегля/гарнитуры/интерлиньяжа: фигура наследует их от лейаута
    или мастера, решает вызывающий."""
    tx_body = shape_element.find(qn("p:txBody"))
    size = family = spacing = None
    run = tx_body.find(".//" + qn("a:r")) if tx_body is not None else None
    r_pr = run.find(qn("a:rPr")) if run is not None else None
    if r_pr is not None:
        if r_pr.get("sz"):
            size = int(r_pr.get("sz")) / 100
        latin = r_pr.find(qn("a:latin"))
        if latin is not None and latin.get("typeface") and not latin.get("typeface").startswith("+"):
            family = latin.get("typeface")
    first_p = tx_body.find(qn("a:p")) if tx_body is not None else None
    p_pr = first_p.find(qn("a:pPr")) if first_p is not None else None
    pct = p_pr.find(qn("a:lnSpc") + "/" + qn("a:spcPct")) if p_pr is not None else None
    if pct is not None and pct.get("val"):
        spacing = int(pct.get("val")) / 100000
    body_pr = tx_body.find(qn("a:bodyPr")) if tx_body is not None else None

    def inset(name: str, default_emu: int) -> float:
        raw = body_pr.get(name) if body_pr is not None else None
        return (int(raw) if raw is not None else default_emu) / 914400

    insets = (inset("lIns", 91440), inset("tIns", 45720), inset("rIns", 91440), inset("bIns", 45720))
    return TextStyle(size_pt=size, family=family, line_spacing=spacing, insets_in=insets)


_TITLE_PH_TYPES = frozenset({"title", "ctrTitle"})


def inherited_text_size(slide, shape_element) -> float | None:
    """Кегль первого уровня, который фигура получает по наследованию, когда
    у её run нет своего `sz`: свой `a:lstStyle` → плейсхолдер лейаута (по
    idx, затем по типу) → стили мастера (`p:titleStyle` для заголовка,
    `p:bodyStyle` для остального). `None`: не плейсхолдер и своего стиля
    нет.

    Нужен замеру: заголовок титула VK Education наследует от лейаута кегль
    крупнее того, что записан в профиле, и текст, «влезший» по профилю, на
    рендере вырастал вверх рамки с `anchor="b"` и налезал на логотип."""
    tx_body = shape_element.find(qn("p:txBody"))
    size = _lvl1_size(tx_body.find(qn("a:lstStyle")) if tx_body is not None else None)
    if size:
        return size
    ph = shape_element.find(".//" + qn("p:nvPr") + "/" + qn("p:ph"))
    if ph is None:
        return None
    ph_type, ph_idx = ph.get("type", "body"), ph.get("idx")
    layout = slide.slide_layout
    master = layout.slide_master
    for owner, by_idx in ((layout, True), (master, False)):
        size = _placeholder_size(owner._element, ph_type, ph_idx if by_idx else None)  # noqa: SLF001
        if size:
            return size
    styles = master._element.find(qn("p:txStyles"))  # noqa: SLF001
    if styles is None:
        return None
    style = styles.find(qn("p:titleStyle") if ph_type in _TITLE_PH_TYPES else qn("p:bodyStyle"))
    return _lvl1_size(style)


def _placeholder_size(root, ph_type: str, ph_idx: str | None) -> float | None:
    """Кегль первого уровня плейсхолдера того же idx (или типа) в лейауте
    или мастере: его `a:lstStyle`, иначе явный `sz` его первого run."""
    found = None
    for el in root.iter(qn("p:ph")):
        if ph_idx is not None and el.get("idx") == ph_idx:
            found = el
            break
        el_type = el.get("type", "body")
        same = el_type == ph_type or (el_type in _TITLE_PH_TYPES and ph_type in _TITLE_PH_TYPES)
        if found is None and same:
            found = el
    if found is None:
        return None
    body = found.getparent().getparent().getparent().find(qn("p:txBody"))
    if body is None:
        return None
    size = _lvl1_size(body.find(qn("a:lstStyle")))
    run_pr = body.find(".//" + qn("a:rPr"))
    if not size and run_pr is not None and run_pr.get("sz"):
        size = int(run_pr.get("sz")) / 100
    return size


def _lvl1_size(list_style) -> float | None:
    if list_style is None:
        return None
    def_rpr = list_style.find(qn("a:lvl1pPr") + "/" + qn("a:defRPr"))
    if def_rpr is not None and def_rpr.get("sz"):
        return int(def_rpr.get("sz")) / 100
    return None


def allow_wrap(shape_element) -> bool:
    """Включает перенос строк, если пример его запретил (`wrap="none"`):
    короткая подпись примера жила в одну строку, а наш текст длиннее и без
    переноса уехал бы вбок через соседние карточки. Возвращает, пришлось
    ли включать."""
    tx_body = shape_element.find(qn("p:txBody"))
    body_pr = tx_body.find(qn("a:bodyPr")) if tx_body is not None else None
    if body_pr is None or body_pr.get("wrap") != "none":
        return False
    body_pr.set("wrap", "square")
    return True


def set_text_size(shape_element, size_pt: float) -> None:
    """Проставляет один кегль всем run, переносам и концам абзацев фигуры
    после ужимания текст не должен остаться пёстрым по размеру."""
    tx_body = shape_element.find(qn("p:txBody"))
    if tx_body is None:
        return
    value = str(int(round(size_pt * 100)))
    # run без своего a:rPr (текст пустого плейсхолдера наследует всё)
    # получает его: иначе кегль некуда записать.
    for run in tx_body.iter(qn("a:r")):
        if run.find(qn("a:rPr")) is None:
            run.insert(0, etree.Element(qn("a:rPr")))
    for tag in ("a:rPr", "a:endParaRPr"):
        for el in tx_body.iter(qn(tag)):
            el.set("sz", value)


# ---------------------------------------------------------------------------
# Родная таблица примера
# ---------------------------------------------------------------------------

# Атрибуты объединения ячеек. Объединение примера относится к его данным
# («Обсуждение доклада» на три столбца), к нашим строкам оно не подходит.
_MERGE_ATTRS = ("gridSpan", "rowSpan", "hMerge", "vMerge")

# Поля ячейки по умолчанию OOXML (`a:tcPr`), EMU: 0.1″ по бокам, 0.05″
# сверху и снизу.
_CELL_MAR_DEFAULTS = {"marL": 91440, "marR": 91440, "marT": 45720, "marB": 45720}


@dataclass(frozen=True)
class CellStyle:
    """Чем будет нарисован текст ячейки: нужен замеру высоты строки."""
    size_pt: float | None
    family: str | None
    line_spacing: float | None
    margins_in: tuple[float, float, float, float]  # слева, сверху, справа, снизу


class _CellText:
    def __init__(self, text: str) -> None:
        self.text = text
        self.bullet = False


def native_table(frame_element):
    """`a:tbl` рамки `p:graphicFrame` или `None`, если в рамке не таблица."""
    if frame_element is None or frame_element.tag != qn("p:graphicFrame"):
        return None
    return frame_element.find(qn("a:graphic") + "/" + qn("a:graphicData") + "/" + qn("a:tbl"))


def _cells(tr) -> list:
    return tr.findall(qn("a:tc"))


def _mostly(flags: list[bool]) -> bool:
    return bool(flags) and sum(flags) * 2 > len(flags)


def fill_native_table(
    frame_element, rows: list[list[str]], *, is_highlight: Callable[[object], bool] = lambda _tc: False,
    align: list[str] | None = None,
) -> None:
    """Заполняет родную таблицу примера нашими строками (первая: шапка).

    Правится таблица примера на месте, а не рисуется своя: заливки, линии,
    стиль таблицы, кегль и цвет ячеек остаются дизайнерскими. Число
    столбцов и строк подгоняется под данные: лишние удаляются, недостающие
    клонируются с последнего столбца и последней строки ТЕЛА (шапка
    остаётся одна и остаётся шапкой). Строки и столбцы, выделенные
    дизайнером (`is_highlight` у большинства их ячеек: тёмная заливка
    «Итого», строка-акцент), в нашу таблицу не переносятся: выделение
    относилось к данным примера, а наша первая строка тела от него стала
    бы второй шапкой. Если выделено всё, берётся как есть.

    Текст ячейки заменяется с сохранением `a:rPr` её первого run (тот же
    приём, что `bind_text`); ячейка без своего кегля (пустой угол шапки)
    получает его у соседки по строке, иначе наш текст в ней вышел бы
    кеглем по умолчанию 18 pt. `align` ("l"/"ctr"/"r" по столбцам)
    переписывает выравнивание абзацев: у примера шапка бывает выровнена
    вправо целиком, под числа, и наш текстовый первый столбец «Показатель»
    повис бы справа. Геометрию (ширины, высоты, рамку) ставит
    `set_native_table_geometry`."""
    tbl = native_table(frame_element)
    if tbl is None:
        raise ValueError("fill_native_table: в рамке нет a:tbl")
    if not rows or not any(rows[0]):
        raise ValueError("fill_native_table: у таблицы нет шапки")
    grid = tbl.find(qn("a:tblGrid"))
    trs = tbl.findall(qn("a:tr"))
    if grid is None or not trs:
        raise ValueError("fill_native_table: у таблицы примера нет сетки или строк")
    n_cols = max(len(r) for r in rows)
    n_body = len(rows) - 1

    _unmerge(trs)

    grid_cols = grid.findall(qn("a:gridCol"))
    body = trs[1:]
    col_flags = [
        _mostly([is_highlight(_cells(tr)[i]) for tr in body if i < len(_cells(tr))])
        for i in range(len(grid_cols))
    ]
    plain_cols = [i for i, hl in enumerate(col_flags) if not hl] or list(range(len(grid_cols)))
    order = plain_cols[:n_cols]
    order += [plain_cols[-1]] * (n_cols - len(order))
    new_cols = [copy.deepcopy(grid_cols[i]) for i in order]
    for gc in grid_cols:
        grid.remove(gc)
    for gc in new_cols:
        grid.append(gc)
    for tr in trs:
        tcs = _cells(tr)
        fresh = [copy.deepcopy(tcs[min(i, len(tcs) - 1)]) for i in order]
        for tc in tcs:
            tr.remove(tc)
        tail = tr.find(qn("a:extLst"))
        for tc in fresh:
            _insert_before(tr, tc, tail)

    header, body = trs[0], trs[1:]
    plain_rows = [tr for tr in body if not _mostly([is_highlight(tc) for tc in _cells(tr)])] or body
    kept = plain_rows[:n_body]
    proto = copy.deepcopy(plain_rows[-1] if plain_rows else header)
    for tr in body:
        if not any(tr is k for k in kept):
            tbl.remove(tr)
    last = kept[-1] if kept else header
    for _ in range(n_body - len(kept)):
        new_tr = copy.deepcopy(proto)
        last.addnext(new_tr)
        last = new_tr

    for tr, values in zip(tbl.findall(qn("a:tr")), rows):
        tcs = _cells(tr)
        row_rpr = next((r for tc in tcs for r in tc.iter(qn("a:rPr")) if r.get("sz")), None)
        for c, tc in enumerate(tcs):
            _fill_cell(tc, str(values[c]) if c < len(values) else "", row_rpr)
            if align is not None and c < len(align):
                for p in tc.iter(qn("a:p")):
                    p_pr = p.find(qn("a:pPr"))
                    if p_pr is None:
                        p_pr = etree.Element(qn("a:pPr"))
                        p.insert(0, p_pr)
                    p_pr.set("algn", align[c])


def _unmerge(trs: list) -> None:
    """Снимает объединения ячеек. Ячейка-продолжение (`hMerge`/`vMerge`)
    получает оформление ячейки, в которую была влита: на экране это была
    одна ячейка, и после разъединения полоса не должна стать пёстрой."""
    above: list = []
    for tr in trs:
        row: list = []
        for i, tc in enumerate(_cells(tr)):
            origin = None
            if tc.get("hMerge") and row:
                origin = row[-1]
            elif tc.get("vMerge") and i < len(above):
                origin = above[i]
            if origin is not None:
                src = origin.find(qn("a:tcPr"))
                own = tc.find(qn("a:tcPr"))
                if src is not None:
                    if own is not None:
                        tc.remove(own)
                    tc.append(copy.deepcopy(src))
            for attr in _MERGE_ATTRS:
                tc.attrib.pop(attr, None)
            row.append(origin if origin is not None else tc)
        above = row


def _fill_cell(tc, text: str, row_rpr) -> None:
    tx_body = tc.find(qn("a:txBody"))
    if tx_body is None:
        tx_body = etree.Element(qn("a:txBody"))
        etree.SubElement(tx_body, qn("a:bodyPr"))
        etree.SubElement(tx_body, qn("a:lstStyle"))
        etree.SubElement(tx_body, qn("a:p"))
        tc.insert(0, tx_body)
    run = tx_body.find(".//" + qn("a:r"))
    if row_rpr is not None:
        if run is None:
            p = tx_body.find(qn("a:p"))
            if p is None:
                p = etree.SubElement(tx_body, qn("a:p"))
            run = etree.Element(qn("a:r"))
            etree.SubElement(run, qn("a:t"))
            _insert_before(p, run, p.find(qn("a:endParaRPr")))
        r_pr = run.find(qn("a:rPr"))
        if r_pr is None:
            run.insert(0, copy.deepcopy(row_rpr))
        elif not r_pr.get("sz"):
            r_pr.set("sz", row_rpr.get("sz"))
    _bind_tx_body(tx_body, [_CellText(text)], None)


def table_cell_styles(frame_element) -> list[list[CellStyle]]:
    """Стиль каждой ячейки таблицы по строкам (после `fill_native_table`)."""
    tbl = native_table(frame_element)
    out = []
    for tr in tbl.findall(qn("a:tr")) if tbl is not None else []:
        row = []
        for tc in _cells(tr):
            r_pr = tc.find(".//" + qn("a:r") + "/" + qn("a:rPr"))
            size = int(r_pr.get("sz")) / 100 if r_pr is not None and r_pr.get("sz") else None
            family = None
            latin = r_pr.find(qn("a:latin")) if r_pr is not None else None
            if latin is not None and latin.get("typeface") and not latin.get("typeface").startswith("+"):
                family = latin.get("typeface")
            pct = tc.find(".//" + qn("a:pPr") + "/" + qn("a:lnSpc") + "/" + qn("a:spcPct"))
            spacing = int(pct.get("val")) / 100000 if pct is not None and pct.get("val") else None
            tc_pr = tc.find(qn("a:tcPr"))

            def mar(name: str, pr=tc_pr) -> float:
                raw = pr.get(name) if pr is not None else None
                return (int(raw) if raw is not None else _CELL_MAR_DEFAULTS[name]) / 914400

            row.append(CellStyle(size, family, spacing, (mar("marL"), mar("marT"), mar("marR"), mar("marB"))))
        out.append(row)
    return out


def template_row_heights_emu(frame_element) -> list[int]:
    """Высоты строк таблицы, как их задал пример (`a:tr/@h`, EMU)."""
    tbl = native_table(frame_element)
    return [int(tr.get("h") or 0) for tr in tbl.findall(qn("a:tr"))] if tbl is not None else []


def set_table_text_size(frame_element, row_sizes_pt: list[float | None]) -> None:
    """Проставляет кегль всем run и концам абзацев каждой строки таблицы
    (`None`: строка остаётся как есть)."""
    tbl = native_table(frame_element)
    for tr, size in zip(tbl.findall(qn("a:tr")), row_sizes_pt):
        if size is None:
            continue
        value = str(int(round(size * 100)))
        for run in tr.iter(qn("a:r")):
            if run.find(qn("a:rPr")) is None:
                run.insert(0, etree.Element(qn("a:rPr")))
        for tag in ("a:rPr", "a:endParaRPr"):
            for el in tr.iter(qn(tag)):
                el.set("sz", value)


def set_native_table_geometry(
    frame_element, box: Box, col_widths_emu: list[int], row_heights_emu: list[int], canvas: Canvas,
) -> None:
    """Ширины столбцов, высоты строк и рамка `p:graphicFrame` одним
    движением. Рамка обязана совпасть с суммой столбцов и строк: у
    примеров из Google Slides она записана как 3 000 000 × 3 000 000 EMU
    при таблице втрое шире, а аудит (L01/L02) видит именно рамку."""
    tbl = native_table(frame_element)
    for gc, w in zip(tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol")), col_widths_emu):
        gc.set("w", str(int(w)))
    for tr, h in zip(tbl.findall(qn("a:tr")), row_heights_emu):
        tr.set("h", str(int(h)))
    xfrm = frame_element.find(qn("p:xfrm"))
    if xfrm is None:
        xfrm = etree.Element(qn("p:xfrm"))
        frame_element.find(qn("p:nvGraphicFramePr")).addnext(xfrm)
    off = xfrm.find(qn("a:off"))
    if off is None:
        off = etree.SubElement(xfrm, qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    if ext is None:
        ext = etree.SubElement(xfrm, qn("a:ext"))
    off.set("x", str(round(box.left * canvas.width_emu)))
    off.set("y", str(round(box.top * canvas.height_emu)))
    ext.set("cx", str(int(sum(col_widths_emu))))
    ext.set("cy", str(int(sum(row_heights_emu))))


# ---------------------------------------------------------------------------
# Удаление лишнего
# ---------------------------------------------------------------------------


def set_shape_box(element, box: Box, canvas: Canvas) -> None:
    """Пишет фигуре её собственный `a:xfrm` по коробке в долях холста.
    Плейсхолдер без `a:xfrm` (наследовал коробку от лейаута) получает
    явный: иначе сузить его нельзя."""
    sp_pr = element.find(qn("p:spPr"))
    if sp_pr is None:
        sp_pr = etree.SubElement(element, qn("p:spPr"))
    xfrm = sp_pr.find(qn("a:xfrm"))
    if xfrm is None:
        xfrm = etree.Element(qn("a:xfrm"))
        sp_pr.insert(0, xfrm)
    off = xfrm.find(qn("a:off"))
    if off is None:
        off = etree.SubElement(xfrm, qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    if ext is None:
        ext = etree.SubElement(xfrm, qn("a:ext"))
    off.set("x", str(int(round(box.left * canvas.width_emu))))
    off.set("y", str(int(round(box.top * canvas.height_emu))))
    ext.set("cx", str(int(round(box.width * canvas.width_emu))))
    ext.set("cy", str(int(round(box.height * canvas.height_emu))))


def remove_shape(element) -> None:
    """Удаляет фигуру; опустевшая группа уходит вслед за ней, иначе в файле
    остаётся невидимая пустая рамка группы, которую ловит аудит."""
    parent = element.getparent()
    if parent is None:
        return
    parent.remove(element)
    while parent is not None and parent.tag == qn("p:grpSp"):
        if any(child.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")) for child in parent):
            break
        grand = parent.getparent()
        if grand is None:
            break
        grand.remove(parent)
        parent = grand


def prune_unfilled(
    slide, decor_unfilled: Iterable[DecorShape], slots_unfilled: Iterable[PatternSlot], canvas: Canvas,
    *, keep: Iterable = (), protect: Iterable[Box] = (),
) -> int:
    """Убирает из клона декор незаполненных единиц повтора и фигуры
    незаполненных слотов. Незаполненный текстовый слот удаляется целиком, а
    не опустошается: пустая рамка с текстом примера: это ровно та
    заглушка, которую ловит аудит I02, а пустая без текста: пустая
    карточка на слайде. `keep`: фигуры, в которые уже лёг текст: они не
    удаляются, даже если коробка совпала (карточка стоит на своей плашке).
    `protect`: коробки декора, который остаётся: у автофигуры PowerPoint
    почти всегда есть пустой `p:txBody`, и плашка заполненной карточки,
    совпавшая коробкой с пустым слотом, иначе ушла бы вместе с ним.
    Фигура ищется сначала по id исходной фигуры (`clone_map`), и тогда
    удаляется ровно она; коробка остаётся запасным путём для декора и
    слотов без id или с id, которого в клоне нет. `protect` нужен только
    запасному пути: фигура, найденная по id, чужой плашкой быть не может.
    Возвращает, сколько фигур удалено."""
    keep_ids = {id(el) for el in keep}
    refs = slide_refs(slide, canvas)
    by_id = clone_map(refs)
    decor_boxes: list[Box] = []
    slot_items: list[PatternSlot] = []
    by_id_hits: dict[int, bool] = {}  # id(элемента) → это слот (иначе декор)
    for d in decor_unfilled:
        ref = by_id.get(d.source_shape_id) if d.source_shape_id else None
        if ref is not None and ref.kind == d.kind:
            by_id_hits.setdefault(id(ref.element), False)
        else:
            decor_boxes.append(d.box)
    for s in slots_unfilled:
        ref = by_id.get(s.source_shape_id) if s.source_shape_id else None
        if ref is not None and _kind_fits(s.role, ref):
            by_id_hits[id(ref.element)] = True
        else:
            slot_items.append(s)
    protected = list(protect)
    removed = 0
    emptied_groups: list = []
    for ref in refs:
        if id(ref.element) in keep_ids or ref.element.getparent() is None:
            continue
        if id(ref.element) in by_id_hits:
            if by_id_hits[id(ref.element)]:
                group = _top_group(ref.element)
                if group is not None:
                    emptied_groups.append(group)
            remove_shape(ref.element)
            removed += 1
            continue
        if ref.box is None:
            continue
        hit = any(same_box(box, ref.box) for box in decor_boxes)
        slot_hit = False
        if not hit and not (
            not shape_text(ref.element).strip() and any(same_box(b, ref.box) for b in protected)
        ):
            slot_hit = hit = any(_kind_fits(s.role, ref) and same_box(s.box, ref.box) for s in slot_items)
        if hit:
            group = _top_group(ref.element) if slot_hit else None
            if group is not None:
                emptied_groups.append(group)
            remove_shape(ref.element)
            removed += 1
    # Карточка-группа, из которой ушли все надписи, остаётся пустой рамкой
    # с маркерами-точками и значками (VK Tech: четыре пустые карточки с
    # «+» и серыми точками на месте пунктов). Такая группа уходит целиком,
    # по тому же правилу, что и единица повтора: рисуется целиком или никак.
    for group in emptied_groups:
        if group.getparent() is None or any(shape_text(sp).strip() for sp in group.iter(qn("p:sp"))):
            continue
        if any(id(el) in keep_ids for el in group.iter()):
            continue
        remove_shape(group)
        removed += 1
    return removed


def _top_group(element):
    """Группа верхнего уровня (прямой ребёнок `p:spTree`), в которой лежит
    фигура, или `None`, если фигура сама на верхнем уровне."""
    top = None
    parent = element.getparent()
    while parent is not None and parent.tag == qn("p:grpSp"):
        top = parent
        parent = parent.getparent()
    return top


def shape_text(element) -> str:
    tx_body = element.find(qn("p:txBody"))
    if tx_body is None:
        return ""
    return "".join(t.text or "" for t in tx_body.iter(qn("a:t")))


def remove_stray_text(slide, canvas: Canvas, *, keep: Iterable = (), badge_boxes: Iterable[Box] = ()) -> int:
    """Убирает фигуры, в которых остался текст примера, но которые не
    стали ни слотом, ни значком декора: майнинг не превращает в слот
    каждую надпись (перекрытые, за полями, слишком мелкие), а в клоне они
    остались бы текстом-рыбой шаблона. Значки (номер шага в кружке) живут
    в декоре и сохраняются по коробке; самозаполняемые плейсхолдеры (номер
    слайда, дата, колонтитул) не трогаются."""
    keep_ids = {id(el) for el in keep}
    badges = list(badge_boxes)
    removed = 0
    for ref in slide_refs(slide, canvas):
        if ref.kind != "shape" or id(ref.element) in keep_ids:
            continue
        if not shape_text(ref.element).strip():
            continue
        if ref.is_placeholder and ref.ph_type in _AUTO_PH_TYPES:
            continue
        if ref.box is not None and any(same_box(b, ref.box) for b in badges):
            continue
        remove_shape(ref.element)
        removed += 1
    return removed


def remove_sample_frames(slide, *, keep: Iterable = ()) -> int:
    """Убирает таблицы и графики примера: в них данные образца («Показатель
    1», «Категория 2»), к содержанию колонки они отношения не имеют. Свои
    таблицу и график сборка кладёт отдельно (`builder._place_visual`).
    `keep`: рамки, уже заполненные нашими данными (родная таблица примера,
    `fill_native_table`)."""
    keep_ids = {id(el) for el in keep}
    removed = 0
    for frame in list(slide._element.iter(qn("p:graphicFrame"))):  # noqa: SLF001
        if id(frame) in keep_ids:
            continue
        data = frame.find(qn("a:graphic") + "/" + qn("a:graphicData"))
        uri = data.get("uri", "") if data is not None else ""
        if uri.endswith("/table") or uri.endswith("/chart"):
            remove_shape(frame)
            removed += 1
    return removed


def remove_in_box(slide, box: Box, canvas: Canvas, *, keep: Iterable = ()) -> int:
    """Убирает всё, что стоит ровно на коробке `box`: место под график,
    который сборка кладёт поверх, освобождается от картинки-образца."""
    keep_ids = {id(el) for el in keep}
    removed = 0
    for ref in slide_refs(slide, canvas):
        if ref.box is not None and id(ref.element) not in keep_ids and same_box(box, ref.box):
            remove_shape(ref.element)
            removed += 1
    return removed


# ---------------------------------------------------------------------------
# Фото пользователя
# ---------------------------------------------------------------------------


def replace_picture(slide, pic_element, data: bytes, frame_box: Box, canvas: Canvas) -> None:
    """Подменяет картинку фигуры `p:pic` на `data`, не трогая саму фигуру
    (обрезка по форме, рамка, тень, эффекты шаблона остаются). Фото
    вписывается «с обрезкой» (`a:srcRect`) по пропорциям рамки: растянутое
    лицо хуже срезанного края. svg-двойник примера в `a:extLst` удаляется,
    иначе PowerPoint показал бы его вместо нового растра."""
    _image_part, rid = slide.part.get_or_add_image_part(io.BytesIO(data))
    blip_fill = pic_element.find(qn("p:blipFill"))
    blip = blip_fill.find(qn("a:blip")) if blip_fill is not None else None
    if blip is None:
        raise ValueError("replace_picture: у фигуры нет a:blip")
    blip.set(qn("r:embed"), rid)
    blip.attrib.pop(qn("r:link"), None)
    for ext in blip.findall(qn("a:extLst")):
        blip.remove(ext)

    old_rect = blip_fill.find(qn("a:srcRect"))
    if old_rect is not None:
        blip_fill.remove(old_rect)
    with Image.open(io.BytesIO(data)) as img:
        img_w, img_h = img.size
    frame_w = frame_box.width * canvas.width_emu
    frame_h = frame_box.height * canvas.height_emu
    if not (img_w and img_h and frame_w > 0 and frame_h > 0):
        return
    img_aspect, frame_aspect = img_w / img_h, frame_w / frame_h
    rect = etree.Element(qn("a:srcRect"))
    if img_aspect > frame_aspect:
        cut = round((1 - frame_aspect / img_aspect) / 2 * 100000)
        rect.set("l", str(cut))
        rect.set("r", str(cut))
    else:
        cut = round((1 - img_aspect / frame_aspect) / 2 * 100000)
        rect.set("t", str(cut))
        rect.set("b", str(cut))
    blip.addnext(rect)


# ---------------------------------------------------------------------------
# Имена частей пакета
# ---------------------------------------------------------------------------


def fix_duplicate_partnames(prs) -> int:
    """Даёт уникальные имена частям пакета, чьи имена совпали. Вызывается
    перед сохранением.

    Откуда совпадения. Картинки примеров после `_clear_sample_slides` ни
    на что не ссылаются, и python-pptx считает их имена свободными: новая
    картинка (фото пользователя, картинка старого пути) получает, скажем,
    `image16.jpg`. Потом следующий клон снова связывается со старой
    `image16.jpg` примера, и в архиве оказываются две записи с одним
    именем. PowerPoint такой файл открывает только через «восстановление».

    Переименование безопасно: связи и `[Content_Types].xml` python-pptx
    строит из `part.partname` в момент сохранения. Возвращает число
    переименованных частей."""
    parts = list(prs.part.package.iter_parts())
    taken = {str(part.partname) for part in parts}
    seen: set[str] = set()
    renamed = 0
    for part in parts:
        name = str(part.partname)
        if name not in seen:
            seen.add(name)
            continue
        match = _PARTNAME_RE.match(name)
        stem, ext = (match.group(1), match.group(3)) if match else (name, "")
        number = 1
        while f"{stem}{number}{ext}" in taken:
            number += 1
        new_name = f"{stem}{number}{ext}"
        part.partname = PackURI(new_name)
        taken.add(new_name)
        seen.add(new_name)
        renamed += 1
    return renamed
