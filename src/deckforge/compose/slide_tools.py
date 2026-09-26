"""Инструменты сборки слайда, которые отдаются агенту.

Зачем. До этого модель писала текст вслепую: она не знала ни какой кегль
несёт выбранная раскладка, ни какой высоты в ней текстовый блок, ни что
получится, когда код всё это уложит. Отсюда весь класс поломок, которые
мы ловили 23 сентября 2026: заголовок, упавший с 48 пунктов на 16, потому
что не влез; слайд, заполненный на 2% холста; список из одного пункта в
блоке размером в пол-слайда. Код знал про эти беды (аудит их честно
отмечал), но узнавал о них ПОСЛЕ того, как модель уже всё решила.

Эти функции переворачивают порядок: модель может СНАЧАЛА спросить, что
получится, и переписать текст, пока он ещё черновик.

Граница слоёв (`tests/plan/test_no_pptx_import.py`, докстрока `plan/
spec.py`). Модуль живёт в `compose/`, а не в `plan/`, потому что работает с
python-pptx и координатами. Модель он не зовёт ни разу — вызывает её
агентный цикл в `plan/writer.py`, для которого эти функции остаются чёрным
ящиком: на вход текст и плоские числа, на выход плоские числа. Ни одна
координата, ни один EMU наружу не возвращается — иначе граница «план не
знает ни одной координаты» протекла бы через ответ инструмента.

Почему инструменты СМЫСЛОВЫЕ, а не геометрические. Агенту не даётся ни
одного способа назвать координату, цвет или кегль: он выбирает раскладку по
её номеру и отдаёт текст. Пока он физически не может сказать «нарисуй
прямоугольник вот здесь», он не может и нарушить шаблон — а это
единственная гарантия, ради которой вся конструкция и строится (ТЗ:
решение адаптируется к произвольному шаблону, а не рисует поверх него).
"""
from __future__ import annotations
from dataclasses import replace
from pathlib import Path

from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.audit.deterministic import audit_slide_layout, slide_fill_ratio
from deckforge.compose.builder import (
    _clear_sample_slides, _pattern_from_model, _remove_last_slide, place_slide,
)
from deckforge.ooxml.geometry import Canvas
from deckforge.plan.spec import SlideSpec
from deckforge.template.profile import TemplateProfile

# Роли слотов, по которым видно, что раскладка умеет держать. Названия — те
# же, что мининг ставит в `PatternSlot.role`; наружу уходят человеческими
# словами (см. `_holds`), потому что список ролей — деталь разбора, а
# агенту нужен ответ на вопрос «влезет ли сюда картинка».
_HOLDS = {
    "bullets": ("bullet", "body"),
    "cards": ("card_title", "card_body"),
    "kpi": ("kpi_value", "kpi_label"),
    "quote": ("quote",),
    "image": ("image",),
    "chart": ("chart",),
    "table": ("table",),
    "subhead": ("subhead",),
    "source": ("source",),
}


def _holds(pattern) -> list[str]:
    roles = {slot.role for slot in pattern.slots}
    return sorted(name for name, needed in _HOLDS.items() if roles.intersection(needed))


def _slot_guide(pattern) -> list[dict]:
    """Что писать в каждое место раскладки: схема слотов, которую модель
    сняла при разборе шаблона (`purpose`, `content_hint`, `max_words`).
    Без неё писатель знал только вместимость в знаках и не мог понять, что
    узкое место над карточкой — номер шага, а широкое под ним — описание.
    Номера и постоянный текст шаблона не отдаются: писать туда нечего.
    Одинаковые описания карточек одного ряда сливаются в одно, иначе
    шесть карточек дали бы шесть одинаковых строк. Без схемы — пусто."""
    guide: list[dict] = []
    seen: set[tuple] = set()
    for slot in pattern.slots:
        if slot.keeps_sample_text or slot.purpose is None:
            continue
        key = (slot.purpose, slot.content_hint, slot.max_words)
        if key in seen:
            continue
        seen.add(key)
        guide.append({"purpose": slot.purpose, "content_hint": slot.content_hint, "max_words": slot.max_words})
    return guide


def list_layouts(profile: TemplateProfile, *, kind: str | None = None) -> list[dict]:
    """Каталог раскладок шаблона в том виде, в каком его читает агент.

    `kind` сужает список до одного вида; без него отдаются все — агент
    вправе посмотреть, что вообще умеет шаблон, прежде чем решать, в какой
    форме подавать материал.

    Наружу уходит только то, что помогает выбрать: номер раскладки, её вид,
    что она держит, сколько элементов и знаков вмещает и тёмная ли она
    (влияет на то, как будет читаться текст). Координат, цветов и кеглей в
    ответе нет — их выбирает код."""
    out: list[dict] = []
    for model in profile.patterns:
        if kind is not None and model.kind != kind:
            continue
        pattern = _pattern_from_model(model)
        cap = pattern.capacity
        out.append({
            "layout_id": pattern.pattern_id,
            "kind": pattern.kind,
            "holds": _holds(pattern),
            "max_items": cap.max_items,
            "max_bullets": cap.max_bullets,
            "max_chars_per_item": cap.max_chars_per_item,
            "max_rows": cap.max_rows,
            "max_cols": cap.max_cols,
            "max_series": cap.max_series,
            "is_dark": pattern.is_dark,
            "places": _slot_guide(pattern),
        })
    out.sort(key=lambda row: (row["kind"], -row["max_items"], row["layout_id"]))
    return out


def try_slide(
    slide_spec: SlideSpec, layout_id: str, profile: TemplateProfile, template_path: Path,
    *, bullet_char: str = "•",
) -> dict:
    """Собирает слайд начерно, показывает вердикт и выбрасывает черновик.

    Это главный инструмент: агент отдаёт написанный текст и номер
    раскладки, получает обратно ответ на вопрос «что получилось» — и может
    переписать текст короче или взять другую раскладку, ПОКА слайд ещё не
    в презентации.

    Возвращает:
    - `ok` — прошёл ли слайд строгие проверки без находок уровня ошибки;
    - `fill_percent` — сколько процентов холста занято содержанием (ТЗ
      считает браком меньше 25% и больше 75%, отсюда и число);
    - `findings` — находки строгих проверок человеческими словами;
    - `notes` — что код был вынужден сделать с текстом сам (усечение,
      невлезший блок, отсутствующий слот) — то, о чём агент иначе не
      узнает вовсе.

    Черновик рисуется в ОТДЕЛЬНОЙ презентации, открытой от того же
    шаблона, и удаляется до возврата: инструмент ничего не меняет в колоде,
    сколько бы раз его ни звали. Цена — открытие .pptx на вызов; шаблоны
    тяжёлые (до 18 МБ), поэтому звать его на каждое слово не стоит, и это
    сказано агенту в его задании."""
    pattern = _find_pattern(profile, layout_id)
    if pattern is None:
        return {
            "ok": False, "fill_percent": 0, "findings": [],
            "notes": [f"раскладки {layout_id!r} в этом шаблоне нет — возьми номер из списка раскладок"],
        }

    canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)
    audit_config = AuditConfig.load()

    prs = Presentation(str(template_path))
    _clear_sample_slides(prs)
    # Черновик пишет находки в СВОЮ копию спека: инструмент вправе быть
    # вызванным десять раз подряд, и накопить десять одинаковых находок в
    # том спеке, который потом пойдёт в колоду, он не должен.
    trial = replace(slide_spec, findings=[])
    place_slide(prs, trial, pattern, profile, audit_config, bullet_char=bullet_char)
    errors = audit_slide_layout(prs.slides[-1], canvas, profile, audit_config, index=slide_spec.index)
    # Заполненность берётся ТЕМ ЖЕ расчётом, что и проверка D05 (`audit.
    # deterministic.slide_fill_ratio`), а не оценкой `fits()` по спеку до
    # отрисовки: первая же живая проверка инструмента показала, что они
    # расходятся вдвое (7% против 14%), и агент правил бы слайд по числу,
    # которого аудит потом не подтвердит.
    fill = slide_fill_ratio(prs.slides[-1], canvas, profile, index=slide_spec.index)
    _remove_last_slide(prs)

    return {
        "ok": not errors,
        "fill_percent": round(fill * 100),
        "findings": [f"{f.check_id}: {f.message}" for f in errors],
        "notes": list(trial.findings),
    }


def _find_pattern(profile: TemplateProfile, layout_id: str):
    for model in profile.patterns:
        if model.pattern_id == layout_id:
            return _pattern_from_model(model)
    return None
