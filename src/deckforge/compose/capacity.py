"""Вместимость места раскладки тем же замером, которым клон решает,
лёг ли текст (задача V5).

Раньше предел места в контракте писателя и проверка клона считались
по-разному. Контракт брал `PatternSlot.max_chars`: площадную оценку
(`patterns.estimate_slot_chars`: 0,55 em на знак, межстрочный 1,25, вся
рамка без полей, кегль из профиля). Клон мерил `textfit.measure` реальным
кеглем фигуры (у заголовка он унаследован от лейаута), гарнитурой run
(тело карточек VK Education набрано Arial, а не Play шаблона), межстрочным
абзаца и рамкой за вычетом полей. К тому же заголовок карточки без своего
места клон ставит жирной первой строкой тела, а контракт отдавал телу всю
рамку. Писатель укладывался в контракт, а клон отказывал `FONT_BUDGET`, и
лестница уходила на запасные раскладки (живой прогон 28 сентября 2026:
dense slide42 x3, slide17 x4, ни одного клона выбранной раскладки).

Здесь одна функция `clone_text_frame` описывает, чем и где клон нарисует
текст, и её зовут оба: `builder._fit_cloned_text` при сборке и
`measured_fits` при планировании. Наружу, в `pattern/` и `plan/`, уходят
только плоские числа (`SlotFit`): слова и знаки, без координат.

Шаблон читается по `TemplateProfile.source_path`. Нет файла (профиль
собран в тесте руками, шаблон удалён после загрузки): замера нет, и
форма раскладки берёт прежнюю оценку профиля."""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path

from deckforge.compose.clone import inherited_text_size, match_slots, sample_slides_by_number, text_style
from deckforge.compose.failure import FontBudget, font_budget
from deckforge.compose.textfit import measure
from deckforge.ooxml.geometry import Canvas
from deckforge.template.typography import fill_scale_gaps

EMU_PER_INCH = 914400
# Межстрочный, когда ни фигура, ни её предки его не задают: тот же, что у
# аудита L03 и у `builder._fit_cloned_text`.
DEFAULT_LINE_SPACING = 1.2
# Допуск сравнения высоты текста с рамкой, дюймы (`builder._FIT_TOLERANCE_IN`).
FIT_TOLERANCE_IN = 0.01
# Ступени шкалы, по которым ужимается текст, без "micro": пол "caption".
_SHRINK_STEPS = ("display", "h1", "h2", "body", "caption")

_TEXT_ROLES = frozenset({"headline", "subhead", "body", "bullet", "card_body", "card_title", "quote",
                         "kpi_value", "kpi_label", "caption"})
_UNIT_BODY_ROLES = frozenset({"card_body", "bullet", "body"})

# Текст, которым меряется вместимость: деловая русская проза средней
# длины слова (6,5 знака со пробелом, как у `CHARS_PER_WORD` 7), без
# длинных терминов, которые переносились бы раньше обычного. Перенос по
# словам зависит от длины слов, поэтому мерить надо похожим текстом, а не
# «ааааа».
_FILLER = (
    "Заявка ждёт согласования почти весь срок, а работа над ней занимает "
    "меньше часа. После запуска маршрутизации медиана сократилась до шести "
    "часов, ошибок стало втрое меньше, и отдел закупок получил отчёт по "
    "каждому этапу. Следующий шаг: правила подмены на время отпуска, обучение "
    "согласующих и контроль сроков в одном окне для всех подразделений."
).split()
# Заголовок карточки, который клон ставит жирной первой строкой тела,
# когда своего места у заголовка нет: два слова, как пишет писатель.
_TITLE_FILLER = "Ошибки маршрутизации"


@dataclass(frozen=True)
class TextFrame:
    """Чем и где клон нарисует текст фигуры: кегль примера (свой у run или
    унаследованный), гарнитура, межстрочный, рамка без полей в дюймах и
    кегли, до которых клону разрешено ужаться (`sizes`, первый: кегль
    примера)."""
    size_pt: float
    family: str
    line_spacing: float
    width_in: float
    height_in: float
    sizes: tuple[float, ...]

    @property
    def allowed_in(self) -> float:
        """Высота, в которой текст считается помещающимся. Рамка бывает ниже
        одной строки собственного кегля (пустой плейсхолдер заголовка
        ЛЦТ2026: 0,41" при 20 pt): PowerPoint рисует строку поверх рамки,
        и одна строка кеглем примера в ней считается помещающейся."""
        one_line = measure("Xg", self.family, self.size_pt, self.width_in, line_spacing=self.line_spacing).height_in
        return max(self.height_in, one_line)

    def fitting_size(self, text: str) -> float | None:
        """Самый крупный разрешённый кегль, которым `text` ложится, или
        `None`, если не ложится ни одним."""
        allowed = self.allowed_in + FIT_TOLERANCE_IN
        for size in self.sizes:
            if measure(text, self.family, size, self.width_in, line_spacing=self.line_spacing).height_in <= allowed:
                return size
        return None


def shrink_sequence(profile, slot_size_pt: float) -> list[float]:
    """Кегли-кандидаты по убыванию, начиная с кегля слота, затем ступени
    шкалы шаблона не выше него и не ниже подписи. Кегли профиля
    нормированы к эталонному холсту и денормируются здесь один раз
    (`TemplateProfile.denorm_pt`/`type_scale_pt`)."""
    size_pt = profile.denorm_pt(slot_size_pt)
    raw_steps = {name: profile.type_scale_pt(name, 0.0) for name in _SHRINK_STEPS}
    # Пол тот же, что у автопочинки (`TemplateProfile.min_font_pt`): сборка и
    # починка не должны расходиться в том, какой кегль ещё читается.
    caption_pt = profile.min_font_pt()

    seq = [size_pt]
    for name in _SHRINK_STEPS:
        value = raw_steps[name]
        if 0 < value < seq[-1] - 0.05:
            seq.append(value)
    if caption_pt > 0 and abs(seq[-1] - caption_pt) > 0.05:
        seq.append(caption_pt)
    # Между далёкими ступенями (44 -> 16 у VK Education) промежуточные
    # кегли, иначе заголовок ужимается сразу до кегля текста.
    filled = [v for v in fill_scale_gaps(seq) if v <= size_pt + 0.05 and v >= caption_pt - 0.05]
    return sorted(set(filled) | {size_pt}, reverse=True)


def clone_text_frame(
    slide, element, box, profile, family: str, canvas: Canvas, slot_size_pt: float,
    *, budget: FontBudget | None,
) -> TextFrame | None:
    """Рамка текста фигуры `element` на слайде `slide` так, как её видит
    клон. `budget=None`: без предела ужимания, вся шкала до подписи.
    `None`: у рамки нет площади."""
    style = text_style(element)
    # Кегль, которым PowerPoint нарисует текст: свой у run, иначе
    # унаследованный от лейаута/мастера, и только если его нет, из профиля.
    size = style.size_pt or inherited_text_size(slide, element) or profile.denorm_pt(slot_size_pt)
    left_in, top_in, right_in, bottom_in = style.insets_in
    width_in = box.width * canvas.width_emu / EMU_PER_INCH - left_in - right_in
    height_in = box.height * canvas.height_emu / EMU_PER_INCH - top_in - bottom_in
    if width_in <= 0 or height_in <= 0 or size <= 0:
        return None
    sizes = [size] + [s for s in shrink_sequence(profile, slot_size_pt) if s < size - 0.05]
    if budget is not None:
        floor = budget.floor_pt(size, sizes[1:])
        sizes = [s for s in sizes if s >= floor - 0.05] or [size]
    return TextFrame(
        size_pt=size, family=style.family or family, line_spacing=style.line_spacing or DEFAULT_LINE_SPACING,
        width_in=width_in, height_in=height_in, sizes=tuple(sizes),
    )


# ---------------------------------------------------------------------------
# Вместимость раскладок шаблона
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class SlotFit:
    """Сколько текста место держит по замеру клона. `words`/`chars`:
    предел, который клон ещё примет, ужав кегль в пределах бюджета;
    `native_words`/`native_chars`: кеглем примера, без ужимания.
    `titled`: замер оставил первую строку под жирный заголовок карточки
    (своего места у заголовка в единице нет)."""
    words: int
    chars: int
    native_words: int
    native_chars: int
    titled: bool = False


def _prefix(words: int) -> str:
    return " ".join(_FILLER[i % len(_FILLER)] for i in range(words))


def _fit_words(frame: TextFrame, *, native: bool, title: str | None) -> int:
    """Наибольшее число слов текста-образца, которое ложится в рамку.
    Высота текста растёт с числом слов монотонно, поэтому двоичный поиск."""
    sizes = frame.sizes[:1] if native else frame.sizes
    capped = TextFrame(frame.size_pt, frame.family, frame.line_spacing, frame.width_in, frame.height_in, sizes)

    def fits(n: int) -> bool:
        body = _prefix(n)
        return capped.fitting_size(f"{title}\n{body}" if title else body) is not None

    lo, hi = 0, 8
    while fits(hi) and hi < 400:
        lo, hi = hi, hi * 2
    while hi - lo > 1:
        mid = (lo + hi) // 2
        if fits(mid):
            lo = mid
        else:
            hi = mid
    return lo


def _titled(pattern, slot) -> bool:
    """Тело единицы повтора, у которой нет места под заголовок: заголовок
    карточки клон поставит жирной первой строкой этого же места
    (`blocks.assign_content_with_drops`)."""
    repeat = pattern.repeat
    if repeat is None or repeat.count < 2 or slot.role not in _UNIT_BODY_ROLES:
        return False
    roles = set(repeat.slot_roles)
    return slot.role in roles and "card_title" not in roles


def _slot_fit(frame: TextFrame, titled: bool) -> SlotFit:
    title = _TITLE_FILLER if titled else None
    words = _fit_words(frame, native=False, title=title)
    native = _fit_words(frame, native=True, title=title)
    return SlotFit(
        words=words, chars=len(_prefix(words)), native_words=native, native_chars=len(_prefix(native)),
        titled=titled,
    )


def measured_fits(profile) -> dict[str, dict[int, SlotFit]]:
    """`pattern_id` → номер слота в `Pattern.slots` → `SlotFit` для всех
    текстовых мест, которые клон найдёт на слайде-примере. Пустой словарь:
    файла шаблона нет, мерить нечем."""
    path = Path(getattr(profile, "source_path", "") or "")
    if not path.name or not path.is_file():
        return {}
    try:
        stamp = path.stat().st_mtime_ns
    except OSError:
        return {}
    # Ключ: файл, его версия и слоты профиля. Два профиля одного файла
    # (разбор с моделью и без) делят замер, только если слоты совпали.
    signature = tuple(
        (p.pattern_id, tuple((s.role, s.source_shape_id, s.size_pt) for s in p.slots)) for p in profile.patterns
    )
    key = (str(path), stamp, signature)
    cached = _CACHE.get(key)
    if cached is None:
        cached = _measure_profile(profile, str(path))
        if len(_CACHE) >= 8:
            _CACHE.pop(next(iter(_CACHE)))
        _CACHE[key] = cached
    return cached


_CACHE: dict[tuple, dict[str, dict[int, SlotFit]]] = {}


def _measure_profile(profile, path: str) -> dict[str, dict[int, SlotFit]]:
    from pptx import Presentation

    # Отложенный импорт: `builder` сам зовёт этот модуль при сборке.
    from deckforge.compose.builder import _pattern_from_model, _primary_family

    canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)
    family = _primary_family(profile)
    budget = font_budget()
    try:
        prs = Presentation(path)
    except Exception:  # noqa: BLE001: битый или чужой файл: замера нет, форма возьмёт оценку профиля
        return {}
    sources = sample_slides_by_number(prs)
    result: dict[str, dict[int, SlotFit]] = {}
    for model in profile.patterns:
        if not model.source_slide_index:
            continue
        source = sources.get(model.source_slide_index[0])
        if source is None:
            continue
        pattern = _pattern_from_model(model)
        matched = match_slots(source, pattern.slots, canvas)
        fits: dict[int, SlotFit] = {}
        for i, slot in enumerate(pattern.slots):
            ref = matched.get(i)
            if ref is None or ref.box is None or slot.role not in _TEXT_ROLES:
                continue
            frame = clone_text_frame(
                source, ref.element, ref.box, profile, family, canvas, slot.size_pt,
                budget=budget.for_slot(slot.role, pattern.kind),
            )
            if frame is not None:
                fits[i] = _slot_fit(frame, _titled(pattern, slot))
        if fits:
            result[model.pattern_id] = fits
    return result
