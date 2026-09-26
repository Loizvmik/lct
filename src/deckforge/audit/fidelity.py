"""Метрики верности шаблону (Template Fidelity) — задача T,
`deckforge_updated_architecture.md`, раздел 16: набор ДЕТЕРМИНИРОВАННЫХ
чисел о том, насколько ГОТОВАЯ колода выглядит собранной в шаблоне, а не
рядом с ним. Дополняет, а не заменяет, два уже существующих аудита:
`audit.deterministic` (24 проверки одного слайда на нарушение вёрстки) и
`audit.visual` (оценка по картинке моделью) — оба ищут дефекты КОНКРЕТНОГО
слайда, здесь же — насколько колода в целом держится композиционного языка
шаблона (доля слайдов клоном, сохранность фигур примера, отклонение
геометрии, разнообразие раскладок, использование родных ассетов).

`template_fidelity(pptx_path, deck_spec, profile) -> FidelityReport` —
единственная публичная функция. Ничего не вызывает модель — метрики нужны
именно как детерминированный, воспроизводимый ориентир поверх/рядом с
вероятностной оценкой по картинке (см. раздел 16 брифа: "vision-based
оценка должна оставаться дополнительным уровнем, а не единственным
критерием качества").

Типографика/палитра переиспользуют результат `audit.deterministic.
run_deterministic` (T01/T02/T03), а не повторяют те же проверки другим
кодом — тот же принцип границ слоёв, что и везде в проекте."""
from __future__ import annotations

import hashlib
import math
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.audit.deterministic import run_deterministic, slide_fill_ratio
from deckforge.compose.builder import CLONE_MARK_PREFIX
from deckforge.compose.clone import clone_map, slide_refs
from deckforge.ooxml.geometry import Canvas
from deckforge.ooxml.ns import qn
from deckforge.ooxml.package import PptxPackage
from deckforge.ooxml.walk import walk_shapes
from deckforge.plan.spec import DeckSpec

# Роли слотов, чья коробка не текстовая (картинка/таблица/график) — бриф
# просит отклонение геометрии только «текстовых фигур клона», см. роли в
# `template/patterns.py::PatternSlot.role`/docs/architecture-deep-dive.md
# раздел 3.1 ("роль по типографической ступени...").
_NON_TEXT_SLOT_ROLES = frozenset({"image", "icon", "table", "chart"})

# Находки, которыми переиспользуется детерминированный аудит для метрики
# «типографика/палитра» — см. докстроку модуля.
_STYLE_CHECK_IDS = frozenset({"T01", "T02", "T03"})


@dataclass
class FidelityReport:
    """Все доли — числа 0..1 (или `None`, если метрика не посчитана — тогда
    причина в `.notes`, см. раздел 16 брифа: параллельная задача R ещё не
    добавила `source_density`, шаблон переехал на диске и т.п. — честная
    деградация одной метрики не должна ронять остальные)."""

    # Доля слайдов готовой колоды, собранных клоном примера (не с нуля).
    native_clone_rate: float
    # Доля фигур клонов (слоты + декор с `source_shape_id`), реально
    # найденных на итоговом слайде — сколько от намайненного паттерна
    # дожило до готового файла.
    native_shape_preservation: float | None
    # Доля текстовых фигур колоды без находок T01 (шрифт)/T02 (кегль)/T03
    # (цвет) детерминированного аудита.
    typography_palette_compliance: float | None
    # Среднее расстояние коробок текстовых фигур клона от коробок слотов
    # паттерна (доли холста: среднее |Δleft|+|Δtop|+|Δwidth|+|Δheight| / 4).
    mean_geometry_deviation: float | None
    # Число уникальных `pattern_id`, применённых в колоде (`SlideSpec.
    # pattern_id`), к числу слайдов.
    pattern_diversity: float | None
    # Энтропия Шеннона распределения `pattern_id` по слайдам, нормированная
    # к [0, 1] (1 — все раскладки использованы поровну, 0 — одна на всю
    # колоду).
    pattern_entropy: float | None
    # Доля картинок на слайдах (кроме декора без картинки), чьи байты
    # совпадают (md5) с медиа шаблона — то есть родная графика шаблона, а
    # не фото пользователя.
    native_asset_usage: float | None
    # |плотность готовой колоды - source_density паттерна|, усреднённая по
    # слайдам, у чьего паттерна это поле есть (параллельная задача R).
    density_delta: float | None
    summary: str
    notes: list[str] = field(default_factory=list)
    # Задача V4: доля площади содержательных объектов, которые человек может
    # править в PowerPoint (текст, родная таблица, родной график, фото
    # пользователя), против растра на месте содержания. Декор шаблона в
    # знаменатель не входит: его никто не редактирует и не должен.
    editable_content_coverage: float | None = None
    # Счётчики объектов по видам (`EDITABILITY_KINDS`).
    editable_counts: dict[str, int] = field(default_factory=dict)


# Виды объектов для метрики редактируемости. Первые четыре редактируемы,
# `template_decor` в метрику не входит, `raster_content` не редактируемо.
EDITABILITY_KINDS: tuple[str, ...] = (
    "editable_text", "native_table", "native_chart", "user_image", "template_decor", "raster_content",
)
_EDITABLE_KINDS = frozenset({"editable_text", "native_table", "native_chart", "user_image"})

# Картинка не из шаблона, закрывающая почти весь слайд, на котором больше
# нет ни текста, ни таблицы, ни графика: слайд собран одной картинкой, а
# это ровно то, что ТЗ запрещает. Порог с запасом на поля и обрезку.
_SLIDE_AS_IMAGE_AREA = 0.85

_CHART_URI_MARK = "drawingml/2006/chart"


def _clone_pattern_id(slide) -> str | None:
    """Читает метку пути сборки (`compose.clone.mark_slide`,
    `p:cSld/@name`) — `None`, если слайд собран не клоном (с нуля) или
    метки нет вовсе (посторонний файл)."""
    c_sld = slide._element.find(qn("p:cSld"))  # noqa: SLF001 — тот же приём, что и mark_slide/audit_slide_layout
    name = c_sld.get("name") if c_sld is not None else None
    if name and name.startswith(CLONE_MARK_PREFIX):
        return name[len(CLONE_MARK_PREFIX):]
    return None


def _native_shape_preservation_and_geometry(
    prs, canvas: Canvas, clone_pattern_by_index: dict[int, object], patterns_by_id: dict[str, object],
    notes: list[str],
) -> tuple[float | None, float | None]:
    shape_total = 0
    shape_hit = 0
    deviations: list[float] = []
    missing_patterns: set[str] = set()

    for index, pattern_id in clone_pattern_by_index.items():
        pattern = patterns_by_id.get(pattern_id)
        if pattern is None:
            missing_patterns.add(pattern_id)
            continue
        slide = prs.slides[index]
        refs = [r for r in slide_refs(slide, canvas) if r.box is not None]
        cmap = clone_map(refs)

        source_ids = [s.source_shape_id for s in pattern.slots if s.source_shape_id]
        source_ids += [d.source_shape_id for d in pattern.decor if d.source_shape_id]
        if source_ids:
            shape_total += len(source_ids)
            shape_hit += sum(1 for sid in source_ids if sid in cmap)

        for slot in pattern.slots:
            if slot.role in _NON_TEXT_SLOT_ROLES or not slot.source_shape_id:
                continue
            ref = cmap.get(slot.source_shape_id)
            if ref is None or ref.box is None:
                continue
            box = ref.box
            deviations.append((
                abs(box.left - slot.box.left) + abs(box.top - slot.box.top)
                + abs(box.width - slot.box.width) + abs(box.height - slot.box.height)
            ) / 4)

    if missing_patterns:
        notes.append(
            f"{len(missing_patterns)} раскладок клонов не найдены в профиле шаблона "
            f"({', '.join(sorted(missing_patterns))}) — профиль в кэше устарел относительно колоды."
        )
    if not clone_pattern_by_index:
        notes.append("Ни один слайд не собран клоном — сохранность фигур и отклонение геометрии не посчитаны.")

    preservation = shape_hit / shape_total if shape_total else None
    deviation = sum(deviations) / len(deviations) if deviations else None
    return preservation, deviation


def _typography_palette_compliance(
    pptx_path: Path, prs, canvas: Canvas, profile, config: AuditConfig, notes: list[str],
) -> float | None:
    findings = run_deterministic(pptx_path, profile, config)
    style_hits = {
        (f.slide_index, (f.shape_ref or "").split(":", 1)[0])
        for f in findings if f.check_id in _STYLE_CHECK_IDS
    }

    text_shape_keys: set[tuple[int, str]] = set()
    with PptxPackage.open(pptx_path) as pkg:
        for index, slide in enumerate(prs.slides):
            slide_part = str(slide.part.partname).lstrip("/")
            root = pkg.xml(slide_part)
            for ref in walk_shapes(root, canvas, include_groups=False):
                if ref.kind != "shape" or ref.box is None:
                    continue
                tx_body = ref.element.find(qn("p:txBody"))
                if tx_body is None:
                    continue
                text = "".join(t.text or "" for t in tx_body.iter(qn("a:t")))
                if text.strip():
                    text_shape_keys.add((index, ref.shape_id))

    if not text_shape_keys:
        notes.append("Типографика/палитра не посчитаны: в колоде нет ни одной текстовой фигуры.")
        return None
    clean = text_shape_keys - style_hits
    return len(clean) / len(text_shape_keys)


def _pattern_diversity_and_entropy(deck_spec: DeckSpec, notes: list[str]) -> tuple[float | None, float | None]:
    pattern_ids = [s.pattern_id for s in deck_spec.slides if s.pattern_id]
    if not pattern_ids:
        notes.append("Разнообразие раскладок не посчитано: ни один слайд плана не несёт pattern_id.")
        return None, None

    diversity = len(set(pattern_ids)) / len(pattern_ids)
    counts = Counter(pattern_ids)
    total = len(pattern_ids)
    raw_entropy = -sum((c / total) * math.log2(c / total) for c in counts.values())
    # Нормировка к [0, 1]: энтропия одного распределения на K корзин не
    # может превысить log2(K) — без нормировки колода из трёх слайдов трёх
    # разных раскладок и колода из пятнадцати слайдов пятнадцати разных
    # раскладок получили бы разные числа при одинаковом смысле «максимально
    # разнообразно». Одна уникальная раскладка на всю колоду — 0, не деление
    # на ноль (log2(1) = 0, ветка ниже не нужна, но log2(1)=0 и раз one bucket
    # даёт raw_entropy=0 — обе части дроби 0, оставляем 0 явно).
    max_entropy = math.log2(len(counts)) if len(counts) > 1 else 0.0
    entropy = raw_entropy / max_entropy if max_entropy > 0 else 0.0
    return diversity, entropy


def _template_media_hashes(profile) -> tuple[set[str] | None, str | None]:
    """md5 медиа шаблона либо причина, почему их нет."""
    source_path = Path(profile.source_path) if profile.source_path else None
    if source_path is None or not source_path.exists():
        return None, (
            "путь к исходному файлу шаблона недоступен "
            "(profile.source_path пуст либо файл переехал на диске с момента разбора)"
        )
    try:
        with PptxPackage.open(source_path) as tpl_pkg:
            return {m.md5 for m in tpl_pkg.media()}, None
    except Exception:  # noqa: BLE001 — шаблон недоступен/повреждён, метрика честно пропускается
        return None, "исходный файл шаблона не открылся"


def _picture_bytes(pkg, rels: dict[str, str], ref) -> bytes | None:
    blip_fill = ref.element.find(qn("p:blipFill"))
    blip = blip_fill.find(qn("a:blip")) if blip_fill is not None else None
    rid = blip.get(qn("r:embed")) if blip is not None else None
    part_name = rels.get(rid) if rid is not None else None
    if part_name is None:
        return None
    try:
        return pkg.part(part_name)
    except KeyError:
        return None


def _frame_kind(ref) -> str:
    graphic_data = ref.element.find(f"{qn('a:graphic')}/{qn('a:graphicData')}")
    if graphic_data is None:
        return "raster_content"
    if graphic_data.find(qn("a:tbl")) is not None:
        return "native_table"
    if _CHART_URI_MARK in (graphic_data.get("uri") or ""):
        return "native_chart"
    # SmartArt правится в PowerPoint как текст; OLE-объект и прочее нет.
    if "diagram" in (graphic_data.get("uri") or ""):
        return "editable_text"
    return "raster_content"


def _area(box) -> float:
    if box is None:
        return 0.0
    width = max(0.0, min(box.right, 1.0) - max(box.left, 0.0))
    height = max(0.0, min(box.bottom, 1.0) - max(box.top, 0.0))
    return width * height


def _editability(
    pptx_path: Path, prs, canvas: Canvas, template_hashes: set[str] | None, notes: list[str],
) -> tuple[float | None, dict[str, int]]:
    """Доля редактируемой площади содержания и счётчики по видам.

    Картинка с байтами из медиа шаблона считается декором шаблона (иконка,
    плашка, фон), прочие картинки фотографиями пользователя. Без медиа
    шаблона отличить их нельзя, и все картинки идут фотографиями: метрика
    тогда завышена, о чём пишется заметка."""
    counts = {kind: 0 for kind in EDITABILITY_KINDS}
    editable_area = 0.0
    content_area = 0.0
    with PptxPackage.open(pptx_path) as pkg:
        for slide in prs.slides:
            slide_part = str(slide.part.partname).lstrip("/")
            root = pkg.xml(slide_part)
            rels = pkg.rels(slide_part)
            items: list[tuple[str, float]] = []
            for ref in walk_shapes(root, canvas, include_groups=False):
                kind: str | None = None
                if ref.kind == "shape":
                    tx_body = ref.element.find(qn("p:txBody"))
                    text = "".join(t.text or "" for t in tx_body.iter(qn("a:t"))) if tx_body is not None else ""
                    kind = "editable_text" if text.strip() else "template_decor"
                elif ref.kind == "graphic_frame":
                    kind = _frame_kind(ref)
                elif ref.kind == "picture":
                    data = _picture_bytes(pkg, rels, ref)
                    if data is None:
                        continue
                    native = template_hashes is not None and hashlib.md5(data).hexdigest() in template_hashes
                    kind = "template_decor" if native else "user_image"
                elif ref.kind == "connector":
                    kind = "template_decor"
                if kind is not None:
                    items.append((kind, _area(ref.box)))
            # Слайд одной картинкой: единственное содержание — растр почти во
            # весь холст. Такая картинка не фото, а вёрстка, залитая в растр.
            content = [(k, a) for k, a in items if k != "template_decor"]
            if len(content) == 1 and content[0][0] == "user_image" and content[0][1] >= _SLIDE_AS_IMAGE_AREA:
                items = [("raster_content", a) if k == "user_image" else (k, a) for k, a in items]
            for kind, area in items:
                counts[kind] += 1
                if kind == "template_decor":
                    continue
                content_area += area
                if kind in _EDITABLE_KINDS:
                    editable_area += area
    if template_hashes is None and counts["user_image"]:
        notes.append(
            "Редактируемость: медиа шаблона недоступны, все картинки посчитаны фотографиями пользователя."
        )
    if content_area <= 0:
        notes.append("Редактируемость не посчитана: в колоде нет содержательных объектов с площадью.")
        return None, counts
    return editable_area / content_area, counts


def _native_asset_usage(
    pptx_path: Path, prs, canvas: Canvas, template_hashes: set[str] | None, why_missing: str | None,
    notes: list[str],
) -> float | None:
    if template_hashes is None:
        notes.append(f"Использование нативных ассетов не посчитано: {why_missing}.")
        return None

    total = 0
    native = 0
    with PptxPackage.open(pptx_path) as pkg:
        for slide in prs.slides:
            slide_part = str(slide.part.partname).lstrip("/")
            root = pkg.xml(slide_part)
            rels = pkg.rels(slide_part)
            for ref in walk_shapes(root, canvas, include_groups=False):
                if ref.kind != "picture":
                    continue
                data = _picture_bytes(pkg, rels, ref)
                if data is None:
                    continue
                total += 1
                if hashlib.md5(data).hexdigest() in template_hashes:
                    native += 1

    if total == 0:
        notes.append("Использование нативных ассетов не посчитано: в колоде нет ни одной картинки.")
        return None
    return native / total


def _density_delta(
    prs, canvas: Canvas, profile, clone_pattern_by_index: dict[int, object],
    patterns_by_id: dict[str, object], notes: list[str],
) -> float | None:
    deltas: list[float] = []
    has_field = False
    for index, pattern_id in clone_pattern_by_index.items():
        pattern = patterns_by_id.get(pattern_id)
        if pattern is None:
            continue
        source_density = getattr(pattern, "source_density", None)
        if source_density is None:
            continue
        has_field = True
        ratio = slide_fill_ratio(prs.slides[index], canvas, profile, index=index)
        deltas.append(abs(ratio - source_density))

    if not has_field:
        notes.append(
            "Дельта плотности не посчитана: у паттернов профиля нет поля `source_density` "
            "(добавляет параллельная задача R)."
        )
        return None
    return sum(deltas) / len(deltas) if deltas else None


def _fmt_pct(value: float | None) -> str:
    return f"{value:.0%}" if value is not None else "н/д"


def _summarize(report_values: dict) -> str:
    parts = [
        f"клоном собрано {report_values['native_clone_rate']:.0%} слайдов",
        f"фигур примера сохранено {_fmt_pct(report_values['native_shape_preservation'])}",
        f"типографика/палитра без находок — {_fmt_pct(report_values['typography_palette_compliance'])}",
        (
            f"среднее отклонение геометрии {report_values['mean_geometry_deviation']:.1%}"
            if report_values['mean_geometry_deviation'] is not None else "отклонение геометрии — н/д"
        ),
        (
            f"разнообразие раскладок {_fmt_pct(report_values['pattern_diversity'])} "
            f"(энтропия {report_values['pattern_entropy']:.2f})"
            if report_values['pattern_entropy'] is not None else "разнообразие раскладок — н/д"
        ),
        f"нативных ассетов {_fmt_pct(report_values['native_asset_usage'])}",
        f"редактируемо {_fmt_pct(report_values.get('editable_content_coverage'))} площади содержания",
        (
            f"дельта плотности к паттерну {report_values['density_delta']:.1%}"
            if report_values['density_delta'] is not None else "дельта плотности — н/д"
        ),
    ]
    return "Верность шаблону: " + "; ".join(parts) + "."


def template_fidelity(pptx_path: Path, deck_spec: DeckSpec, profile) -> FidelityReport:
    """Метрики верности готовой колоды (`pptx_path`) шаблону, разобранному в
    `profile`. `deck_spec` — вариант плана, по которому колода собрана
    (`SlideSpec.pattern_id` на слайд, см. `compose.builder.build_deck`) —
    нужен для разнообразия раскладок: клоном собраны не все слайды (слайд
    может быть собран с нуля, `compose/builder.py::_build_from_scratch`), а
    `pattern_id` план несёт для обоих путей сборки."""
    pptx_path = Path(pptx_path)
    prs = Presentation(str(pptx_path))
    canvas = Canvas(profile.canvas_width_emu, profile.canvas_height_emu)
    patterns_by_id = {p.pattern_id: p for p in profile.patterns}
    notes: list[str] = []

    n_slides = len(prs.slides)
    clone_pattern_by_index: dict[int, str] = {}
    for index, slide in enumerate(prs.slides):
        pattern_id = _clone_pattern_id(slide)
        if pattern_id is not None:
            clone_pattern_by_index[index] = pattern_id
    native_clone_rate = len(clone_pattern_by_index) / n_slides if n_slides else 0.0

    native_shape_preservation, mean_geometry_deviation = _native_shape_preservation_and_geometry(
        prs, canvas, clone_pattern_by_index, patterns_by_id, notes,
    )
    typography_palette_compliance = _typography_palette_compliance(
        pptx_path, prs, canvas, profile, AuditConfig.load(), notes,
    )
    pattern_diversity, pattern_entropy = _pattern_diversity_and_entropy(deck_spec, notes)
    template_hashes, why_missing = _template_media_hashes(profile)
    native_asset_usage = _native_asset_usage(pptx_path, prs, canvas, template_hashes, why_missing, notes)
    editable_coverage, editable_counts = _editability(pptx_path, prs, canvas, template_hashes, notes)
    density_delta = _density_delta(prs, canvas, profile, clone_pattern_by_index, patterns_by_id, notes)

    values = {
        "native_clone_rate": native_clone_rate,
        "native_shape_preservation": native_shape_preservation,
        "typography_palette_compliance": typography_palette_compliance,
        "mean_geometry_deviation": mean_geometry_deviation,
        "pattern_diversity": pattern_diversity,
        "pattern_entropy": pattern_entropy,
        "native_asset_usage": native_asset_usage,
        "density_delta": density_delta,
        "editable_content_coverage": editable_coverage,
    }
    return FidelityReport(**values, summary=_summarize(values), notes=notes, editable_counts=editable_counts)
