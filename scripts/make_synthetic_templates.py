"""Три синтетических шаблона для проверки пайплайна на незнакомом .pptx.

Датасет даёт четыре шаблона, на защите придёт пятый, которого никто не
видел. Эти три собраны так, чтобы бить по разным допущениям разбора и
сборки: (a) строгий корпоративный 16:9 с таблицей, родной диаграммой и
колонтитулами; (b) 4:3, где карточки собраны группами, заливки градиентные,
есть поворот, соединители и шрифт, которого нет в системе; (c) почти
пустой: три слайда, пустая тема шрифтов, без картинок.

Детерминированный: одинаковый запуск даёт побайтно одинаковые файлы (время
в zip и свойства документа зафиксированы), поэтому файлы лежат в git и
тесты не пересобирают их на лету.

Запуск: `python scripts/make_synthetic_templates.py [каталог]`, по
умолчанию `fixtures/templates-synthetic/`.
"""
from __future__ import annotations

import copy
import datetime as dt
import io
import sys
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "fixtures" / "templates-synthetic"
FIXED_TIME = dt.datetime(2026, 1, 1, 12, 0, 0)
ZIP_TIME = (2026, 1, 1, 12, 0, 0)

NAVY = RGBColor(0x1F, 0x2A, 0x44)
TEAL = RGBColor(0x00, 0x87, 0x8A)
SAND = RGBColor(0xF2, 0xEE, 0xE6)
ORANGE = RGBColor(0xE8, 0x6A, 0x33)
GREY = RGBColor(0x5A, 0x5F, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ---------------------------------------------------------------------------
# Общие приёмы
# ---------------------------------------------------------------------------


def _theme_part(prs):
    master = prs.slide_masters[0]
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            return rel.target_part
    raise RuntimeError("у мастера нет темы")


def _edit_theme(prs, edit) -> None:
    """Правит XML темы на месте. python-pptx не даёт API к теме, а часть
    хранит blob, поэтому переписываем его целиком."""
    part = _theme_part(prs)
    root = etree.fromstring(part.blob)
    edit(root)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _set_theme_fonts(major: str, minor: str):
    def edit(root) -> None:
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for tag, face in (("majorFont", major), ("minorFont", minor)):
            node = root.find(f".//a:fontScheme/a:{tag}", ns)
            node.find("a:latin", ns).set("typeface", face)
            # Кириллица и прочие письменности идут через latin: убираем
            # отдельные гарнитуры, чтобы тема несла ровно две.
            for extra in node.findall("a:font", ns):
                node.remove(extra)
    return edit


def _set_theme_colors(dk2: str, lt2: str, accents: list[str]):
    def edit(root) -> None:
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = root.find(".//a:clrScheme", ns)
        for name, value in [("dk2", dk2), ("lt2", lt2)] + [(f"accent{i + 1}", v) for i, v in enumerate(accents)]:
            slot = scheme.find(f"a:{name}", ns)
            for child in list(slot):
                slot.remove(child)
            etree.SubElement(slot, qn("a:srgbClr")).set("val", value)
    return edit


def _scale_master_x(prs, factor: float) -> None:
    """Встроенный шаблон python-pptx свёрстан под 4:3. На 16:9 его
    плейсхолдеры сжались бы в левые три четверти холста, поэтому
    растягиваем мастер и лейауты по горизонтали."""
    parts = [prs.slide_masters[0]] + list(prs.slide_layouts)
    # Геометрию всех фигур читаем ДО записи: плейсхолдер лейаута без своего
    # a:xfrm наследует её у мастера, и после правки мастера растянулся бы
    # дважды. Записываем все четыре числа: при записи одного `left`
    # python-pptx создаёт xfrm с нулевой высотой.
    frames = [
        (shape, shape.left, shape.top, shape.width, shape.height)
        for holder in parts for shape in holder.shapes
        if shape.left is not None and shape.width is not None
    ]
    for shape, left, top, width, height in frames:
        shape.left = Emu(int(left * factor))
        shape.top = Emu(top)
        shape.width = Emu(int(width * factor))
        shape.height = Emu(height)


def _add_slide_number(slide) -> None:
    """Колонтитул с номером: копия плейсхолдера номера из лейаута с полем
    `slidenum`. Так делает PowerPoint, когда включают «Номер слайда»."""
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.type is not None and int(ph.placeholder_format.type) == 13:
            sp = copy.deepcopy(ph._element)
            txbody = sp.find(qn("p:txBody"))
            for p in txbody.findall(qn("a:p")):
                txbody.remove(p)
            p = etree.SubElement(txbody, qn("a:p"))
            fld = etree.SubElement(p, qn("a:fld"))
            fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
            fld.set("type", "slidenum")
            etree.SubElement(fld, qn("a:rPr")).set("lang", "ru-RU")
            etree.SubElement(fld, qn("a:t")).text = "‹#›"
            slide.shapes._spTree.append(sp)
            return


def _text(shape, text: str, *, size: int | None = None, bold: bool = False,
          color: RGBColor | None = None, font: str | None = None) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    for p in tf.paragraphs:
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            r.font.bold = bold
            if color is not None:
                r.font.color.rgb = color
            if font:
                r.font.name = font


def _box(slide, shape_type, x, y, w, h, *, fill: RGBColor | None = None, line: bool = False):
    shape = slide.shapes.add_shape(shape_type, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if not line:
        shape.line.fill.background()
    return shape


def _placeholder_png(width: int, height: int, seed: int) -> bytes:
    """Картинка-заглушка: полосы и круг, без шума, чтобы байты не зависели
    от версии Pillow сильнее, чем нужно."""
    img = Image.new("RGB", (width, height), (0x1F, 0x2A, 0x44))
    draw = ImageDraw.Draw(img)
    for i in range(0, width, 40):
        shade = 0x30 + (i * 7 + seed * 13) % 0x50
        draw.rectangle([i, 0, i + 20, height], fill=(shade, 0x87, 0x8A))
    r = min(width, height) // 4
    draw.ellipse([width // 2 - r, height // 2 - r, width // 2 + r, height // 2 + r], fill=(0xE8, 0x6A, 0x33))
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=False)
    return buf.getvalue()


def _freeze(prs) -> None:
    cp = prs.core_properties
    cp.author = "DeckForge synthetic"
    cp.last_modified_by = "DeckForge synthetic"
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.revision = 1


def _save_deterministic(prs, path: Path) -> None:
    """python-pptx пишет в zip текущее время; перепаковываем с постоянным,
    чтобы файл в git не менялся от перезапуска генератора."""
    _freeze(prs)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o644 << 16
            dst.writestr(fixed, src.read(info.filename))
    path.write_bytes(out.getvalue())


def _layout(prs, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise KeyError(name)


# ---------------------------------------------------------------------------
# (a) строгий корпоративный 16:9
# ---------------------------------------------------------------------------


def make_corporate(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    _scale_master_x(prs, 12192000 / 9144000)
    _edit_theme(prs, _set_theme_fonts("Georgia", "Arial"))
    _edit_theme(prs, _set_theme_colors("1F2A44", "F2EEE6", ["00878A", "E86A33", "1F2A44", "8AA1B1", "C9B79C", "5A5F6B"]))
    W, H = 12192000, 6858000
    M = 609600

    # 1. Обложка
    s = prs.slides.add_slide(_layout(prs, "Title Slide"))
    _box(s, MSO_SHAPE.RECTANGLE, 0, 0, W // 3, H, fill=NAVY)
    s.shapes.title.text = "Название доклада в две строки"
    s.placeholders[1].text = "Подзаголовок, имя докладчика, дата"
    _add_slide_number(s)

    # 2. Разделитель
    s = prs.slides.add_slide(_layout(prs, "Section Header"))
    _box(s, MSO_SHAPE.RECTANGLE, 0, H - 914400, W, 914400, fill=TEAL)
    s.shapes.title.text = "01. Раздел"
    s.placeholders[1].text = "Короткая подводка к разделу"
    _add_slide_number(s)

    # 3. Две колонки
    s = prs.slides.add_slide(_layout(prs, "Two Content"))
    s.shapes.title.text = "Две колонки: было и стало"
    s.placeholders[1].text = "Было: ручная маршрутизация\nСреднее время три дня\nОчередь растёт к концу квартала"
    s.placeholders[2].text = "Стало: автоматическая маршрутизация\nСреднее время полдня\nОчередь ровная"
    _add_slide_number(s)

    # 4. Три карточки с иконками-фигурами
    s = prs.slides.add_slide(_layout(prs, "Title Only"))
    s.shapes.title.text = "Три карточки с иконками"
    card_w = (W - 2 * M - 2 * 304800) // 3
    top = 1828800
    icons = (MSO_SHAPE.OVAL, MSO_SHAPE.DIAMOND, MSO_SHAPE.HEXAGON)
    for i, icon in enumerate(icons):
        x = M + i * (card_w + 304800)
        _box(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, top, card_w, 3657600, fill=SAND)
        _box(s, icon, x + 228600, top + 228600, 548640, 548640, fill=TEAL)
        head = s.shapes.add_textbox(Emu(x + 228600), Emu(top + 914400), Emu(card_w - 457200), Emu(548640))
        _text(head, f"Заголовок карточки {i + 1}", size=20, bold=True, color=NAVY)
        body = s.shapes.add_textbox(Emu(x + 228600), Emu(top + 1554480), Emu(card_w - 457200), Emu(1828800))
        _text(body, "Описание карточки в две-три строки текста о сути пункта", size=14, color=GREY)
    _add_slide_number(s)

    # 5. Таблица 4x3
    s = prs.slides.add_slide(_layout(prs, "Title Only"))
    s.shapes.title.text = "Таблица сравнения"
    table = s.shapes.add_table(4, 3, Emu(M), Emu(1828800), Emu(W - 2 * M), Emu(3200400)).table
    cells = [["Показатель", "До", "После"], ["Время", "3 дня", "0,5 дня"], ["Ошибки", "12%", "3%"], ["Эскалации", "40", "9"]]
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    _add_slide_number(s)

    # 6. Родная диаграмма bar
    s = prs.slides.add_slide(_layout(prs, "Title Only"))
    s.shapes.title.text = "Диаграмма по кварталам"
    data = CategoryChartData()
    data.categories = ["I кв.", "II кв.", "III кв.", "IV кв."]
    data.add_series("Заявки", (120, 150, 170, 210))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(M), Emu(1600200), Emu(W - 2 * M), Emu(4343400), data)
    _add_slide_number(s)

    # 7. Фото-слайд с картинкой-заглушкой
    s = prs.slides.add_slide(_layout(prs, "Title Only"))
    s.shapes.title.text = "Фото и подпись"
    s.shapes.add_picture(io.BytesIO(_placeholder_png(960, 720, 1)), Emu(W // 2), Emu(0), Emu(W // 2), Emu(H))
    cap = s.shapes.add_textbox(Emu(M), Emu(2743200), Emu(W // 2 - 2 * M), Emu(1828800))
    _text(cap, "Подпись к фотографии в несколько строк: что на ней и зачем", size=18, color=GREY)
    _add_slide_number(s)

    # 8. Финал
    s = prs.slides.add_slide(_layout(prs, "Title Slide"))
    _box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=NAVY)
    # Фон ушёл вперёд плейсхолдеров: переносим его в начало дерева.
    tree = s.shapes._spTree
    bg = tree[-1]
    tree.remove(bg)
    tree.insert(2, bg)
    s.shapes.title.text = "Спасибо!"
    s.placeholders[1].text = "Контакты: почта, телефон"
    _add_slide_number(s)

    _save_deterministic(prs, path)


# ---------------------------------------------------------------------------
# (b) 4:3 с группами
# ---------------------------------------------------------------------------


def _gradient(shape, a: RGBColor, b: RGBColor, angle: float = 90.0) -> None:
    shape.fill.gradient()
    shape.fill.gradient_angle = angle
    stops = shape.fill.gradient_stops
    stops[0].color.rgb = a
    stops[1].color.rgb = b


def make_grouped(path: Path) -> None:
    prs = Presentation()  # встроенный шаблон уже 4:3
    _edit_theme(prs, _set_theme_fonts("Bogus Sans", "Bogus Sans"))
    W, H = 9144000, 6858000
    M = 457200

    # 1. Обложка с градиентным фоном и повёрнутой плашкой
    s = prs.slides.add_slide(_layout(prs, "Title Slide"))
    bg = _box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _gradient(bg, NAVY, TEAL)
    tree = s.shapes._spTree
    tree.remove(bg._element)
    tree.insert(2, bg._element)
    tag = _box(s, MSO_SHAPE.RECTANGLE, W - 2286000, 457200, 1828800, 365760, fill=ORANGE)
    tag.rotation = 15.0
    _text(tag, "ЧЕРНОВИК", size=14, bold=True, color=WHITE)
    s.shapes.title.text = "Обложка на градиенте"
    s.placeholders[1].text = "Подзаголовок шрифтом, которого нет"

    # 2. Три карточки-группы
    s = prs.slides.add_slide(_layout(prs, "Title Only"))
    s.shapes.title.text = "Карточки собраны группами"
    card_w = (W - 2 * M - 2 * 182880) // 3
    for i in range(3):
        x = M + i * (card_w + 182880)
        group = s.shapes.add_group_shape()
        plate = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(1828800), Emu(card_w), Emu(3657600))
        _gradient(plate, SAND, WHITE, 270.0)
        plate.line.fill.background()
        dot = group.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x + 182880), Emu(2011680), Emu(457200), Emu(457200))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ORANGE
        dot.line.fill.background()
        head = group.shapes.add_textbox(Emu(x + 182880), Emu(2651760), Emu(card_w - 365760), Emu(457200))
        _text(head, f"Пункт {i + 1}", size=18, bold=True, color=NAVY)
        body = group.shapes.add_textbox(Emu(x + 182880), Emu(3200400), Emu(card_w - 365760), Emu(1828800))
        _text(body, "Текст карточки внутри группы, две строки", size=12, color=GREY)

    # 3. Схема из соединителей
    s = prs.slides.add_slide(_layout(prs, "Title Only"))
    s.shapes.title.text = "Процесс: шаги и стрелки"
    step_w, step_h, gap = 1828800, 914400, 548640
    y = 3017520
    boxes = []
    for i in range(4):
        x = M + i * (step_w + gap)
        box = _box(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, step_w, step_h, fill=TEAL)
        _text(box, f"Шаг {i + 1}", size=16, bold=True, color=WHITE)
        boxes.append(box)
    for a, b in zip(boxes, boxes[1:]):
        conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(0), Emu(0), Emu(1), Emu(1))
        conn.begin_connect(a, 3)
        conn.end_connect(b, 1)
        conn.line.color.rgb = NAVY
        conn.line.width = Pt(2)
    elbow = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, Emu(0), Emu(0), Emu(1), Emu(1))
    elbow.begin_connect(boxes[0], 2)
    elbow.end_connect(boxes[3], 2)

    # 4. Текстовый слайд с градиентной полосой
    s = prs.slides.add_slide(_layout(prs, "Title and Content"))
    bar = _box(s, MSO_SHAPE.RECTANGLE, 0, H - 365760, W, 365760)
    _gradient(bar, ORANGE, NAVY, 0.0)
    s.shapes.title.text = "Текстовый слайд"
    s.placeholders[1].text = "Первый тезис\nВторой тезис\nТретий тезис"

    # 5. Финал
    s = prs.slides.add_slide(_layout(prs, "Title Slide"))
    s.shapes.title.text = "Вопросы?"
    s.placeholders[1].text = "Спасибо за внимание"

    _save_deterministic(prs, path)


# ---------------------------------------------------------------------------
# (c) почти пустой
# ---------------------------------------------------------------------------


def _drop_theme_fonts(root) -> None:
    """Тема без шрифтов: у мажорной и минорной гарнитуры пустое имя. Схема
    OOXML требует сами элементы, поэтому их не удаляем, а обнуляем."""
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for tag in ("majorFont", "minorFont"):
        node = root.find(f".//a:fontScheme/a:{tag}", ns)
        for child in list(node):
            node.remove(child)
        for sub in ("latin", "ea", "cs"):
            etree.SubElement(node, qn(f"a:{sub}")).set("typeface", "")


def make_bare(path: Path) -> None:
    prs = Presentation()
    _edit_theme(prs, _drop_theme_fonts)
    s = prs.slides.add_slide(_layout(prs, "Title Slide"))
    s.shapes.title.text = "Заголовок"
    s.placeholders[1].text = "Подзаголовок"
    s = prs.slides.add_slide(_layout(prs, "Title and Content"))
    s.shapes.title.text = "Слайд с текстом"
    s.placeholders[1].text = "Один пункт\nДругой пункт"
    s = prs.slides.add_slide(_layout(prs, "Title Slide"))
    s.shapes.title.text = "Конец"
    _save_deterministic(prs, path)


TEMPLATES = {
    "corporate-16x9.pptx": make_corporate,
    "grouped-4x3.pptx": make_grouped,
    "bare-3-slides.pptx": make_bare,
}


def main(argv: list[str]) -> int:
    out = Path(argv[1]) if len(argv) > 1 else DEFAULT_OUT
    out.mkdir(parents=True, exist_ok=True)
    for name, make in TEMPLATES.items():
        make(out / name)
        print(out / name)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
