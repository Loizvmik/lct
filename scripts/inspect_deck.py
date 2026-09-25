"""Сводка по собранной презентации — то, что нужно для разбора глазами.

Аудит отвечает на вопрос «есть ли брак по счётным признакам». Этот скрипт
отвечает на другой: «как слайд устроен» — сколько на нём текста, декора,
картинок, какой кегль у заголовка, какую долю холста занимает содержание.
Именно этих чисел не хватало 25 сентября, когда пришлось полдня выяснять,
почему заголовок вышел мелким.

    .venv/bin/python scripts/inspect_deck.py <файл.pptx> [профиль.json]

Без профиля считается только то, что видно в самом файле; с профилем
добавляется заполненность холста тем же расчётом, что у проверки D05.
"""
from __future__ import annotations
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from deckforge.audit.deterministic import slide_fill_ratio  # noqa: E402
from deckforge.ooxml.geometry import Canvas  # noqa: E402
from deckforge.template.profile import TemplateProfile  # noqa: E402


def _kind(shape) -> str:
    return str(shape.shape_type).split()[0].lower()


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = Path(sys.argv[1])
    prs = Presentation(str(path))

    profile = canvas = None
    if len(sys.argv) > 2:
        profile = TemplateProfile.model_validate_json(Path(sys.argv[2]).read_text(encoding="utf-8"))
        canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)

    print(f"{path.name}: слайдов {len(prs.slides)}")
    print(f"{'№':>3} {'фигур':>6} {'текст':>6} {'карт':>5} {'плашек':>7} {'кегли':>16} {'занято':>7} {'путь':>12}  заголовок")
    empty = 0
    for i, slide in enumerate(prs.slides, start=1):
        kinds = Counter(_kind(s) for s in slide.shapes)
        texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        sizes = sorted({
            r.font.size.pt
            for s in slide.shapes if s.has_text_frame
            for p in s.text_frame.paragraphs for r in p.runs if r.font.size
        })
        fill = ""
        if canvas is not None:
            ratio = slide_fill_ratio(slide, canvas, profile, index=i - 1)
            fill = f"{ratio:.0%}"
            if ratio < 0.25:
                empty += 1
        head = texts[0][:46] if texts else "— пусто —"
        # Путь сборки: клон слайда-примера сборка метит в имени слайда
        # (`compose.builder.CLONE_MARK_PREFIX`), слайд без метки собран с нуля.
        name = slide._element.cSld.get("name") or ""
        path_label = "клон " + name.rsplit(":", 1)[-1] if name.startswith("deckforge:clone:") else "с нуля"
        print(
            f"{i:>3} {len(slide.shapes):>6} {len(texts):>6} {kinds.get('picture', 0):>5} "
            f"{kinds.get('auto_shape', 0):>7} {str(sizes)[:16]:>16} {fill:>7} {path_label:>12}  {head}"
        )
    if canvas is not None:
        print(f"\nслайдов заполнено меньше четверти: {empty} из {len(prs.slides)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
