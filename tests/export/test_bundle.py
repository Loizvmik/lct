"""Тест `deckforge.export.bundle.export_bundle` (Task 14, Step 1 брифа:
`test_bundle_has_all_three_formats`). Требует LibreOffice — `.pdf`/`.png`
идут через `render.soffice` (см. её докстроку про изоляцию параллельных
вызовов и шрифт шаблона)."""
from __future__ import annotations

import pytest

from deckforge.export.bundle import export_bundle
from deckforge.render.soffice import soffice_available


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice не установлен")
def test_bundle_has_all_three_formats(PROFILE, PPTX, tmp_path):
    bundle = export_bundle(PPTX, PROFILE, tmp_path)
    assert bundle.pptx.exists() and bundle.pdf.exists() and bundle.html.exists()


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice не установлен")
def test_bundle_includes_deck_spec_when_given(DECK, PROFILE, PPTX, tmp_path):
    bundle = export_bundle(PPTX, PROFILE, tmp_path, deck_spec=DECK)
    text = bundle.html.read_text(encoding="utf-8")
    assert DECK.slides[1].headline in text


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice не установлен")
def test_bundle_renders_one_png_per_slide(PROFILE, PPTX, tmp_path):
    from pptx import Presentation

    bundle = export_bundle(PPTX, PROFILE, tmp_path)
    assert len(bundle.pngs) == len(Presentation(str(PPTX)).slides)


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice не установлен")
def test_bundle_converts_pptx_to_pdf_only_once(PROFILE, PPTX, tmp_path, monkeypatch):
    """PDF и PNG-превью раньше рендерились из pptx по отдельности (`to_pdf`
    внутри `to_pngs` конвертировал файл заново) — на трёх вариантах это 6
    вызовов soffice вместо 3, замер: 302.8с против лимита ТЗ в 300с."""
    import deckforge.render.soffice as soffice_mod

    calls = []
    orig_run_soffice = soffice_mod._run_soffice
    monkeypatch.setattr(
        soffice_mod, "_run_soffice",
        lambda *a, **kw: calls.append(a) or orig_run_soffice(*a, **kw),
    )

    export_bundle(PPTX, PROFILE, tmp_path)

    assert len(calls) == 1
