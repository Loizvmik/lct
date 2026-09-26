"""Тесты `deckforge.export.html.to_html` (Task 14, Step 1 брифа).

`PROFILE.palette.roles`/`PROFILE.type_scale.families` в исходном тексте
брифа — фактические имена в `TemplateProfile` (Task 8+) другие:
`palette_roles` — плоский словарь на самой модели, не вложенный объект
`.palette.roles`; `type_scale.families` совпадает дословно. Тесты ниже
написаны по реальному API, тем же способом, каким расходится буква брифа с
кодом и в других тестовых модулях проекта (см. `tests/compose/conftest.py`
и соседей)."""
from __future__ import annotations

from deckforge.audit.fidelity import template_fidelity
from deckforge.audit.visual import VisualAuditResult
from deckforge.export.html import to_html


def test_html_is_a_single_self_contained_file(DECK, PROFILE, PPTX, tmp_path):
    html = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html")
    text = html.read_text(encoding="utf-8")
    assert '<script src="http' not in text
    assert '<link rel="stylesheet" href="http' not in text
    assert "://" not in text.replace("http://www.w3.org", "")  # namespaces/атрибуты SVG не в счёт


def test_html_slides_are_text_not_screenshots(DECK, PROFILE, PPTX, tmp_path):
    """То же требование, что и к pptx: слайд не должен быть картинкой."""
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    assert DECK.slides[1].headline in text
    for item in DECK.slides[1].blocks[0].items:
        assert item in text
    assert "<img" not in text or "data:image" in text  # если картинка есть — она вшита, не по ссылке


def test_html_uses_template_tokens(DECK, PROFILE, PPTX, tmp_path):
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    assert PROFILE.palette_roles["brand"].lower() in text.lower()
    assert PROFILE.type_scale.families[0] in text


def test_html_renders_table_content_as_native_table(DECK, PROFILE, PPTX, tmp_path):
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    assert "<table" in text
    assert "стрелки" in text and "страницы" in text


def test_html_renders_chart_data_as_text_not_raster(DECK, PROFILE, PPTX, tmp_path):
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    assert "<svg" in text
    # числа графика лежат текстом внутри SVG/скрытой таблицы данных, не только пикселями
    assert "VK Tech" in text and "WorkSpace" in text


def test_html_has_keyboard_navigation_and_overview(DECK, PROFILE, PPTX, tmp_path):
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    assert "ArrowRight" in text and "ArrowLeft" in text
    assert "overview" in text


def test_html_slide_count_matches_pptx(DECK, PROFILE, PPTX, tmp_path):
    from pptx import Presentation

    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    n_slides = len(Presentation(str(PPTX)).slides)
    assert text.count('<section class="slide"') == n_slides * 2  # экран + печать (см. докстроку html.py)


def test_html_without_visual_argument_has_no_score_markup(DECK, PROFILE, PPTX, tmp_path):
    """Старое поведение буквально: без `visual` (по умолчанию) экспорт не
    знает про PPTEval вовсе — та же честная деградация, что и у остальных
    необязательных надстроек проекта."""
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    # CSS-правила `.slide-score`/`#deck-score` — часть статичного шаблона
    # документа и присутствуют всегда (as-is стиль на случай, если аудит
    # прогонят позже и пересоберут те же слайды); проверяем отсутствие
    # именно РАЗМЕТКИ (элементов), а не факта упоминания класса в CSS.
    assert 'class="slide-score"' not in text
    assert 'id="deck-score"' not in text


def test_html_shows_ppteval_scores_when_visual_result_is_passed(DECK, PROFILE, PPTX, tmp_path):
    """Задача G: оценка того слайда, чей `headline` совпал с индексом
    `slide_scores` (`DECK.slides[1]` — "Где уходит время"), должна попасть в
    бейдж рядом с этим слайдом, а сводка по колоде — куда-то в документ
    целиком."""
    visual = VisualAuditResult(
        slide_scores={1: {"content": 4, "design": 3, "why": "ясно, но тесно"}},
        deck_score={"coherence": 5, "why": "переходы логичны"},
        content_avg=4.0,
        design_avg=3.0,
    )
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html", visual=visual).read_text(encoding="utf-8")
    assert 'class="slide-score"' in text
    assert "ясно, но тесно" in text
    assert 'id="deck-score"' in text
    assert "содержание 4.0" in text
    assert "дизайн 3.0" in text
    assert "связность 5" in text


def test_html_score_badge_appears_only_for_the_matched_slide(DECK, PROFILE, PPTX, tmp_path):
    """Слайд без записи в `slide_scores` не получает бейдж вовсе — только
    тот, чей индекс реально сопоставился (`_match_spec_slide`)."""
    visual = VisualAuditResult(slide_scores={1: {"content": 5, "design": 5, "why": "образцовый слайд"}})
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html", visual=visual).read_text(encoding="utf-8")
    # Ровно два бейджа на слайд 1 (экранная версия + печатная — та же
    # причина, что и у `test_html_slide_count_matches_pptx`).
    assert text.count('class="slide-score"') == 2


def test_html_carries_speaker_notes(DECK, PROFILE, PPTX, tmp_path):
    """Заказчик ждёт на выходе «готовые слайды и текст к каждому слайду»
    (уточнение от 23 сентября 2026): по слайдам на защите рассказывают, и
    текст докладчика должен читаться рядом со слайдом, а не только на
    странице заметок .pptx, которую в браузере не открыть."""
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    assert DECK.slides[0].speaker_notes in text


def test_html_without_fidelity_argument_has_no_fidelity_markup(DECK, PROFILE, PPTX, tmp_path):
    """Задача T: та же честная деградация, что и у PPTEval (см. `test_html_
    without_visual_argument_has_no_score_markup`) — без `fidelity` (по
    умолчанию) блок «Верность шаблону» отсутствует в разметке."""
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html").read_text(encoding="utf-8")
    assert 'id="template-fidelity"' not in text


def test_html_shows_fidelity_summary_when_passed(DECK, PROFILE, PPTX, tmp_path):
    """Задача T: `audit.fidelity.template_fidelity` на том же PPTX/DECK,
    переданный в `to_html`, попадает в документ отдельным блоком (см.
    `_fidelity_html`, не трогает разметку PPTEval/бюджета)."""
    report = template_fidelity(PPTX, DECK, PROFILE)
    text = to_html(DECK, PROFILE, PPTX, tmp_path / "deck.html", fidelity=report).read_text(encoding="utf-8")
    assert 'id="template-fidelity"' in text
    assert report.summary in text
