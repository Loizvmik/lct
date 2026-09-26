"""Лестница отказов сборки (задача U, раздел 11 архитектуры): клон,
запасная раскладка, сокращение текста под контракт, разбиение слайда,
и только потом сборка с нуля. Модель фейковая, сети нет; шаблон VK
Education из `dataset/templates/`, профиль разбирается без модели."""
from __future__ import annotations
import functools
import json
import threading
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.compose import builder
from deckforge.compose.clone import sample_slides_by_number
from deckforge.ooxml.geometry import Canvas
from deckforge.pattern.forms import forms_of
from deckforge.plan.spec import BulletBlock, Card, CardBlock, DeckSpec, SlideSpec
from deckforge.plan.variants import Variant
from deckforge.provider.base import LLMProvider
import deckforge.provider.scheduler as scheduler_module
from deckforge.provider.scheduler import ModelScheduler
from deckforge.template.profile import TemplateProfile
from deckforge.workflow.repair import SlideRepairer, contract_for_slide


TEMPLATE = Path("dataset/templates/Шаблон презентации VK Education.pptx")
_LONG = " ".join(["длинное пояснение про очередь заявок и ожидание согласующего"] * 5)
_SHORT = ["Медиана 31,5 часа, работа 28 минут", "Ожидание 98,5% времени"]


@functools.lru_cache(maxsize=None)
def _profile() -> TemplateProfile:
    return TemplateProfile.from_file(TEMPLATE, cache_dir=None)


@pytest.fixture(scope="module")
def profile():
    return _profile()


class _FakeRepair:
    """`compose.builder.SlideRepair` без модели: записывает вызовы и
    отдаёт `result(slide)` или `None`."""

    def __init__(self, result=None):
        self.result, self.calls = result, []

    def shorten(self, slide, pattern_id, problems):
        self.calls.append((slide.index, pattern_id, list(problems)))
        return self.result(slide) if self.result else None


def _short(slide: SlideSpec) -> SlideSpec:
    return replace(slide, blocks=[BulletBlock(items=list(_SHORT))])


def _long_slide() -> SlideSpec:
    return SlideSpec(
        index=0, kind="bullets", headline="Заявка ждёт 98,5% времени", pattern_id="slide13",
        blocks=[BulletBlock(items=[_LONG, _LONG, _LONG])],
    )


def _ladder(profile, slide: SlideSpec, repair, *, room: int = 5, pattern_ids=("slide13",)):
    """Лестница одного слайда на одной раскладке: запасных нет, чтобы
    ступень сокращения была единственным способом собрать клон."""
    prs = Presentation(str(TEMPLATE))
    sources = sample_slides_by_number(prs)
    builder._clear_sample_slides(prs)
    patterns = [
        builder._pattern_from_model(m) for pid in pattern_ids for m in profile.patterns if m.pattern_id == pid
    ]
    canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)
    outcome = builder._place_with_ladder(
        prs, slide, patterns, profile, canvas, AuditConfig.load(), forms=forms_of(profile), repair=repair,
        room=room, bullet_char="•", user_photos=None, image_bytes=None, source_slides=sources,
    )
    return outcome, prs


def test_first_clone_is_the_first_rung(profile):
    slide = replace(_long_slide(), blocks=[BulletBlock(items=list(_SHORT))])
    repair = _FakeRepair(_short)

    outcome, prs = _ladder(profile, slide, repair)

    assert outcome.rung == "clone"
    assert repair.calls == [], "клон принят сразу: чинить нечего"
    assert len(prs.slides) == 1


def test_shortened_text_is_accepted_and_replaces_the_slide_text(profile):
    slide = _long_slide()
    repair = _FakeRepair(_short)

    outcome, prs = _ladder(profile, slide, repair)

    assert outcome.rung == "shorten"
    assert len(repair.calls) == 1
    _index, pattern_id, problems = repair.calls[0]
    assert pattern_id == "slide13"
    # Причина отказа: переполнение, найденное аудитом (L03) или пределом
    # ужимания кегля при подгонке (задача V3), с указанием места.
    assert any("L03" in p or "не помещается" in p for p in problems), "модель должна знать, почему клон отклонён"
    assert slide.meta["failure_codes"].startswith("TEXT_OVERFLOW/")
    assert slide.meta["ladder_path"].startswith("clone>roomier>shorten")
    assert slide.blocks[0].items == _SHORT, "в колоде тот текст, что лёг на слайд"
    assert len(prs.slides) == 1
    name = prs.slides[0]._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld",
    ).get("name") or ""
    assert name.startswith(builder.CLONE_MARK_PREFIX)
    assert any("текст сокращён" in n for n in outcome.notes)


def test_scratch_comes_only_after_repair_was_tried(profile):
    slide = _long_slide()
    repair = _FakeRepair(None)

    outcome, prs = _ladder(profile, slide, repair)

    assert len(repair.calls) == 1, "сокращение пробуется до сборки с нуля"
    assert outcome.rung == "scratch"
    assert len(prs.slides) == 1
    assert outcome.notes[-1].endswith("(раскладка 'slide13').") and "сборка с нуля" in outcome.notes[-1]


def _cards_slide(n: int) -> SlideSpec:
    return SlideSpec(
        index=0, kind="cards", headline="Восемь шагов раскатки", pattern_id="slide21",
        blocks=[CardBlock(items=[Card(title=f"Шаг {i}", body=f"Короткое описание шага {i}") for i in range(n)])],
        speaker_notes="Заметка докладчика",
    )


def test_too_many_units_split_the_slide_instead_of_cutting_text(profile):
    spec = DeckSpec(title="t", language="ru", slides=[_cards_slide(8)])
    repair = _FakeRepair(_short)

    builder.build_deck(spec, profile, TEMPLATE, Variant.dense, repair=repair)

    assert repair.calls == [], "переполнение по единицам не чинится сокращением: карточки унесли бы факты"
    assert [len(s.blocks[0].items) for s in spec.slides] == [4, 4]
    first, second = spec.slides
    assert second.headline == first.headline, "заголовок-вывод один на оба слайда"
    assert "Продолжение" in second.speaker_notes and "Продолжение" not in second.headline
    assert [s.index for s in spec.slides] == [0, 1]
    counts = builder.ladder_counts(spec)
    assert counts["split"] == 1 and counts["scratch"] == 0
    assert any("нужно 8 единиц, у лучшей раскладки максимум 4" in f for f in first.findings)


def test_no_split_when_the_deck_is_at_the_slide_limit(profile):
    spec = DeckSpec(title="t", language="ru", slides=[_cards_slide(8)])

    builder.build_deck(spec, profile, TEMPLATE, Variant.dense, repair=_FakeRepair(None), max_slides=1)

    assert len(spec.slides) == 1
    assert builder.ladder_counts(spec)["split"] == 0
    assert any("разделить нельзя" in f for f in spec.slides[0].findings)


def test_every_slide_is_counted_on_exactly_one_rung(profile):
    slides = [_cards_slide(8), replace(_long_slide(), index=1)]
    spec = DeckSpec(title="t", language="ru", slides=slides)

    builder.build_deck(spec, profile, TEMPLATE, Variant.dense, repair=_FakeRepair(_short))

    counts = builder.ladder_counts(spec)
    assert sum(counts.values()) == len(spec.slides) == 3
    assert all(any("лестница сборки" in f for f in s.findings) for s in spec.slides)


# ---------------------------------------------------------------------------
# SlideRepairer: сокращение через писателя и его общий предел вызовов
# ---------------------------------------------------------------------------


class _ShortenLLM(LLMProvider):
    """Отвечает коротким слайдом под контракт и меряет одновременность."""

    def __init__(self):
        self.requests: list[dict] = []
        self.concurrent = self.max_concurrent = 0
        self._lock = threading.Lock()

    def complete(self, messages, *, schema=None, max_tokens=4096, temperature=0.3) -> str:
        payload = json.loads(messages[1]["content"])
        with self._lock:
            self.requests.append(payload)
            self.concurrent += 1
            self.max_concurrent = max(self.max_concurrent, self.concurrent)
        try:
            import time
            time.sleep(0.05)
            return json.dumps({
                "headline": "Заявка ждёт 98,5% времени",
                "blocks": [{"type": "bullets", "items": list(_SHORT)}],
                "source_note": "Данные: замер процесса",
            }, ensure_ascii=False)
        finally:
            with self._lock:
                self.concurrent -= 1


def test_repairer_shortens_through_the_writer_contract(profile):
    llm = _ShortenLLM()
    repairer = SlideRepairer(profile=profile, llm=llm, sources=[], style="dense")

    result = repairer.shorten(_long_slide(), "slide13", ["L03: не помещается в рамку"])

    assert result is not None and result.blocks[0].items == _SHORT
    assert result.pattern_id == "slide13"
    request = llm.requests[0]
    assert request["layout_rejected"] is True
    assert "L03: не помещается в рамку" in request["contract_problems"]
    assert request["contract"]["blocks"][0]["type"] == "bullets"
    assert repairer.calls == 1 and repairer.accepted == 1


def test_repairer_respects_its_call_limit_and_works_without_a_model(profile):
    llm = _ShortenLLM()
    repairer = SlideRepairer(profile=profile, llm=llm, sources=[], max_calls=1)

    assert repairer.shorten(_long_slide(), "slide13", []) is not None
    assert repairer.shorten(_long_slide(), "slide13", []) is None
    assert len(llm.requests) == 1
    assert SlideRepairer(profile=profile, llm=None, sources=[]).shorten(_long_slide(), "slide13", []) is None


def test_repair_calls_share_the_writer_limit(profile, monkeypatch):
    """Починка трёх стилей идёт через тот же планировщик процесса, что и
    письмо: больше его лимита одновременных вызовов модели не бывает."""
    monkeypatch.setattr(scheduler_module, "_DEFAULT", ModelScheduler(2))
    llm = _ShortenLLM()
    repairers = [SlideRepairer(profile=profile, llm=llm, sources=[], max_calls=10) for _ in range(3)]
    threads = [
        threading.Thread(target=r.shorten, args=(_long_slide(), "slide13", [])) for r in repairers for _ in range(3)
    ]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    assert len(llm.requests) == 9
    assert llm.max_concurrent <= 2


def test_contract_for_a_slide_keeps_its_unit_count(profile):
    contract = contract_for_slide(_long_slide(), "slide13", profile, "dense")

    assert contract.pattern_id == "slide13"
    assert contract.slots[0].block == "bullets" and contract.slots[0].count == 3, "единиц столько, сколько на слайде"


def test_planner_backups_come_right_after_the_chosen_layout(profile):
    """Вторая ступень берёт запасные планировщика в его порядке, а не
    собственный рейтинг сборки."""
    patterns = [builder._pattern_from_model(m) for m in profile.patterns]
    bullets = [p.pattern_id for p in patterns if p.kind == "bullets"]
    assert len(bullets) >= 3
    slide = replace(_long_slide(), pattern_id=bullets[0], alternatives=(bullets[2], bullets[1]))

    order = [p.pattern_id for p in builder._resolve_pattern(slide, patterns, profile, Variant.dense)]

    assert order[:3] == [bullets[0], bullets[2], bullets[1]]


def test_lost_blocks_count_content_without_a_slot(profile):
    """Сборка с нуля не принимает раскладку, куда не легли карточки, даже
    без находок аудита: `_lost_blocks` считает такие части."""
    grid = builder._grid_from_model(profile.grid)
    by_id = {m.pattern_id: builder._pattern_from_model(m) for m in profile.patterns}
    photos_only = next(pid for pid, p in by_id.items() if not any(s.role in ("body", "card_body") for s in p.slots))

    assert builder._lost_blocks(_cards_slide(4), by_id["slide21"], grid) == 0
    assert builder._lost_blocks(_cards_slide(4), by_id[photos_only], grid) >= 1


def test_no_model_calls_in_emergency_mode_or_without_time(profile):
    from deckforge.workflow.budget import BudgetPolicy, RunBudget, RunMode
    from deckforge.workflow.repair import repairer_for

    llm = _ShortenLLM()
    roomy = RunBudget.from_policy(BudgetPolicy(budget_seconds=300.0))
    assert repairer_for(roomy, profile, llm, [], "dense").llm is llm

    emergency = RunBudget.from_policy(BudgetPolicy(budget_seconds=300.0))
    emergency.mode = RunMode.EMERGENCY
    assert repairer_for(emergency, profile, llm, [], "dense").llm is None

    late = RunBudget.from_policy(BudgetPolicy(budget_seconds=10.0))
    assert repairer_for(late, profile, llm, [], "dense").llm is None
