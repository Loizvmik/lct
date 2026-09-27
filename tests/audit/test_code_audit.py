"""Задача T2: аудит не судит фрагмент кода как текст слайда (D02, T01,
T02), а L03 мерит у него ширину строки без переносов."""
from __future__ import annotations
from pathlib import Path

from pptx.util import Inches, Pt

from deckforge.audit.config import AuditConfig
from deckforge.audit.deterministic import run_deterministic
from deckforge.compose.code import CODE_SHAPE_NAME, add_code_block
from deckforge.ooxml.geometry import Box
from deckforge.plan.spec import CodeBlock
from deckforge.template.profile import TemplateProfile

TEMPLATE = Path("dataset/templates/VK Tech шаблон.pptx")
PROFILE = TemplateProfile.from_file(TEMPLATE, cache_dir=None)
CONFIG = AuditConfig.load()
CODE = CodeBlock(language="python", code="\n".join(
    f"    step_{i} = approve(request, level={i})  # согласующий уровня {i}" for i in range(12)
))


def _code_findings(findings):
    return [f for f in findings if f.shape_ref and CODE_SHAPE_NAME in f.shape_ref]


def _add_code_box(slide, width_in: float, size_pt: float, text: str):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(width_in), Inches(4.0))
    box.name = CODE_SHAPE_NAME
    tf = box.text_frame
    tf.word_wrap = False
    for i, line in enumerate(text.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size_pt)
        run.font.name = "Courier New"


def test_code_block_is_not_judged_as_slide_text(blank_deck):
    slide = blank_deck.add_slide()
    add_code_block(slide, Box(0.05, 0.2, 0.9, 0.7), CODE, PROFILE, 1.0)
    findings = run_deterministic(blank_deck.save_as("code.pptx"), PROFILE, CONFIG)
    ids = {f.check_id for f in _code_findings(findings)}
    assert not ids & {"D02", "T01", "T02", "L03", "L04"}, ids


def test_L03_catches_a_code_line_wider_than_its_frame(blank_deck):
    slide = blank_deck.add_slide()
    _add_code_box(slide, 3.0, 18, "x = 1\n" + "y = " + "a + " * 30 + "b")
    findings = _code_findings(run_deterministic(blank_deck.save_as("wide.pptx"), PROFILE, CONFIG))
    l03 = [f for f in findings if f.check_id == "L03"]
    assert len(l03) == 1 and "шире рамки" in l03[0].message
    assert not {f.check_id for f in findings} & {"T01", "T02"}


def test_short_code_in_a_wide_frame_passes_L03(blank_deck):
    slide = blank_deck.add_slide()
    _add_code_box(slide, 9.0, 14, "x = 1\ny = 2")
    findings = _code_findings(run_deterministic(blank_deck.save_as("ok.pptx"), PROFILE, CONFIG))
    assert not [f for f in findings if f.check_id in ("L03", "L04")]
