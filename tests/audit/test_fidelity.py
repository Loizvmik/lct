"""Тесты `audit.fidelity.template_fidelity` — задача T (метрики верности
шаблону). Не импортируется из `conftest.py` намеренно (тот же принцип, что
и в `tests/audit/test_deterministic.py`/`tests/compose/test_builder.py`:
`tests/` — не пакет, каждый модуль объявляет свои TEMPLATE/PROFILE)."""
from __future__ import annotations
from pathlib import Path

import pytest

from deckforge.audit.fidelity import EDITABILITY_KINDS, template_fidelity
from deckforge.compose.builder import Variant, build_deck
from deckforge.plan.spec import Card, CardBlock, DeckSpec, SlideSpec
from deckforge.template.profile import TemplateProfile

# VK Education: `slide21` — четыре кружка с номерами над четырьмя
# описаниями, тот же паттерн, которым `tests/compose/test_clone.py::
# test_build_deck_marks_cloned_slides` уже доказал, что клон реально
# срабатывает (см. её докстроку) — воспроизводим тот же рецепт, а не
# гадаем содержание для клона заново.
EDU_TEMPLATE = Path("dataset/templates/Шаблон презентации VK Education.pptx")
EDU_PROFILE = TemplateProfile.from_file(EDU_TEMPLATE, cache_dir=None)

# VK Tech — для сценариев, не завязанных конкретно на клон (аудит T01-T03,
# колода с нуля), тот же файл, что и `tests/audit/conftest.py`.
TECH_TEMPLATE = Path("dataset/templates/VK Tech шаблон.pptx")
TECH_PROFILE = TemplateProfile.from_file(TECH_TEMPLATE, cache_dir=None)


def _cards_spec(n: int) -> SlideSpec:
    return SlideSpec(
        index=0, kind="two_col", headline="Что нужно для раскатки", pattern_id="slide21",
        blocks=[CardBlock(items=[Card(body=f"Пункт номер {i + 1} плана раскатки") for i in range(n)])],
    )


def test_fidelity_on_cloned_deck_reports_high_clone_rate_and_shape_preservation():
    spec = DeckSpec(title="Клон", language="ru", slides=[_cards_spec(2)])
    pptx_path = build_deck(spec, EDU_PROFILE, EDU_TEMPLATE, Variant.dense, clone_examples=True)

    report = template_fidelity(pptx_path, spec, EDU_PROFILE)

    assert report.native_clone_rate == 1.0
    assert report.native_shape_preservation is not None
    assert 0.0 < report.native_shape_preservation <= 1.0
    assert report.mean_geometry_deviation is not None
    # Клон подгоняет незаполненные единицы повтора и рамку текста — коробка
    # заголовка/тела не обязана буквально совпасть с коробкой слота
    # паттерна, но остаётся заведомо близкой (доли холста), не где попало.
    assert report.mean_geometry_deviation < 0.2
    # Одна колода, один pattern_id — разнообразие максимальное (единица),
    # энтропия распределения одной корзины нормирована в ноль (см. докстроку
    # `_pattern_diversity_and_entropy`).
    assert report.pattern_diversity == 1.0
    assert report.pattern_entropy == 0.0
    assert "Верность шаблону" in report.summary


def test_fidelity_on_scratch_deck_reports_zero_clone_rate_and_notes_why():
    spec = DeckSpec(title="С нуля", language="ru", slides=[_cards_spec(2)])
    pptx_path = build_deck(spec, EDU_PROFILE, EDU_TEMPLATE, Variant.dense, clone_examples=False)

    report = template_fidelity(pptx_path, spec, EDU_PROFILE)

    assert report.native_clone_rate == 0.0
    assert report.native_shape_preservation is None
    assert report.mean_geometry_deviation is None
    assert any("клон" in note for note in report.notes)


def test_density_delta_is_computed_against_the_pattern_source_density():
    """У паттернов есть `source_density` (задача R): дельта плотности
    считается как число; когда поля нет ни у одного паттерна, метрика честно
    пропускается с пометкой, а соседние метрики не страдают."""
    spec = DeckSpec(title="Клон", language="ru", slides=[_cards_spec(2)])
    pptx_path = build_deck(spec, EDU_PROFILE, EDU_TEMPLATE, Variant.dense, clone_examples=True)

    report = template_fidelity(pptx_path, spec, EDU_PROFILE)
    assert report.density_delta is not None
    assert report.native_clone_rate == 1.0

    stripped = EDU_PROFILE.model_copy(update={
        "patterns": [p.model_copy(update={"source_density": None}) for p in EDU_PROFILE.patterns],
    })
    report_without = template_fidelity(pptx_path, spec, stripped)
    assert report_without.density_delta is None
    assert any("source_density" in note for note in report_without.notes)


def test_pattern_diversity_is_none_without_pattern_ids():
    """`SlideSpec.pattern_id` не проставлен ни на одном слайде — например,
    план, собранный до `apply_variant`."""
    spec = DeckSpec(
        title="Без раскладки", language="ru",
        slides=[SlideSpec(index=0, kind="section", headline="Заголовок")],
    )
    pptx_path = build_deck(spec, TECH_PROFILE, TECH_TEMPLATE, Variant.dense, clone_examples=False)

    report = template_fidelity(pptx_path, spec, TECH_PROFILE)

    assert report.pattern_diversity is None
    assert report.pattern_entropy is None
    assert any("pattern_id" in note for note in report.notes)


def test_typography_palette_compliance_is_a_fraction_in_unit_interval():
    spec = DeckSpec(title="С нуля", language="ru", slides=[_cards_spec(2)])
    pptx_path = build_deck(spec, EDU_PROFILE, EDU_TEMPLATE, Variant.dense, clone_examples=False)

    report = template_fidelity(pptx_path, spec, EDU_PROFILE)

    assert report.typography_palette_compliance is not None
    assert 0.0 <= report.typography_palette_compliance <= 1.0


def test_native_asset_usage_is_none_when_the_deck_has_no_pictures():
    spec = DeckSpec(title="Без картинок", language="ru", slides=[_cards_spec(2)])
    pptx_path = build_deck(spec, EDU_PROFILE, EDU_TEMPLATE, Variant.dense, clone_examples=False)

    report = template_fidelity(pptx_path, spec, EDU_PROFILE)

    if report.native_asset_usage is not None:
        assert 0.0 <= report.native_asset_usage <= 1.0
    else:
        assert any("картин" in note for note in report.notes)


def test_report_summary_mentions_every_metric_word():
    spec = DeckSpec(title="Клон", language="ru", slides=[_cards_spec(2)])
    pptx_path = build_deck(spec, EDU_PROFILE, EDU_TEMPLATE, Variant.dense, clone_examples=True)

    report = template_fidelity(pptx_path, spec, EDU_PROFILE)

    for word in ("клоном", "фигур", "типографика", "разнообразие", "нативных", "плотности"):
        assert word in report.summary, report.summary


# --- Задача V4: метрика редактируемости. ---


def test_editability_on_cloned_and_scratch_decks():
    """Колода из текста (клоном и с нуля): всё содержание редактируемо,
    счётчики по видам есть, текстовых объектов больше нуля."""
    for clone in (True, False):
        spec = DeckSpec(title="Редактируемость", language="ru", slides=[_cards_spec(2)])
        pptx_path = build_deck(spec, EDU_PROFILE, EDU_TEMPLATE, Variant.dense, clone_examples=clone)

        report = template_fidelity(pptx_path, spec, EDU_PROFILE)

        assert report.editable_content_coverage == 1.0, (clone, report.editable_counts, report.notes)
        assert report.editable_counts["editable_text"] > 0
        assert report.editable_counts["raster_content"] == 0
        assert set(report.editable_counts) == set(EDITABILITY_KINDS)
        assert "редактируемо 100%" in report.summary


def _synthetic_deck(tmp_path: Path) -> Path:
    """Четыре слайда: текст, родная таблица, родной график и слайд одной
    картинкой во весь холст."""
    import io

    from PIL import Image
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu

    prs = Presentation()
    width, height = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    s.shapes.add_textbox(Emu(0), Emu(0), Emu(width // 2), Emu(height // 4)).text_frame.text = "Текст"
    s = prs.slides.add_slide(blank)
    s.shapes.add_table(2, 2, Emu(0), Emu(0), Emu(width // 2), Emu(height // 4))
    s = prs.slides.add_slide(blank)
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(0), Emu(0), Emu(width // 2), Emu(height // 4), data)
    s = prs.slides.add_slide(blank)
    png = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 10, 10)).save(png, format="PNG")
    png.seek(0)
    s.shapes.add_picture(png, Emu(0), Emu(0), Emu(width), Emu(height))

    path = tmp_path / "synthetic.pptx"
    prs.save(str(path))
    return path


def test_editability_counts_native_objects_and_flags_slide_as_image(tmp_path):
    from pptx import Presentation

    from deckforge.audit.fidelity import _editability
    from deckforge.ooxml.geometry import Canvas

    path = _synthetic_deck(tmp_path)
    prs = Presentation(str(path))
    canvas = Canvas(prs.slide_width, prs.slide_height)
    notes: list[str] = []

    coverage, counts = _editability(path, prs, canvas, set(), notes)

    assert counts["editable_text"] == 1
    assert counts["native_table"] == 1
    assert counts["native_chart"] == 1
    assert counts["raster_content"] == 1, "картинка во весь пустой слайд не фото, а вёрстка в растре"
    assert counts["user_image"] == 0
    # Три объекта по 1/8 холста редактируемы, растр во весь холст нет.
    assert coverage == pytest.approx(3 * 0.125 / (3 * 0.125 + 1.0), abs=1e-3)


def test_template_picture_is_decor_not_content(tmp_path):
    """Та же картинка, если её байты есть в медиа шаблона, считается
    декором и в знаменатель не входит."""
    import hashlib

    from pptx import Presentation

    from deckforge.audit.fidelity import _editability
    from deckforge.ooxml.geometry import Canvas
    from deckforge.ooxml.package import PptxPackage

    path = _synthetic_deck(tmp_path)
    with PptxPackage.open(path) as pkg:
        hashes = {hashlib.md5(pkg.part(n)).hexdigest() for n in pkg.names() if n.startswith("ppt/media/")}
    prs = Presentation(str(path))
    coverage, counts = _editability(path, prs, Canvas(prs.slide_width, prs.slide_height), hashes, [])

    assert counts["template_decor"] >= 1
    assert counts["raster_content"] == 0
    assert coverage == 1.0
