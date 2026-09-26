"""Тесты `deckforge.audit.visual` (Task 12, Приложение 1 ТЗ, «Валидация
контента»).

Офлайн-тесты (без сети) проверяют инфраструктуру, которая обязана работать
одинаково при любом ответе модели: валидацию схемой, честную деградацию без
(мультимодальной) модели, устойчивость к невалидному ответу, число вызовов
модели на колоду (N слайдов + 1 коллаж, не 11×N — докстрока `visual.py`).

`@live`-тесты (`YANDEX_API_KEY` в окружении) идут в реальную модель по
реальным превью, отрендеренным LibreOffice (`soffice_available()`) — сеть и
недетерминированность ответа означают, что они НЕ гоняются в
`uv run pytest tests/ -v` без ключа и намеренно допускают некоторый допуск в
assert'ах (модель может ошибиться, см. докстроку `audit.visual`: «ТЗ
дословно: недетерминированная проверка... может дать разный ответ»)."""
from __future__ import annotations
import json
import os
from pathlib import Path

import pytest
from PIL import Image

from deckforge.audit.visual import (
    CHECK_IDS, DECK_LEVEL_CHECK_IDS, PER_SLIDE_CHECK_IDS, VisualAuditResult,
    _build_collage, _build_shape_manifest, _build_slide_prompt, _extract_scores, _findings_from_answers,
    _load_agent_prompt, _parse_answer, _parse_scores, _supports_vision, run_visual,
)
from deckforge.compose.builder import Variant, build_deck
from deckforge.plan.outline import SourceDoc
from deckforge.plan.spec import BulletBlock, DeckSpec, SlideSpec
from deckforge.provider.base import VisionProvider
from deckforge.render.soffice import soffice_available, to_pngs
from deckforge.template.profile import TemplateProfile

live = pytest.mark.skipif(not os.getenv("YANDEX_API_KEY"), reason="нет YANDEX_API_KEY")
needs_render = pytest.mark.skipif(not soffice_available(), reason="LibreOffice не установлен")

TEMPLATE = Path("dataset/templates/VK Tech шаблон.pptx")
PROFILE = TemplateProfile.from_file(TEMPLATE, cache_dir=None)


# ---------------------------------------------------------------------------
# Фейковые провайдеры
# ---------------------------------------------------------------------------


class _TextOnlyProvider:
    """Не мультимодальна и не наследует `VisionProvider` — ни `ask_image`,
    ни `.card.vision`. То же самое, что реальный текстовый провайдер без
    vision-возможности."""


class _AllOkVisionProvider(VisionProvider):
    """Всегда отвечает "всё хорошо" на запрошенное подмножество вопросов —
    считает, сколько раз её вызвали (`calls`), нужен тестам числа вызовов и
    устойчивости к параллельным потокам."""

    def __init__(self):
        self.calls: list[str] = []

    def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
        self.calls.append(prompt)
        keys = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])["answer_only_keys"]
        return json.dumps({k: {"ok": True} for k in keys})


class _BrokenJsonProvider(VisionProvider):
    """Мультимодальна (наследует `VisionProvider`), но отвечает не JSON —
    тест устойчивости пайплайна к невалидному ответу модели."""

    def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
        return "конечно! вот мой ответ без единой фигурной скобки"


class _CardVisionProvider:
    """Не наследует `VisionProvider`, но несёт `.card.vision` — тот же
    духк-тайпинг, каким `YandexProvider` сообщает о себе через
    `provider.registry.ModelCard` (`_supports_vision`, docstring
    `visual.py`)."""

    class _Card:
        def __init__(self, vision: bool):
            self.vision = vision

    def __init__(self, vision: bool):
        self.card = self._Card(vision)


def _tiny_png(color=(10, 20, 30)) -> Path:
    import tempfile
    path = Path(tempfile.mktemp(suffix=".png"))
    Image.new("RGB", (320, 180), color).save(path)
    return path


def _tiny_spec(n: int) -> DeckSpec:
    return DeckSpec(
        title="Тестовая колода",
        language="ru",
        slides=[
            SlideSpec(index=i, kind="bullets", headline=f"Слайд {i}", blocks=[BulletBlock(items=["пункт"])])
            for i in range(n)
        ],
    )


# ---------------------------------------------------------------------------
# AGENT.md
# ---------------------------------------------------------------------------


def test_agent_prompt_has_frontmatter_with_version():
    meta, body = _load_agent_prompt()
    assert meta["name"] == "content-auditor"
    assert meta["version"]
    assert meta["model_role"] == "content_audit"
    assert body


def test_agent_prompt_contains_all_eleven_questions_verbatim():
    _meta, body = _load_agent_prompt()
    questions = [
        "Заголовок содержит вывод, а не просто называет тему?",
        "Содержимое слайда соответствует заголовку?",
        "Слайд пересказывается одним предложением?",
        "Все цифры и факты со слайда есть в исходных материалах?",
        "На слайде есть содержание, а не только заголовок?",
        "Картинки и иконки относятся к теме слайда?",
        "Нет служебного мусора: реплик спикера, кусков промпта?",
        "Текст без опечаток?",
        "Вся колода на одном языке?",
        "Все строки таблицы и элементы легенды работают на мысль слайда?",
        "Соседние слайды связаны между собой по логике?",
    ]
    for q in questions:
        assert q in body


def test_check_ids_cover_eleven_questions_split_slide_vs_deck():
    assert len(CHECK_IDS) == 11
    assert set(PER_SLIDE_CHECK_IDS) | set(DECK_LEVEL_CHECK_IDS) == set(CHECK_IDS)
    assert set(PER_SLIDE_CHECK_IDS) & set(DECK_LEVEL_CHECK_IDS) == set()
    assert DECK_LEVEL_CHECK_IDS == ("C09", "C11")


# ---------------------------------------------------------------------------
# _supports_vision
# ---------------------------------------------------------------------------


def test_supports_vision_none_is_false():
    assert _supports_vision(None) is False


def test_supports_vision_false_for_text_only_provider():
    assert _supports_vision(_TextOnlyProvider()) is False


def test_supports_vision_true_for_vision_provider_subclass():
    assert _supports_vision(_AllOkVisionProvider()) is True


def test_supports_vision_reads_card_attribute():
    assert _supports_vision(_CardVisionProvider(vision=True)) is True
    assert _supports_vision(_CardVisionProvider(vision=False)) is False


# ---------------------------------------------------------------------------
# _parse_answer — валидация схемой
# ---------------------------------------------------------------------------


def test_parse_answer_accepts_valid_json():
    raw = json.dumps({"C01": {"ok": True}, "C05": {"ok": False, "where": "нет буллетов"}})
    result = _parse_answer(raw, ("C01", "C05"))
    assert result == {"C01": (True, None), "C05": (False, "нет буллетов")}


def test_parse_answer_strips_markdown_json_fence():
    """Живой прогон обязательной проверки задачи: qwen3.6 оборачивает ответ
    в ```json ... ``` вопреки AGENT.md ("никакого текста вне JSON") — это
    оформление, не сломанный формат, снимается перед разбором JSON."""
    raw = "```json\n" + json.dumps({"C01": {"ok": True}}) + "\n```"
    assert _parse_answer(raw, ("C01",)) == {"C01": (True, None)}


def test_parse_answer_strips_bare_markdown_fence_without_json_tag():
    raw = "```\n" + json.dumps({"C01": {"ok": True}}) + "\n```"
    assert _parse_answer(raw, ("C01",)) == {"C01": (True, None)}


def test_parse_answer_rejects_non_json():
    with pytest.raises(ValueError, match="JSON"):
        _parse_answer("не json", ("C01",))


def test_parse_answer_rejects_missing_key():
    with pytest.raises(ValueError, match="C05"):
        _parse_answer(json.dumps({"C01": {"ok": True}}), ("C01", "C05"))


def test_parse_answer_rejects_non_bool_ok():
    with pytest.raises(ValueError):
        _parse_answer(json.dumps({"C01": {"ok": "да"}}), ("C01",))


def test_parse_answer_rejects_top_level_non_object():
    with pytest.raises(ValueError, match="объект"):
        _parse_answer(json.dumps([1, 2, 3]), ("C01",))


# ---------------------------------------------------------------------------
# _parse_scores / _extract_scores — оценки PPTEval (задача G), необязательные
# ---------------------------------------------------------------------------


def test_parse_scores_accepts_valid_values():
    assert _parse_scores({"content": 4, "design": 3, "why": "ок"}, ("content", "design")) == {
        "content": 4, "design": 3, "why": "ок",
    }


def test_parse_scores_drops_out_of_range_int():
    assert _parse_scores({"content": 7}, ("content",)) is None


def test_parse_scores_drops_non_int_value():
    assert _parse_scores({"content": "четыре"}, ("content",)) is None


def test_parse_scores_drops_bool_value():
    """`bool` — подкласс `int` в Python; `True`/`False` не оценка 1-5."""
    assert _parse_scores({"content": True}, ("content",)) is None


def test_parse_scores_drops_non_string_why():
    result = _parse_scores({"content": 4, "why": 123}, ("content",))
    assert result == {"content": 4}


def test_parse_scores_returns_none_for_non_dict_input():
    assert _parse_scores("не объект", ("content",)) is None
    assert _parse_scores(None, ("content",)) is None


def test_parse_scores_returns_none_when_nothing_valid():
    assert _parse_scores({"content": -1, "why": ""}, ("content",)) is None


def test_extract_scores_reads_scores_key_from_raw_answer():
    raw = json.dumps({"C01": {"ok": True}, "scores": {"content": 5, "design": 5, "why": "идеально"}})
    assert _extract_scores(raw, ("content", "design")) == {"content": 5, "design": 5, "why": "идеально"}


def test_extract_scores_returns_none_when_scores_key_missing():
    raw = json.dumps({"C01": {"ok": True}})
    assert _extract_scores(raw, ("content", "design")) is None


def test_extract_scores_returns_none_on_unparseable_raw():
    """Невалидный JSON не должен ронять извлечение оценок — та же честная
    деградация, что и у остального модуля (докстрока `_extract_scores`)."""
    assert _extract_scores("не json вовсе", ("content",)) is None
    assert _extract_scores(None, ("content",)) is None


# ---------------------------------------------------------------------------
# _findings_from_answers
# ---------------------------------------------------------------------------


def test_findings_from_answers_skips_ok_true():
    findings = _findings_from_answers(0, {"C01": (True, None)})
    assert findings == []


def test_findings_from_answers_emits_for_ok_false():
    findings = _findings_from_answers(2, {"C05": (False, "только заголовок")})
    assert len(findings) == 1
    f = findings[0]
    assert f.check_id == "C05"
    assert f.slide_index == 2
    assert "только заголовок" in f.message
    assert f.fixable is False


# ---------------------------------------------------------------------------
# Манифест текстовых фигур (задача M: адресность находок) —
# `_build_shape_manifest`/`_findings_from_answers(manifest_roles=...)`
# ---------------------------------------------------------------------------


def test_findings_from_answers_uses_manifest_id_as_shape_ref():
    """Модель ответила `where` РОВНО id из манифеста слайда — находка несёт
    `shape_ref` в том же формате `"id:метка"`, что и `audit.deterministic`."""
    findings = _findings_from_answers(0, {"C01": (False, "42")}, {"42": "headline"})
    assert len(findings) == 1
    assert findings[0].shape_ref == "42:headline"
    assert "headline" in findings[0].message


def test_findings_from_answers_falls_back_to_raw_where_when_id_unknown():
    """`where` не совпал ни с одним id манифеста (модель описала место
    словами, как раньше, или манифест пуст) — старое поведение дословно."""
    findings = _findings_from_answers(0, {"C01": (False, "в правом верхнем углу")}, {"42": "headline"})
    assert findings[0].shape_ref is None
    assert "в правом верхнем углу" in findings[0].message


def test_build_slide_prompt_carries_shape_manifest_payload():
    spec = _tiny_spec(1)
    manifest = [{"id": "7", "role": "bullet", "text": "Пункт один"}]
    prompt = _build_slide_prompt("тело промпта", 0, 1, spec, spec.slides[0], [], "", manifest)
    payload = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])
    assert payload["shape_manifest"] == manifest


def test_build_slide_prompt_shape_manifest_defaults_to_empty_list():
    spec = _tiny_spec(1)
    prompt = _build_slide_prompt("тело промпта", 0, 1, spec, spec.slides[0], [], "")
    payload = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])
    assert payload["shape_manifest"] == []


def test_build_shape_manifest_matches_headline_and_bullet_roles():
    """Роль определяется сопоставлением текста реальной фигуры с текстом
    плана (заголовок/буллет), не разбором координат раскладки — граница
    `audit`/`compose`/`template` не нарушена (см. докстроку модуля)."""
    spec = DeckSpec(
        title="Колода", language="ru",
        slides=[SlideSpec(
            index=0, kind="bullets", headline="Наш главный вывод квартала",
            blocks=[BulletBlock(items=["Пункт первый про рост", "Пункт второй про удержание"])],
        )],
    )
    built = build_deck(spec, PROFILE, TEMPLATE, Variant.dense)
    from pptx import Presentation

    slide = Presentation(str(built)).slides[0]
    manifest, roles = _build_shape_manifest(slide, spec.slides[0])
    assert any(m["role"] == "headline" for m in manifest), manifest
    assert any(m["role"] == "bullet" for m in manifest), manifest
    assert set(roles.values()) >= {"headline", "bullet"}
    # Первые 60 знаков, не весь текст — манифест не должен раздувать промпт.
    assert all(len(m["text"]) <= 60 for m in manifest)


def test_build_shape_manifest_is_empty_without_pptx_slide():
    spec = _tiny_spec(1)
    assert _build_shape_manifest(None, spec.slides[0]) == ([], {})


def test_run_visual_sets_shape_ref_when_model_echoes_manifest_id(tmp_path):
    """Сквозной путь: `run_visual(pptx_path=...)` строит манифест, модель
    отвечает id этой фигуры в `where`, находка получает `shape_ref`."""
    spec = DeckSpec(
        title="Колода", language="ru",
        slides=[SlideSpec(
            index=0, kind="bullets", headline="Наш главный вывод",
            blocks=[BulletBlock(items=["Пункт один"])],
        )],
    )
    built = build_deck(spec, PROFILE, TEMPLATE, Variant.dense)
    pngs = [_tiny_png()]

    class _EchoManifestIdProvider(VisionProvider):
        def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
            payload = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])
            manifest = payload.get("shape_manifest") or []
            headline_id = next((m["id"] for m in manifest if m["role"] == "headline"), None)
            keys = payload["answer_only_keys"]
            answer = {k: {"ok": True} for k in keys}
            if headline_id is not None and "C01" in answer:
                answer["C01"] = {"ok": False, "where": headline_id}
            return json.dumps(answer)

    try:
        result = run_visual(
            pngs, spec, PROFILE, _EchoManifestIdProvider(), max_workers=1,
            pptx_path=built, deck_level=False,
        )
    finally:
        pngs[0].unlink(missing_ok=True)
    assert result.skipped_reason is None
    c01 = [f for f in result if f.check_id == "C01"]
    assert c01, "модель должна была найти headline в манифесте и ответить «нет» по C01"
    assert c01[0].shape_ref is not None and c01[0].shape_ref.endswith(":headline")


def test_run_visual_without_pptx_path_keeps_old_where_behavior(tmp_path):
    """`pptx_path=None` (по умолчанию) — манифест пуст, `where` остаётся
    текстом модели как есть, `shape_ref` не проставляется (иначе как
    раньше, бриф дословно)."""
    spec = _tiny_spec(1)
    pngs = [_tiny_png()]

    class _Provider(VisionProvider):
        def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
            payload = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])
            assert payload["shape_manifest"] == []
            keys = payload["answer_only_keys"]
            answer = {k: {"ok": True} for k in keys}
            if "C01" in answer:
                answer["C01"] = {"ok": False, "where": "заголовок слишком общий"}
            return json.dumps(answer)

    try:
        result = run_visual(pngs, spec, PROFILE, _Provider(), max_workers=1, deck_level=False)
    finally:
        pngs[0].unlink(missing_ok=True)
    c01 = [f for f in result if f.check_id == "C01"]
    assert c01 and c01[0].shape_ref is None
    assert "заголовок слишком общий" in c01[0].message


# ---------------------------------------------------------------------------
# Режимы N рискованных слайдов (задача M): FULL 6 / FAST 3 / EMERGENCY 0 —
# бриф просил 8/4, живой прогон на VK Education + queue-latency (реальный
# писатель и реальная модель аудита) показал 169.5с на 8 слайдов при 4
# потоках у dense (порог брифа — 90с), N снижен.
# ---------------------------------------------------------------------------


def test_config_modes_carry_task_m_slide_counts():
    from deckforge.workflow.budget import RunMode, load_policy

    policy = load_policy(Path("config/app.yaml"))
    assert policy.modes[RunMode.FULL].visual_audit_max_slides == 6
    assert policy.modes[RunMode.FAST].visual_audit_max_slides == 3
    assert policy.modes[RunMode.EMERGENCY].visual_audit_max_slides == 0


# ---------------------------------------------------------------------------
# _build_collage
# ---------------------------------------------------------------------------


def test_build_collage_returns_valid_png_taller_than_one_slide():
    paths = [_tiny_png((i * 30, 0, 0)) for i in range(3)]
    try:
        png_bytes = _build_collage(paths)
        import io
        im = Image.open(io.BytesIO(png_bytes))
        assert im.format == "PNG"
        single = Image.open(paths[0])
        assert im.height > single.height  # три слайда подряд, не один
    finally:
        for p in paths:
            p.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# run_visual — офлайн: честная деградация / устойчивость к мусору / число вызовов
# ---------------------------------------------------------------------------


def test_visual_audit_degrades_loudly_without_a_vision_model():
    """Если выбранная модель не мультимодальна, аудит обязан сказать об
    этом, а не молча вернуть пустой список и создать видимость проверки."""
    spec = _tiny_spec(1)
    pngs = [_tiny_png()]
    try:
        report = run_visual(pngs, spec, PROFILE, _TextOnlyProvider())
        assert report.skipped_reason and "мультимодал" in report.skipped_reason
        assert list(report) == []
    finally:
        pngs[0].unlink(missing_ok=True)


def test_visual_audit_degrades_loudly_without_any_provider():
    spec = _tiny_spec(1)
    pngs = [_tiny_png()]
    try:
        report = run_visual(pngs, spec, PROFILE, None)
        assert report.skipped_reason
        assert list(report) == []
    finally:
        pngs[0].unlink(missing_ok=True)


def test_malformed_model_answer_does_not_crash_the_pipeline():
    spec = _tiny_spec(2)
    pngs = [_tiny_png(), _tiny_png()]
    try:
        result = run_visual(pngs, spec, PROFILE, _BrokenJsonProvider(), max_workers=2)
        assert result.skipped_reason is None  # аудит РЕАЛЬНО выполнялся, просто ответы не разобрались
        findings = list(result)
        assert findings, "невалидный ответ модели обязан дать хоть одну находку C00, не тишину"
        assert all(f.check_id == "C00" for f in findings)
    finally:
        for p in pngs:
            p.unlink(missing_ok=True)


class _FailFirstThenOkProvider(VisionProvider):
    """Первый вызов по каждому ОТДЕЛЬНОМУ промпту (слайд/коллаж — у них
    разный текст промпта) отказывает так же, как qwen3.6 на пустом
    reasoning-ответе (`RuntimeError`, не JSON) — второй вызов по тому же
    промпту отвечает валидно. Ключ по ТЕКСТУ промпта, не по счётчику вызовов
    вообще — так тест не зависит от того, в каком порядке параллельные
    потоки пришли к провайдеру, только от того, что у КАЖДОГО отдельного
    запроса будет ровно одна повторная попытка."""

    def __init__(self):
        self.attempts_by_prompt: dict[str, int] = {}

    def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
        n = self.attempts_by_prompt.get(prompt, 0) + 1
        self.attempts_by_prompt[prompt] = n
        if n == 1:
            raise RuntimeError("будто весь бюджет max_tokens ушёл в reasoning, content пуст")
        keys = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])["answer_only_keys"]
        return json.dumps({k: {"ok": True} for k in keys})

    @property
    def total_calls(self) -> int:
        return sum(self.attempts_by_prompt.values())


class _AlwaysFailProvider(VisionProvider):
    """Отказывает на КАЖДОЙ попытке — проверяет, что повтор не превращается
    в бесконечный цикл и что после исчерпания попытки находка всё равно
    ровно одна на слайд/коллаж (не по одной на попытку)."""

    def __init__(self):
        self.call_count = 0

    def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
        self.call_count += 1
        raise RuntimeError("модель недоступна (тест)")


def test_model_silence_gets_one_retry_and_succeeds_on_second_attempt():
    """Бриф задачи: qwen3.6 иногда тратит весь бюджет на рассуждение и не
    отвечает вовсе — именно на самых кривых слайдах, где вердикт нужнее
    всего. Один повторный запрос (не наращивание max_tokens) часто
    срабатывает — второй заход обязан вернуть настоящий вердикт, а не
    C00, если он валиден."""
    spec = _tiny_spec(1)
    pngs = [_tiny_png()]
    provider = _FailFirstThenOkProvider()
    try:
        result = run_visual(pngs, spec, PROFILE, provider, max_workers=1)
        assert result.skipped_reason is None
        assert list(result) == [], "второй (успешный) заход обязан дать вердикт, а не находку C00"
        # слайд: 1 неудачная попытка + 1 удачная; коллаж: тоже отказывает
        # первый раз (тот же провайдер на всё) + удачная — итого 4 попытки.
        assert provider.total_calls == 4
    finally:
        pngs[0].unlink(missing_ok=True)


def test_model_silence_persisting_through_the_retry_still_yields_exactly_one_c00_per_unit():
    """Если и второй заход пуст — находка остаётся C00, РОВНО одна на
    слайд/коллаж (не по одной на попытку, и не молчание)."""
    spec = _tiny_spec(2)
    pngs = [_tiny_png(), _tiny_png((5, 5, 5))]
    provider = _AlwaysFailProvider()
    try:
        result = run_visual(pngs, spec, PROFILE, provider, max_workers=2)
        findings = list(result)
        assert len(findings) == 3, "2 слайда + 1 коллаж = 3 честные находки C00, не 6 (по одной на попытку)"
        assert all(f.check_id == "C00" for f in findings)
        # 2 попытки на КАЖДЫЙ из 3 запросов (2 слайда + коллаж) = 6 вызовов.
        assert provider.call_count == 6
    finally:
        for p in pngs:
            p.unlink(missing_ok=True)


def test_run_visual_makes_n_plus_one_model_calls():
    """N слайдов -> N вызовов по одному на слайд + 1 вызов на коллаж всей
    колоды (C09/C11) — не 11×N (докстрока `visual.py`)."""
    spec = _tiny_spec(4)
    pngs = [_tiny_png((i * 10, 0, 0)) for i in range(4)]
    provider = _AllOkVisionProvider()
    try:
        result = run_visual(pngs, spec, PROFILE, provider, max_workers=4)
        assert result.skipped_reason is None
        assert list(result) == []  # всё "ok": true -> ни одной находки
        assert result.model_calls == 5
        assert len(provider.calls) == 5
        assert result.slides_checked == 4
    finally:
        for p in pngs:
            p.unlink(missing_ok=True)


def test_run_visual_deck_level_call_asks_only_deck_level_keys():
    spec = _tiny_spec(2)
    pngs = [_tiny_png(), _tiny_png((0, 0, 99))]
    provider = _AllOkVisionProvider()
    try:
        run_visual(pngs, spec, PROFILE, provider, max_workers=2)
        keys_asked = [json.loads(p.rsplit("Служебные данные:\n", 1)[1])["answer_only_keys"] for p in provider.calls]
        per_slide_calls = [k for k in keys_asked if set(k) == set(PER_SLIDE_CHECK_IDS)]
        deck_calls = [k for k in keys_asked if set(k) == set(DECK_LEVEL_CHECK_IDS)]
        assert len(per_slide_calls) == 2
        assert len(deck_calls) == 1
    finally:
        for p in pngs:
            p.unlink(missing_ok=True)


def test_run_visual_findings_sorted_with_deck_level_first():
    """`slide_index=None` (проверки на колоду целиком) сортируется как -1 —
    тот же порядок, что и I01 (`slide_index=None`) в `audit.deterministic.
    _stable_sort`."""

    class _Provider(VisionProvider):
        def ask_image(self, png, prompt, *, max_tokens=1024):
            keys = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])["answer_only_keys"]
            return json.dumps({k: {"ok": False, "where": "тест"} for k in keys})

    spec = _tiny_spec(2)
    pngs = [_tiny_png(), _tiny_png((0, 5, 5))]
    try:
        result = run_visual(pngs, spec, PROFILE, _Provider(), max_workers=2)
        findings = list(result)
        assert findings[0].slide_index is None
        assert findings[0].check_id in DECK_LEVEL_CHECK_IDS
        assert all(f.slide_index in (0, None) for f in findings[:len(DECK_LEVEL_CHECK_IDS)])
    finally:
        for p in pngs:
            p.unlink(missing_ok=True)


def test_visual_audit_result_is_iterable_and_sized():
    result = VisualAuditResult(findings=[])
    assert len(result) == 0
    assert list(result) == []


# ---------------------------------------------------------------------------
# Оценки PPTEval (задача G) — необязательные, вне диапазона отбрасываются
# ---------------------------------------------------------------------------


class _ScoredProvider(VisionProvider):
    """Всегда отвечает "всё хорошо" на да/нет-вопросы плюс валидные оценки —
    слайд получает content/design, коллаж получает coherence."""

    def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
        keys = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])["answer_only_keys"]
        payload = {k: {"ok": True} for k in keys}
        if set(keys) == set(PER_SLIDE_CHECK_IDS):
            payload["scores"] = {"content": 4, "design": 3, "why": "ясно, но тесно"}
        else:
            payload["scores"] = {"coherence": 5, "why": "переходы логичны"}
        return json.dumps(payload)


def test_run_visual_collects_slide_and_deck_scores():
    spec = _tiny_spec(2)
    pngs = [_tiny_png(), _tiny_png((0, 0, 99))]
    try:
        result = run_visual(pngs, spec, PROFILE, _ScoredProvider(), max_workers=2)
        assert result.slide_scores == {
            0: {"content": 4, "design": 3, "why": "ясно, но тесно"},
            1: {"content": 4, "design": 3, "why": "ясно, но тесно"},
        }
        assert result.deck_score == {"coherence": 5, "why": "переходы логичны"}
        assert result.content_avg == 4.0
        assert result.design_avg == 3.0
    finally:
        for p in pngs:
            p.unlink(missing_ok=True)


class _NoScoreProvider(VisionProvider):
    """Отвечает валидно на да/нет-вопросы, но вообще не присылает `scores` —
    честный случай "модель не оценила", не ошибка."""

    def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
        keys = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])["answer_only_keys"]
        return json.dumps({k: {"ok": True} for k in keys})


def test_run_visual_without_scores_key_leaves_scores_empty():
    spec = _tiny_spec(1)
    pngs = [_tiny_png()]
    try:
        result = run_visual(pngs, spec, PROFILE, _NoScoreProvider(), max_workers=1)
        assert result.skipped_reason is None
        assert result.slide_scores == {}
        assert result.deck_score is None
        assert result.content_avg is None
        assert result.design_avg is None
    finally:
        pngs[0].unlink(missing_ok=True)


class _GarbageScoreProvider(VisionProvider):
    """Да/нет-ответы валидны, но `scores` — мусор разных сортов: оценка вне
    диапазона 1-5, оценка строкой вместо числа, `why` не строкой. Ни одна
    ось не должна попасть в результат, а сам аудит не должен упасть или
    получить находку C00 — это не невалидный ОТВЕТ (да/нет-часть в порядке),
    это невалидная ОЦЕНКА (бриф задачи G, тот же принцип, что и у C00)."""

    def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
        keys = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])["answer_only_keys"]
        payload = {k: {"ok": True} for k in keys}
        if set(keys) == set(PER_SLIDE_CHECK_IDS):
            payload["scores"] = {"content": 9, "design": "плохо", "why": 12345}
        else:
            payload["scores"] = {"coherence": 0, "why": ""}
        return json.dumps(payload)


def test_run_visual_drops_out_of_range_or_malformed_scores_without_failing_audit():
    spec = _tiny_spec(1)
    pngs = [_tiny_png()]
    try:
        result = run_visual(pngs, spec, PROFILE, _GarbageScoreProvider(), max_workers=1)
        assert result.skipped_reason is None
        assert list(result) == []  # да/нет-ответы валидны -> ни одной находки C00
        assert result.slide_scores == {}  # ни content (вне 1-5), ни design (не int) не прошли
        assert result.deck_score is None  # coherence=0 вне диапазона, why="" пусто
        assert result.content_avg is None
        assert result.design_avg is None
    finally:
        pngs[0].unlink(missing_ok=True)


def test_run_visual_partial_scores_average_only_over_slides_that_have_them():
    """Один слайд из двух получил оценку — среднее считается по одному, не
    делится на два (докстрока `_axis_average`: "не среди ВСЕХ слайдов")."""

    class _PartialProvider(VisionProvider):
        def ask_image(self, png: bytes, prompt: str, *, max_tokens: int = 1024) -> str:
            payload_keys = json.loads(prompt.rsplit("Служебные данные:\n", 1)[1])
            keys = payload_keys["answer_only_keys"]
            payload = {k: {"ok": True} for k in keys}
            if set(keys) == set(PER_SLIDE_CHECK_IDS) and payload_keys["slide_index_1based"] == 1:
                payload["scores"] = {"content": 2, "design": 2, "why": "первый слайд слабый"}
            return json.dumps(payload)

    spec = _tiny_spec(2)
    pngs = [_tiny_png(), _tiny_png((0, 0, 55))]
    try:
        result = run_visual(pngs, spec, PROFILE, _PartialProvider(), max_workers=1)
        assert result.slide_scores == {0: {"content": 2, "design": 2, "why": "первый слайд слабый"}}
        assert result.content_avg == 2.0
        assert result.design_avg == 2.0
    finally:
        for p in pngs:
            p.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# @live — реальная модель, реальный рендер
# ---------------------------------------------------------------------------


def _build_and_render(spec: DeckSpec, tmp_path: Path, name: str) -> list[Path]:
    built = build_deck(spec, PROFILE, TEMPLATE, Variant.dense)
    return to_pngs(built, tmp_path / name)


@pytest.fixture
def vlm():
    from deckforge.provider.yandex import YandexProvider
    # Дедлайн выше дефолта (60с, `DEFAULT_DEADLINE_SECONDS` в
    # `provider/yandex.py`): визуальный аудит стартует с бюджетом
    # `max_tokens`, заметно превышающим потолок остальных ролей конвейера
    # (`_PER_SLIDE_MAX_TOKENS`/`_DECK_LEVEL_MAX_TOKENS`, `audit/visual.py` —
    # живой замер там же), и честный (не зависший) ответ на таком бюджете
    # может занять больше 60с (обязательная проверка задачи: до ~46с на
    # слайде, замер не гарантирует, что это потолок).
    return YandexProvider(
        model="qwen3.6-35b-a3b",
        api_key=os.environ.get("YANDEX_API_KEY"),
        folder_id=os.environ.get("YANDEX_FOLDER_ID"),
        deadline_seconds=180.0,
    )


@live
@needs_render
def test_missing_body_is_caught(vlm, tmp_path):
    """C05: на слайде есть содержание, а не только заголовок — раздел-
    разделитель (`kind="section"`) намеренно несёт только заголовок,
    честный ответ модели на этот вопрос — "нет"."""
    spec = DeckSpec(
        title="Только заголовок",
        language="ru",
        slides=[SlideSpec(index=0, kind="section", headline="Итоги квартала")],
    )
    pngs = _build_and_render(spec, tmp_path, "title-only")
    result = run_visual(pngs, spec, PROFILE, vlm, max_workers=1)
    assert result.skipped_reason is None
    assert "C05" in {f.check_id for f in result}


@live
@needs_render
def test_invented_number_is_caught(vlm, tmp_path):
    """C04: все цифры со слайда есть в исходных материалах — слайд несёт
    полностью выдуманную статистику, которой нет ни в edu-platform
    sources.md, ни где-либо ещё."""
    sources_text = (Path("fixtures/content-packs/edu-platform") / "sources.md").read_text(encoding="utf-8")
    spec = DeckSpec(
        title="Итоги учебной платформы",
        language="ru",
        slides=[
            SlideSpec(
                index=0, kind="bullets", headline="Аномальный рост вовлечённости",
                blocks=[BulletBlock(items=[
                    "Выдуманная метрика: 314 592 активных пользователя за один день",
                    "Рост показателя на 812% за неделю — рекорд платформы",
                ])],
            ),
        ],
    )
    pngs = _build_and_render(spec, tmp_path, "invented-number")
    result = run_visual(
        pngs, spec, PROFILE, vlm, sources=[SourceDoc(name="sources.md", text=sources_text)], max_workers=1,
    )
    assert result.skipped_reason is None
    assert "C04" in {f.check_id for f in result}


@live
@needs_render
def test_clean_slide_passes_most_checks(vlm, tmp_path):
    sources_text = "Доведение курса до конца выросло с 41% до 57% в четвёртом квартале."
    spec = DeckSpec(
        title="Доведение курса выросло",
        language="ru",
        slides=[
            SlideSpec(
                index=0, kind="bullets", headline="Доведение курса выросло с 41% до 57%",
                blocks=[BulletBlock(items=[
                    "Доведение до конца выросло с 41% в первом квартале до 57% в четвёртом",
                    "Рост обеспечили три изменения, внедрённые в октябре",
                ])],
                source_note="Источник: sources.md",
            ),
        ],
    )
    pngs = _build_and_render(spec, tmp_path, "clean")
    result = run_visual(
        pngs, spec, PROFILE, vlm, sources=[SourceDoc(name="sources.md", text=sources_text)], max_workers=1,
    )
    assert result.skipped_reason is None
    assert len(result) <= 2


@live
@needs_render
def test_deterministic_finding_slide_also_flagged_by_visual_audit(vlm, tmp_path):
    """Слайд, на котором детерминированный аудит уже нашёл проблему (пустой
    почти без содержания слайд, вручную опустошённый после сборки) —
    визуальный аудит, глядя на ту же картинку, обязан заметить то же самое
    (C05)."""
    from pptx import Presentation

    spec = DeckSpec(
        title="Будет опустошён",
        language="ru",
        slides=[
            SlideSpec(
                index=0, kind="bullets", headline="Итоги пилота",
                blocks=[BulletBlock(items=["Пилот прошёл успешно, показатели выросли"])],
            ),
        ],
    )
    built = build_deck(spec, PROFILE, TEMPLATE, Variant.dense)

    from deckforge.audit.config import AuditConfig
    from deckforge.audit.deterministic import run_deterministic

    prs = Presentation(str(built))
    slide = prs.slides[0]
    # Опустошаем все текстовые рамки, кроме заголовка — тот же дефект,
    # что ловит D02/I03 детерминированного аудита ("пустой/почти пустой
    # слайд"), и одновременно C05 визуального.
    from deckforge.ooxml.ns import qn
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.text_frame.text.strip() == "Итоги пилота":
            continue
        for p in list(shape.text_frame._txBody.findall(qn("a:p"))):
            for r in list(p.findall(qn("a:r"))):
                p.remove(r)
    emptied = tmp_path / "emptied.pptx"
    prs.save(str(emptied))

    config = AuditConfig.load()
    det_findings = run_deterministic(emptied, PROFILE, config)
    det_ids = {f.check_id for f in det_findings}

    pngs = to_pngs(emptied, tmp_path / "emptied-png")
    vis_result = run_visual(pngs, spec, PROFILE, vlm, max_workers=1)
    assert vis_result.skipped_reason is None
    vis_ids = {f.check_id for f in vis_result}

    # Не требуем БУКВАЛЬНО одного и того же check_id (детерминированный и
    # визуальный аудиты используют разные идентификаторы) — требуем, что
    # ОБА заметили дефект на этом слайде: детерминированный обязан (это его
    # прямая работа), визуальный — по крайней мере C05 ("есть содержание,
    # а не только заголовок").
    assert det_ids, "детерминированный аудит обязан заметить опустошённый слайд"
    assert "C05" in vis_ids, f"визуальный аудит не заметил опустошённый слайд, ответил: {vis_ids}"
