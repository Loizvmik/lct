"""Тесты `deckforge.audit.autofix.apply_fixes` (Task 16 API/веб-интерфейс —
`POST /api/decks/{id}/fix` и `autofix=True` по умолчанию у `POST /api/
decks` опираются на этот модуль, см. его докстроку про границы честности:
покрыты все тринадцать `check_id` с `fixable=True`, только они).

Использует ту же инфраструктуру, что и `test_deterministic.py`
(`deck_with`/`PROFILE`/`CONFIG` из `tests/audit/conftest.py`, тот же приём
внедрения ровно одного дефекта в заведомо чистую колоду) — не своя копия:
`apply_fixes` работает НАД РЕЗУЛЬТАТОМ `run_deterministic`, тестировать его
без настоящих находок того же аудита было бы проверкой другого контракта."""
from __future__ import annotations
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from deckforge.audit.autofix import SUPPORTED_CHECKS, apply_fixes
from deckforge.audit.deterministic import run_deterministic

from .conftest import CONFIG, PROFILE
from .test_deterministic import _add_text, _ids, _make_bulleted


def _findings_by_check(path, check_id: str):
    return [f for f in run_deterministic(path, PROFILE, CONFIG) if f.check_id == check_id]


def test_apply_fixes_moves_out_of_bounds_shape_back_onto_the_slide(deck_with):
    path = deck_with(lambda s: s.shapes.add_textbox(Inches(20), Inches(1), Inches(3), Inches(1)))
    findings = _findings_by_check(path, "L01")
    assert findings, "фикстура обязана воспроизводить L01"

    result = apply_fixes(path, PROFILE, findings, {"f0": findings[0]})
    assert result.applied == ["f0"]
    assert result.changed

    assert "L01" not in _ids(run_deterministic(path, PROFILE, CONFIG))


def test_apply_fixes_shrinks_text_that_overflows_its_frame(deck_with):
    path = deck_with(lambda s: _add_text(s, 0.3, 1.3, 2.0, 0.3, ("слово " * 5).strip(), size_pt=24, family="Play"))
    findings = _findings_by_check(path, "L03")
    assert findings, "фикстура обязана воспроизводить L03"

    result = apply_fixes(path, PROFILE, findings, {"f0": findings[0]})
    assert result.applied == ["f0"]

    assert "L03" not in _ids(run_deterministic(path, PROFILE, CONFIG))


def test_apply_fixes_snaps_off_scale_font_to_the_type_scale(deck_with):
    path = deck_with(lambda s: _add_text(s, 0.3, 1.3, 3.0, 1.5, "Кегль не по шкале", size_pt=37, family="Play"))
    findings = _findings_by_check(path, "T02")
    assert findings, "фикстура обязана воспроизводить T02"

    result = apply_fixes(path, PROFILE, findings, {"f0": findings[0]})
    assert result.applied == ["f0"]

    assert "T02" not in _ids(run_deterministic(path, PROFILE, CONFIG))


def test_apply_fixes_flips_low_contrast_text_to_a_readable_color(deck_with):
    def _fn(slide):
        _add_text(slide, 0.5, 1.5, 2.5, 0.6, "Едва видно", size_pt=20, color_hex="FFFFFF")

    path = deck_with(_fn)
    findings = _findings_by_check(path, "T06")
    assert findings, "белый текст на светлом фоне слайда обязан воспроизвести T06"

    result = apply_fixes(path, PROFILE, findings, {"f0": findings[0]})
    assert result.applied == ["f0"]

    assert "T06" not in _ids(run_deterministic(path, PROFILE, CONFIG))


def test_apply_fixes_skips_findings_outside_the_supported_check_set(deck_with):
    """D01 (слишком много буллетов) размечена `fixable=False` в
    `audit.deterministic` — `apply_fixes` обязана честно пропустить её, а не
    сделать вид, что починила то, для чего нет эвристики."""
    def _fn(slide):
        box = slide.shapes.add_textbox(Inches(0.3), Inches(1.3), Inches(3.5), Inches(3.5))
        tf = box.text_frame
        tf.word_wrap = True
        for i in range(8):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = f"Пункт номер {i}"
            run.font.size = Pt(14)
            _make_bulleted(p)

    path = deck_with(_fn)
    findings = run_deterministic(path, PROFILE, CONFIG)
    unsupported = [f for f in findings if not f.fixable]
    assert unsupported, "фикстура обязана воспроизводить D01 (fixable=False)"

    finding = unsupported[0]
    result = apply_fixes(path, PROFILE, findings, {"f0": finding})
    assert result.applied == []
    assert result.skipped == ["f0"]
    assert not result.changed


def test_apply_fixes_only_touches_the_requested_finding(deck_with):
    """Два независимых дефекта на разных фигурах — просим починить только
    один (L01), второй (T06) обязан остаться нетронутым."""
    def _fn(slide):
        slide.shapes.add_textbox(Inches(20), Inches(1), Inches(3), Inches(1))
        _add_text(slide, 0.5, 3.0, 2.5, 0.6, "Едва видно", size_pt=20, color_hex="FFFFFF")

    path = deck_with(_fn)
    findings = run_deterministic(path, PROFILE, CONFIG)
    ids_before = _ids(findings)
    assert {"L01", "T06"} <= ids_before

    l01 = next(f for f in findings if f.check_id == "L01")
    result = apply_fixes(path, PROFILE, findings, {"only-l01": l01})
    assert result.applied == ["only-l01"]

    ids_after = _ids(run_deterministic(path, PROFILE, CONFIG))
    assert "L01" not in ids_after
    assert "T06" in ids_after


def test_every_fixable_check_id_is_in_supported_checks():
    """Защита от повторения истории I06: находка обещала fixable=True, а
    `apply_fixes` её тихо не чинил (не было ни в `SUPPORTED_CHECKS`, ни в
    `_FIXERS` — `apply_fixes` клал такую находку в skipped молча).

    Статически разбирает `audit/deterministic.py` (AST, не импорт и вызов
    каждой проверки — часть checks нужны специфичные дефекты, которые
    дорого/сложно каждый раз собирать): у КАЖДОГО литерального
    `fixable=True` (что в вызовах `_finding(...)`, что в прямом
    `Finding(...)`) обязан быть check_id из `SUPPORTED_CHECKS`. Ловит и
    будущий check с `fixable=True`, для которого забыли завести фиксер."""
    import ast
    import inspect

    from deckforge.audit import deterministic as det_module

    source = inspect.getsource(det_module)
    tree = ast.parse(source)

    finding_def = next(
        node for node in ast.walk(tree)
        if isinstance(node, ast.FunctionDef) and node.name == "_finding"
    )
    finding_param_names = [a.arg for a in finding_def.args.args]

    def _literal(node):
        return node.value if isinstance(node, ast.Constant) else None

    def _call_args(call: ast.Call, param_names: list[str]) -> dict[str, object]:
        args: dict[str, object] = {}
        for i, arg in enumerate(call.args):
            if i < len(param_names):
                args[param_names[i]] = _literal(arg)
        for kw in call.keywords:
            if kw.arg is not None:
                args[kw.arg] = _literal(kw.value)
        return args

    offenders = []
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call) or not isinstance(node.func, ast.Name):
            continue
        if node.func.id == "_finding":
            args = _call_args(node, finding_param_names)
        elif node.func.id == "Finding":
            args = _call_args(node, [])  # Finding(...) зовётся только с keyword'ами
        else:
            continue
        check_id, fixable = args.get("check_id"), args.get("fixable")
        if fixable is True and isinstance(check_id, str) and check_id not in SUPPORTED_CHECKS:
            offenders.append((check_id, node.lineno))

    assert not offenders, (
        f"check_id с fixable=True вне SUPPORTED_CHECKS (apply_fixes тихо их пропустит): {offenders}"
    )
