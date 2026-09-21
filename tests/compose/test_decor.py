"""Тесты переноса декора раскладки (правка по итогам визуального ревью,
находка №4 отчёта Task 9: мокап телефона на VK Education, слайд 2, залит
сплошным средним цветом картинки вместо экрана — `DecorShape` не несёт
байтов изображения (только средний цвет заливки, см. `template/patterns.
py::_to_decor`), и закрашивать контейнер под фото этим средним цветом хуже,
чем не закрашивать вовсе (см. `decor.py`, обоснование ниже)."""
from __future__ import annotations

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from deckforge.compose.decor import apply_decor
from deckforge.ooxml.geometry import Box
from deckforge.template.patterns import DecorShape

CANVAS_W = 9144000
CANVAS_H = 5143500


def _blank_slide():
    prs = Presentation()
    prs.slide_width = Emu(CANVAS_W)
    prs.slide_height = Emu(CANVAS_H)
    layout = prs.slide_layouts[6]
    return prs.slides.add_slide(layout)


def test_picture_filled_shape_is_not_painted_as_a_solid_block():
    """Мокап телефона/фото-контейнер — автофигура (`kind="shape"`) с
    картиночной заливкой (`fill_kind="picture"`): `fill_hex` в этом случае —
    средний цвет ФОТО (см. `_picture_fill_average_color`), не осознанный
    выбор дизайна. Закрасить им контейнер — нарисовать чужеродную сплошную
    плашку там, где должен быть снимок экрана; декор такого вида не
    переносится вовсе (см. докстроку `apply_decor`)."""
    slide = _blank_slide()
    shape = DecorShape(
        kind="shape", box=Box(0.3, 0.3, 0.2, 0.35), rotation=0.0, flip_h=False, flip_v=False,
        fill_hex="#1F3A93", has_fill=True, fill_kind="picture",
    )
    apply_decor(slide, [shape], CANVAS_W, CANVAS_H)
    assert list(slide.shapes) == []


def test_solid_filled_shape_is_still_painted():
    """Регрессия: обычная цветная плашка (`fill_kind="solid"`) — вид декора,
    который эта правка не должна трогать, — по-прежнему переносится."""
    slide = _blank_slide()
    shape = DecorShape(
        kind="shape", box=Box(0.3, 0.3, 0.2, 0.35), rotation=0.0, flip_h=False, flip_v=False,
        fill_hex="#1F3A93", has_fill=True, fill_kind="solid",
    )
    apply_decor(slide, [shape], CANVAS_W, CANVAS_H)
    autoshapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(autoshapes) == 1
