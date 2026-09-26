"""Сборка с нуля в стиле шаблона (задача U): заголовок цветом и кеглем
заголовков шаблона, как у клонов, текст не мельче body, фото не меньше
четверти холста, разбросанные подписи в один ряд.

Шаблон VK Education из `dataset/templates/`: заголовки примеров цвета не
пишут, наследуют синий от мастера, и сборка с нуля раньше ставила чёрный
48 при синих 36 у клонов. Профиль разбирается без модели."""
from __future__ import annotations
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.compose import builder
from deckforge.compose.clone import sample_slides_by_number
from deckforge.ooxml.geometry import Box, Canvas
from deckforge.plan.spec import BulletBlock, DeckSpec, SlideSpec, TextBlock, Visual
from deckforge.plan.variants import Variant

TEMPLATE = Path("dataset/templates/Шаблон презентации VK Education.pptx")
PHOTO = Path("fixtures/content-packs/queue-latency/photos/approver-laptop-review.jpg")


@pytest.fixture(scope="module")
def profile(profile_fixture):
    return profile_fixture(TEMPLATE.name)


@pytest.fixture(scope="module")
def look(profile):
    prs = Presentation(str(TEMPLATE))
    patterns = [builder._pattern_from_model(m) for m in profile.patterns]
    canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)
    return builder.template_look(profile, patterns, sample_slides_by_number(prs), canvas)


def _slide(**kw) -> SlideSpec:
    base = dict(
        index=0, kind="bullets", headline="Заявка ждёт 98,5% времени", pattern_id="slide13",
        blocks=[BulletBlock(items=["Медиана 31,5 часа", "Ожидание 98,5%"]), TextBlock(text="Ручная маршрутизация"),
                TextBlock(text="Ожидание согласования")],
    )
    base.update(kw)
    return SlideSpec(**base)


def _build(profile, slide: SlideSpec, **kw):
    spec = DeckSpec(title="Облик", language="ru", slides=[slide])
    out = builder.build_deck(spec, profile, TEMPLATE, Variant.visual, clone_examples=False, **kw)
    return Presentation(str(out)).slides[0], spec


def _run_of(slide, text: str):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() == text:
            return shape.text_frame.paragraphs[0].runs[0]
    raise AssertionError(f"нет фигуры с текстом {text!r}")


def test_template_look_is_the_inherited_blue_headline(look):
    assert look.headline_colors[0] == "#0077FF"
    assert 36 <= look.headline_pt <= 48


def test_scratch_headline_takes_template_colour_and_size(profile, look):
    slide, spec = _build(profile, _slide())

    assert builder.ladder_counts(spec)["scratch"] == 1
    run = _run_of(slide, "Заявка ждёт 98,5% времени")
    assert str(run.font.color.rgb) == look.headline_colors[0].lstrip("#")
    assert run.font.size.pt == pytest.approx(look.headline_pt, abs=0.6)


def test_scratch_body_text_is_not_smaller_than_body_step(profile):
    slide, _spec = _build(profile, _slide())

    body_pt = profile.type_scale_pt("body")
    run = _run_of(slide, "Ручная маршрутизация")
    assert run.font.size.pt >= body_pt - 0.05


def test_scratch_photo_is_at_least_a_quarter_of_the_canvas_wide(profile):
    slide, _spec = _build(
        profile, _slide(visual=Visual(kind="photo", photo_name="p.jpg")), user_photos={"p.jpg": PHOTO},
    )

    pictures = [s for s in slide.shapes if s.shape_type == 13]
    widest = max(p.width for p in pictures) / profile.canvas_width_emu
    # Картинка вписывается в коробку «contain»: по ширине не меньше
    # коробки, если фото шире её пропорций, иначе по высоте. Коробка не
    # уже четверти холста, а у горизонтального фото ширина и есть коробка.
    assert widest >= 0.25 - 0.005


def test_scattered_labels_share_one_row(profile):
    slide, _spec = _build(profile, _slide())

    tops = [
        s.top for s in slide.shapes
        if s.has_text_frame and s.text_frame.text.strip() in ("Ручная маршрутизация", "Ожидание согласования")
    ]
    assert len(tops) == 2
    assert abs(tops[0] - tops[1]) <= 0.01 * profile.canvas_height_emu


def test_at_least_wide_grows_from_the_centre_and_stays_in_margins(profile):
    grid = builder._grid_from_model(profile.grid)
    small = Box(left=0.85, top=0.6, width=0.06, height=0.1)

    grown = builder._at_least_wide(small, 0.25, grid)

    assert grown.width == pytest.approx(0.25)
    assert grown.height / grown.width == pytest.approx(small.height / small.width, rel=0.05)
    assert grown.left + grown.width <= 1.0 - grid.margin_right + 1e-9
