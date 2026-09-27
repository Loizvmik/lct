"""Задача T6: три дефекта живого прогона full-coverage.

1. Код на плашке в полслайда обрезался краем (VK Education, раскладка
   slide15): кегль до 9 pt, перенос длинной строки, плашка по высоте кода.
2. Заголовок в две строки при сборке с нуля ложился на подзаголовок
   (WorkSpace, slide17): места под заголовком ниже его набранного низа.
3. Клон таймлайна ЛЦТ2026 (slide25) под четыре карточки оставлял линии
   пятой единицы; таблетка заголовка наезжала на логотипы фона.

Профили разбираются без модели (`cache_dir=None`)."""
from __future__ import annotations
import functools
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.compose import builder
from deckforge.compose.blocks import Paragraph, SlotContent
from deckforge.compose.clone import sample_slides_by_number, shape_text, slide_refs
from deckforge.compose.code import (
    CODE_SHAPE_NAME, MIN_CODE_PT, MONO_ADVANCE, add_code_block, fit_code_lines, wrap_code_lines,
)
from deckforge.compose.colorpick import _relative_luminance
from deckforge.compose.kpi import text_slot
from deckforge.ooxml.geometry import Box, Canvas
from deckforge.plan.spec import Card, CardBlock, CodeBlock, SlideSpec
from deckforge.template.profile import TemplateProfile

TEMPLATES = Path("dataset/templates")
LCT = TEMPLATES / "ЛЦТ2026 Шаблон презентации.pptx"
ROUTE = CodeBlock(language="python", code=(
    "def route(request: Request) -> list[Approver]:\n"
    "    if request.kind == \"purchase\" and request.amount > 5_000_000:\n"
    "        return [head_of(request.unit), cfo()]\n"
    "    if request.kind == \"purchase\":\n"
    "        return [head_of(request.unit)]\n"
    "    return [line_manager(request.author)]"
))


@functools.lru_cache(maxsize=None)
def _lct_profile() -> TemplateProfile:
    return TemplateProfile.from_file(LCT, cache_dir=None)


def _contrast(a: str, b: str) -> float:
    la, lb = sorted((_relative_luminance(a), _relative_luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


# --- 1. код -------------------------------------------------------------------


def test_long_code_line_wraps_at_a_space_with_a_continuation_indent():
    line = '    if request.kind == "purchase" and request.amount > 5_000_000:'
    wrapped = wrap_code_lines([line, "    return x"], 40)
    assert all(len(w) <= 40 for w in wrapped)
    assert wrapped[-1] == "    return x"
    assert wrapped[0].startswith("    if ") and wrapped[1].startswith(" " * 8)
    assert " ".join(w.strip() for w in wrapped[:-1]) == line.strip()


def test_code_goes_down_to_nine_points_before_wrapping():
    lines = ROUTE.code.splitlines()
    longest = max(len(x) for x in lines)
    # Ширина, где самая длинная строка ложится только на 9 pt.
    width = longest * MONO_ADVANCE * 9.2 / 72
    kept, size = fit_code_lines(lines, width, 4.0, 14.0)
    assert size.fits and size.size_pt == MIN_CODE_PT == 9.0 and kept == lines
    narrow, size = fit_code_lines(lines, width * 0.7, 4.0, 14.0)
    assert size.fits and len(narrow) > len(lines)
    assert all(len(x) * MONO_ADVANCE * size.size_pt / 72 <= width * 0.7 + 1e-6 for x in narrow)


def test_scratch_code_plate_fits_its_column_and_hugs_the_code(profile_fixture, new_slide):
    """Правая тёмная половина VK Education slide15: плашка не шире колонки,
    строки не шире плашки, высота по коду, текст контрастен плашке."""
    profile = profile_fixture("Шаблон презентации VK Education.pptx")
    slide = new_slide()
    box = Box(0.545, 0.1, 0.4, 0.85)
    notes = add_code_block(slide, box, ROUTE, profile, 0.0, caption="Маршрутизация")
    assert notes == []
    plate = next(sh for sh in slide.shapes if sh.name == f"{CODE_SHAPE_NAME} plate")
    code = next(sh for sh in slide.shapes if sh.name == CODE_SHAPE_NAME)
    width_emu = profile.canvas_width_emu
    assert plate.left + plate.width <= width_emu
    assert plate.height < box.height * profile.canvas_height_emu / 2, "плашка по коду, не во всю колонку"
    runs = [r for p in code.text_frame.paragraphs for r in p.runs]
    size = runs[0].font.size.pt
    assert size >= MIN_CODE_PT
    for p in code.text_frame.paragraphs:
        assert len(p.text) * MONO_ADVANCE * size / 72 * 914400 <= code.width + 1
    fill = "#" + str(plate.fill.fore_color.rgb)
    assert _contrast(fill, "#" + str(runs[0].font.color.rgb)) >= 4.5


def test_code_plate_never_leaves_the_canvas(profile_fixture, new_slide):
    profile = profile_fixture("Шаблон презентации VK Education.pptx")
    slide = new_slide()
    add_code_block(slide, Box(0.7, 0.1, 0.5, 0.5), ROUTE, profile, 1.0)
    for sh in slide.shapes:
        assert sh.left + sh.width <= profile.canvas_width_emu + 1


# --- 2. заголовок в две строки ------------------------------------------------


def _headline_and_note(profile, headline: str) -> list[SlotContent]:
    norm = profile.canvas_norm
    head = SlotContent(
        text_slot("headline", Box(0.035, 0.062, 0.642, 0.119), 36.0 * norm), "headline", [Paragraph(headline)],
    )
    note = SlotContent(
        text_slot("body", Box(0.035, 0.181, 0.298, 0.094), 14.0 * norm), "body",
        [Paragraph("Пилот: 4 подразделения, 1 240 заявок, 340 согласующих.")],
    )
    return [head, note]


def test_text_under_a_two_line_headline_starts_below_it(profile_fixture):
    """WorkSpace, slide17: «Медиана сократилась с 31,5 до 6,2 часов» в две
    строки кеглем 36 ложилась на подпись под коробкой заголовка."""
    profile = profile_fixture("VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx")
    grid = builder._grid_from_model(profile.grid)
    contents = _headline_and_note(profile, "Медиана сократилась с 31,5 до 6,2 часов")
    bottom, _size = builder._headline_text_bottom(contents[0], profile, None)
    assert bottom > contents[1].slot.box.top, "без правки подпись под второй строкой"

    out = builder._clear_headline(contents, profile, None, grid)

    assert out[1].slot.box.top >= bottom
    assert out[0].slot.box.bottom >= bottom - 1e-6, "коробка заголовка по набранной высоте"
    assert out[1].slot.box.left == contents[1].slot.box.left


def test_one_line_headline_leaves_the_layout_alone(profile_fixture):
    profile = profile_fixture("VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx")
    grid = builder._grid_from_model(profile.grid)
    contents = _headline_and_note(profile, "Итоги пилота")
    out = builder._clear_headline(contents, profile, None, grid)
    assert out[1].slot.box == contents[1].slot.box


# --- 3. таймлайн ЛЦТ2026 ------------------------------------------------------


@pytest.fixture
def lct_timeline():
    """Клон slide25 ЛЦТ2026 под четыре карточки из пяти."""
    profile = _lct_profile()
    pattern = next(builder._pattern_from_model(m) for m in profile.patterns if m.pattern_id == "slide25")
    prs = Presentation(str(LCT))
    sources = sample_slides_by_number(prs)
    builder._clear_sample_slides(prs)
    spec = SlideSpec(
        index=2, kind="timeline", headline="92% времени — ожидание, а не работа согласующих",
        blocks=[CardBlock(items=[
            Card(title=f"Этап {i + 1}", body=f"{i + 1}-я часть времени заявки уходит на этот этап") for i in range(4)
        ])],
    )
    outcome = builder.place_slide_by_clone(prs, spec, pattern, profile, AuditConfig.load(), sources[25])
    assert outcome.reason is None, outcome.reason
    canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)
    return prs.slides[-1], pattern, canvas


def test_timeline_keeps_no_lines_outside_the_filled_units(lct_timeline):
    slide, pattern, canvas = lct_timeline
    cells = builder._unit_cells(pattern)
    assert len(cells) == 5
    centers = sorted(c.left + c.width / 2 for c in cells)
    half = min(b - a for a, b in zip(centers, centers[1:])) / 2
    edge = centers[3] + half
    lines = [r for r in slide_refs(slide, canvas) if r.kind == "connector"]
    assert lines, "линии заполненных единиц остаются"
    stray = [r.box for r in lines if r.box.right > edge]
    assert stray == []
    numbers = sorted(shape_text(r.element).strip() for r in slide_refs(slide, canvas)
                     if shape_text(r.element).strip().isdigit() and r.box.top > 0.4 and r.box.bottom < 0.7)
    assert numbers == ["1", "2", "3", "4"], "кружки-номера заполненных шагов на оси"


def test_title_plate_stops_before_the_background_logos(lct_timeline):
    slide, _pattern, canvas = lct_timeline
    head = next(r for r in slide_refs(slide, canvas) if "92%" in shape_text(r.element))
    plate = min(
        (r for r in slide_refs(slide, canvas)
         if r.kind == "shape" and not shape_text(r.element).strip() and r.box is not None
         and r.box.left <= head.box.left and r.box.top <= head.box.top + 0.005 and r.box.bottom >= head.box.bottom - 0.005),
        key=lambda r: r.box.area,
    )
    marks = builder._background_marks_left(slide, plate.box.top, plate.box.bottom, head.box.left + 0.2)
    assert marks is not None, "логотипы нарисованы в картинке фона"
    assert plate.box.right < marks
