"""Задача T3: показатели, таблица, цитата и команда на каждом шаблоне
датасета. После разбора видов (задача T1) у части шаблонов есть родная
раскладка формы (показатели WorkSpace, цитата и команда VK Education), и
планировщик берёт её; где нет, берёт носитель, и форму строит сборка. Тесты
проверяют итог, а не путь: форма на слайде, крупные числа в ряд, таблица
нарисована, у команды имена, без выхода за холст (L01), наложений (L02) и
захода в поля (L06). Профиль разбирается без модели (`from_file(
cache_dir=None)`), раскладку выбирает тот же расчёт, что в пайплайне
(`candidates_for`)."""
from __future__ import annotations
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.audit.deterministic import run_deterministic
from deckforge.compose import builder
from deckforge.compose.quote import quote_marks
from deckforge.compose.team import initials, member_photos
from deckforge.pattern.candidates import candidates_for, cover_pattern_id, is_closing_pattern
from deckforge.pattern.forms import forms_of
from deckforge.pattern.intent import SlideIntent
from deckforge.plan.normalize import normalize_deck
from deckforge.plan.spec import (
    DeckSpec, Kpi, KpiBlock, QuoteBlock, SlideSpec, TableVisual, TeamBlock, TeamMember, Visual,
    deck_spec_from_debug_dict, deck_spec_to_dict, slide_spec_from_dict,
)
from deckforge.plan.variants import Variant

TEMPLATES = sorted(Path("dataset/templates").glob("*.pptx"))
PHOTOS_DIR = Path("fixtures/content-packs/full-coverage/photos")
PHOTOS = {p.name: p for p in PHOTOS_DIR.glob("*.jpg")}
LAYOUT_ERRORS = {"L01", "L02", "L06"}

_KPI = KpiBlock(items=[
    Kpi("31,5 ч → 6,2 ч", "сквозная медиана согласования"), Kpi("23% → 6%", "заявок с просрочкой SLA"),
    Kpi("4,1", "оценка удобства авторами"),
])
_TABLE = Visual(kind="table", table=TableVisual(rows=[
    ["Критерий", "Ручная", "Автоматическая"],
    ["Кто назначает согласующего", "автор заявки", "правило по типу и сумме"],
    ["Подмена отсутствующего", "вручную, по звонку", "из кадровой системы"],
    ["Время назначения", "до 4 часов", "2 секунды"],
    ["Ошибки назначения", "39% заявок", "4% заявок"],
]))
_QUOTE = QuoteBlock(
    "Впервые заявка проходит быстрее, чем я успеваю о ней вспомнить. Автоматическая маршрутизация сняла "
    "главный источник жалоб авторов.", "Сергей Лебедев, директор по операциям",
)
_TEAM = TeamBlock(items=[
    TeamMember("Анна Смирнова", "руководитель продукта"), TeamMember("Илья Петров", "архитектор"),
    TeamMember("Мария Ким", "аналитик"), TeamMember("Денис Орлов", "разработчик портала"),
])
_FORMS = {
    "kpi": ([_KPI], None, "Пилот сократил медиану в пять раз"),
    "table": ([], _TABLE, "Правило назначает согласующего за секунды"),
    "quote": ([_QUOTE], None, "Заказчик увидел разницу в первый месяц"),
    "team": ([_TEAM], Visual(kind="photo", photo_name="office-review.jpg"), "Пилот вели четыре человека"),
}


def _slide(profile, form: str, index: int = 1) -> SlideSpec:
    forms = forms_of(profile)
    closing = frozenset(p.pattern_id for p in profile.patterns if is_closing_pattern(p, profile))
    intent = SlideIntent(index=index, outline_kind="data", intent="x", items=3, form=form)
    found = candidates_for(intent, profile, forms, position=index, last=index + 1,
                           cover_id=cover_pattern_id(profile), closing_ids=closing)
    kinds = {p.pattern_id: p.kind for p in profile.patterns}
    blocks, visual, headline = _FORMS[form]
    pid = found.pattern_ids[0]
    return SlideSpec(
        index=index, kind=kinds[pid], headline=headline, blocks=list(blocks), visual=visual,
        source_note="Замер пилота, январь–июнь 2026", pattern_id=pid, alternatives=tuple(found.pattern_ids[1:4]),
        meta={"scratch_form": found.scratch_form} if found.scratch_form else {},
    )


@pytest.fixture(scope="module", params=TEMPLATES, ids=lambda p: p.stem[:14])
def built(request, profile_fixture):
    template = request.param
    profile = profile_fixture(template.name)
    spec = DeckSpec(title="Формы с нуля", language="ru", slides=[
        _slide(profile, form, i) for i, form in enumerate(_FORMS, start=1)
    ])
    out = builder.build_deck(spec, profile, template, Variant.dense, user_photos=PHOTOS)
    findings = run_deterministic(out, profile, AuditConfig.load())
    return spec, Presentation(str(out)), findings


def test_every_form_lands_and_plain_scratch_is_not_used(built):
    """Клон родной раскладки мог не пройти, но тогда форма строится с нуля
    (`ladder_path` кончается на `scratch_form`), а не обычной сборкой с
    нуля, которая раскладывала показатели мелким кеглем и теряла таблицу."""
    spec, _prs, _findings = built
    for slide in spec.slides:
        path = slide.meta.get("ladder_path", "")
        assert slide.meta.get("ladder_rung") != "scratch" or path.endswith("scratch_form"), (slide.meta, slide.findings)
        assert not any("не отрисован" in f for f in slide.findings), slide.findings


def test_no_overflow_overlap_or_margin_breach(built):
    _spec, _prs, findings = built
    bad = [f"{f.check_id}: {f.message}" for f in findings if f.check_id in LAYOUT_ERRORS]
    assert not bad, bad


def _texts(slide) -> list[str]:
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


def _paragraphs(slide):
    """(фигура, текст абзаца, кегль первого run) по всем абзацам слайда:
    клон кладёт значение и подпись показателя в одну фигуру двумя
    абзацами, сборка с нуля в две фигуры."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
            yield shape, p.text, sizes[0] if sizes else None


def test_kpi_values_are_large_and_on_one_row(built):
    _spec, prs, _findings = built
    slide = prs.slides[0]
    found = list(_paragraphs(slide))
    values = [(sh, pt) for sh, text, pt in found if text in {k.value for k in _KPI.items}]
    labels = [(sh, pt) for sh, text, pt in found if text in {k.label for k in _KPI.items}]
    assert len(values) >= 2 and len(values) == len(labels), [t for _s, t, _p in found]
    assert len({round(sh.top / 10000) for sh, _pt in values}) == 1
    assert all(v_pt is not None and l_pt is not None for (_s, v_pt), (_t, l_pt) in zip(values, labels))
    assert min(pt for _s, pt in values) >= 1.5 * max(pt for _s, pt in labels)


def test_table_is_drawn_in_the_free_place(built):
    _spec, prs, _findings = built
    tables = [s for s in prs.slides[1].shapes if s.has_table]
    assert len(tables) == 1
    assert len(tables[0].table.rows) == 5


def test_quote_is_italic_in_language_quotes_with_author(built):
    _spec, prs, _findings = built
    texts = _texts(prs.slides[2])
    quote = next(t for t in texts if "Впервые заявка" in t)
    assert quote.startswith("«") and quote.endswith("»")
    assert any(t.startswith("— Сергей Лебедев") for t in texts)
    shape = next(s for s in prs.slides[2].shapes if s.has_text_frame and s.text_frame.text == quote)
    assert shape.text_frame.paragraphs[0].runs[0].font.italic


def test_team_cards_have_names_and_photo_or_initials(built):
    _spec, prs, _findings = built
    slide = prs.slides[3]
    texts = " \n".join(_texts(slide))
    for member in _TEAM.items:
        assert member.name in texts
    # Фото слайда уходит первому участнику по порядку, остальным кружки
    # с инициалами (или фото родного повтора шаблона, если он есть).
    pictures = [s for s in slide.shapes if s.shape_type == 13]
    circles = {"ИП", "МК", "ДО"} <= set(_texts(slide))
    assert pictures and (circles or len(pictures) >= len(_TEAM.items))


def test_member_photo_found_by_caption(tmp_path):
    photo = tmp_path / "portrait.jpg"
    photo.write_bytes((PHOTOS_DIR / "office-review.jpg").read_bytes())
    (tmp_path / "portrait.txt").write_text("Илья Петров на демонстрации портала", encoding="utf-8")
    other = tmp_path / "other.jpg"
    other.write_bytes(photo.read_bytes())
    members = [TeamMember("Анна Смирнова"), TeamMember("Илья Петров")]
    found = member_photos(members, {"portrait.jpg": photo, "other.jpg": other}, ["other.jpg"])
    assert found == [other, photo]
    assert initials("Анна Смирнова") == "АС"


def test_quote_marks_follow_the_language():
    assert quote_marks("Текст", "ru") == ("«", "»")
    assert quote_marks("Plain text", "en") == ("“", "”")


def test_team_block_round_trips_and_normalize_keeps_scratch_forms(profile_fixture):
    slide = slide_spec_from_dict({
        "kind": "bullets", "headline": "Команда пилота", "blocks": [
            {"type": "team", "items": [{"name": "Анна Смирнова", "role": "продукт"}, {"name": "Илья Петров"}]},
        ],
    }, 1)
    assert slide.blocks[0] == TeamBlock(items=[TeamMember("Анна Смирнова", "продукт"), TeamMember("Илья Петров", "")])
    deck = DeckSpec(title="T", language="ru", slides=[slide])
    assert deck_spec_from_debug_dict(deck_spec_to_dict(deck)).slides[0].blocks == slide.blocks
    # Два пункта с числами на шаблоне с показателями нормализация сделала бы
    # показателями; слайд с формой с нуля она не трогает.
    kpi_slide = SlideSpec(index=0, kind="bullets", headline="Показатели пилота", blocks=[_KPI],
                          pattern_id="x", meta={"scratch_form": "kpi"})
    profile = profile_fixture(TEMPLATES[0].name)
    out = normalize_deck(DeckSpec(title="T", language="ru", slides=[kpi_slide]), profile)
    assert out.slides[0] is kpi_slide
