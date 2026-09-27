"""Задача T2: фрагмент кода на слайде моногарнитурой, строка в строку, без
переносов. С нуля своей плашкой, у клона в фигуру слота роли `code`."""
from __future__ import annotations
import functools
from pathlib import Path

from pptx import Presentation

from deckforge.compose.builder import Variant, build_deck
from deckforge.compose.code import CODE_SHAPE_NAME, MIN_CODE_PT, add_code_block, code_size, fill_code_shape
from deckforge.ooxml.geometry import Box
from deckforge.ooxml.ns import qn
from deckforge.plan.spec import CodeBlock, DeckSpec, SlideSpec, Visual
from deckforge.template.profile import TemplateProfile

VKE = Path("dataset/templates/Шаблон презентации VK Education.pptx")
ROUTE = CodeBlock(language="python", code=(
    "def route(request: Request) -> list[Approver]:\n"
    "    if request.kind == \"purchase\" and request.amount > 5_000_000:\n"
    "        return [head_of(request.unit), cfo()]\n"
    "    if request.kind == \"purchase\":\n"
    "        return [head_of(request.unit)]\n"
    "    return [line_manager(request.author)]"
))


@functools.lru_cache(maxsize=None)
def _profile() -> TemplateProfile:
    return TemplateProfile.from_file(VKE, cache_dir=None)


def _code_shapes(slide):
    return [sh for sh in slide.shapes if sh.name == CODE_SHAPE_NAME]


def _assert_is_code(shape, code: CodeBlock, family: str) -> None:
    tf = shape.text_frame
    assert [p.text for p in tf.paragraphs] == code.code.splitlines(), "строка в строку"
    assert tf._txBody.find(qn("a:bodyPr")).get("wrap") == "none", "переносов нет"
    runs = [r for p in tf.paragraphs for r in p.runs if r.text]
    assert runs and {r.font.name for r in runs} == {family}
    assert all(p._p.find(qn("a:pPr")).find(qn("a:buNone")) is not None for p in tf.paragraphs)
    assert all(r.font.size.pt >= MIN_CODE_PT for r in runs)


def test_font_size_goes_down_the_scale_but_not_below_ten_points():
    lines = ROUTE.code.splitlines()
    roomy = code_size(lines, 10.0, 4.0, 16.0)
    assert roomy.fits and roomy.size_pt == 16.0
    tight = code_size(lines, 6.0, 4.0, 16.0)
    assert tight.fits and MIN_CODE_PT <= tight.size_pt < 16.0
    hopeless = code_size(lines, 2.0, 4.0, 16.0)
    assert not hopeless.fits and hopeless.size_pt == MIN_CODE_PT


def test_scratch_code_block_uses_the_template_mono_font_on_a_palette_plate(new_slide):
    profile = _profile()
    assert profile.type_scale.mono == ["Consolas"]
    slide = new_slide()
    notes = add_code_block(slide, Box(0.05, 0.25, 0.9, 0.6), ROUTE, profile, 1.0, caption="Правило закупок")
    assert notes == []
    [code] = _code_shapes(slide)
    _assert_is_code(code, ROUTE, "Consolas")
    names = {sh.name for sh in slide.shapes}
    assert f"{CODE_SHAPE_NAME} plate" in names and f"{CODE_SHAPE_NAME} language" in names
    plate = next(sh for sh in slide.shapes if sh.name == f"{CODE_SHAPE_NAME} plate")
    allowed = {c.upper().lstrip("#") for c in [*profile.palette_roles.values(), *profile.theme.scheme.values()] if c}
    assert str(plate.fill.fore_color.rgb) in allowed
    assert plate.top + plate.height <= code.top + code.height + 914400, "плашка по коду, а не во всю рамку"


def test_code_goes_into_the_designer_shape_of_a_code_slot():
    """Слот роли `code` (задача T1): текст примера заменён кодом, фигура,
    её плашка и цвет остаются."""
    prs = Presentation(str(VKE))
    shape = next(sh for sh in prs.slides[37].shapes if sh.has_text_frame and len(sh.text_frame.text) > 40)
    notes = fill_code_shape(shape._element, ROUTE, _profile(), 7.0, 3.0, 16.0)
    assert notes == []
    assert shape.name == CODE_SHAPE_NAME
    _assert_is_code(shape, ROUTE, "Consolas")


SHORT = CodeBlock(language="python", code=(
    "def route(request):\n"
    "    if request.amount > 5_000_000:\n"
    "        return [head(request), cfo()]\n"
    "    return [head(request)]"
))


def _code_slide(pattern_id: str, profile=None, code: CodeBlock = ROUTE) -> SlideSpec:
    kind = next(p.kind for p in (profile or _profile()).patterns if p.pattern_id == pattern_id)
    return SlideSpec(
        index=0, kind=kind, headline="Маршрут задаёт правило", pattern_id=pattern_id,
        visual=Visual(kind="code", caption="Крупные закупки идут финансовому директору", code=code),
    )


def test_build_deck_puts_the_code_on_the_text_area_of_the_layout():
    spec = DeckSpec(title="t", language="ru", slides=[_code_slide("slide38")])
    prs = Presentation(str(build_deck(spec, _profile(), VKE, Variant.dense)))
    [code] = _code_shapes(prs.slides[0])
    _assert_is_code(code, ROUTE, "Consolas")
    assert not any("не лёг" in f for f in spec.slides[0].findings)


def test_build_deck_clones_a_code_slot_with_the_designer_shape():
    """Синтетический паттерн вида `code`: у slide38 главное текстовое место
    объявлено слотом `code`, как его снимет разбор задачи T1. Рамка там в
    полслайда, поэтому фрагмент короткий: длинный клон отклонил бы по L03
    и собрал слайд с нуля своей плашкой."""
    profile = _profile()
    patterns = []
    for p in profile.patterns:
        if p.pattern_id == "slide38":
            body = max((s for s in p.slots if s.role in ("body", "bullet")), key=lambda s: s.max_chars)
            p = p.model_copy(update={
                "kind": "code",
                "slots": [s.model_copy(update={"role": "code"}) if s is body else s for s in p.slots],
            })
        patterns.append(p)
    coded = profile.model_copy(update={"patterns": patterns})
    spec = DeckSpec(title="t", language="ru", slides=[_code_slide("slide38", coded, SHORT)])
    prs = Presentation(str(build_deck(spec, coded, VKE, Variant.dense)))
    slide = prs.slides[0]
    [code] = _code_shapes(slide)
    _assert_is_code(code, SHORT, "Consolas")
    assert not any(sh.name == f"{CODE_SHAPE_NAME} plate" for sh in slide.shapes), "плашка дизайнера, не наша"
    assert any("клоном" in f for f in spec.slides[0].findings)
