"""Тесты сборки слайда на лейауте шаблона (Task 9, Step 3 брифа).

Фикстуры `SAMPLE_SPEC`/`CARDS_SPEC`/`PROFILE`/`TEMPLATE` бриф не задаёт
дословно (только сигнатуры тестовых функций и их поведение) — собраны здесь
по одному из учебных шаблонов (VK Tech, 20 паттернов: `bullets`×11,
`cards`×4, `two_col`×3, `section`×2 — золотой прогон `test_golden_profiles.
py`) реальным содержанием пакета `fixtures/content-packs/queue-latency`
(цифры и формулировки — из его `sources.md`)."""
from __future__ import annotations
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckforge.audit.config import AuditConfig
from deckforge.compose.builder import (
    Variant, _avoid_decor_overlap, _best_contrast_color, _clear_sample_slides, _is_cosmetic_truncation,
    _local_background_luminance, _overlap_ratio, _pick_pattern, _place_best_candidate, _relative_luminance,
    _SelectionHistory, build_deck, count_embedded_photos, fits,
)
from deckforge.compose.textfit import measure
from deckforge.ooxml.geometry import Box, Canvas
from deckforge.ooxml.package import PptxPackage
from deckforge.plan.spec import BulletBlock, Card, CardBlock, DeckSpec, SlideSpec, TextBlock
from deckforge.template.grid import Grid
from deckforge.template.naming import MIN_CONTRAST, contrast_ratio
from deckforge.template.patterns import Capacity, DecorShape, Pattern, PatternSlot, RepeatSpec
from deckforge.template.profile import TemplateProfile

TEMPLATE = Path("dataset/templates/VK Tech шаблон.pptx")
PROFILE = TemplateProfile.from_file(TEMPLATE)

SAMPLE_SPEC = DeckSpec(
    title="Сокращение времени согласования заявок",
    language="ru",
    slides=[
        SlideSpec(
            index=0, kind="section",
            headline="Автоматическая маршрутизация заявок",
        ),
        SlideSpec(
            index=1, kind="bullets",
            headline="Где уходит время",
            blocks=[
                BulletBlock(items=[
                    "Ожидание первого согласующего — медиана 18 часов",
                    "Ожидание второго согласующего — медиана 11 часов",
                    "Чистая работа людей — 28 минут из 31,5 часа сквозного цикла",
                ]),
            ],
            source_note="Источник: замер 1 240 заявок, март—май 2026",
        ),
        SlideSpec(
            index=2, kind="two_col",
            headline="Пилот дал результат",
            blocks=[
                TextBlock(text="Сквозная медиана сократилась с 31,5 до 6,2 часа."),
                TextBlock(text="Доля переназначений вручную упала с 39% до 4%."),
            ],
        ),
    ],
)

CARDS_SPEC = DeckSpec(
    title="Сокращение времени согласования заявок",
    language="ru",
    slides=[
        SlideSpec(index=0, kind="section", headline="Что нужно для раскатки"),
        SlideSpec(
            index=1, kind="bullets", headline="Риски",
            blocks=[BulletBlock(items=["Задержка данных кадровой системы", "4 подразделения со своими регламентами"])],
        ),
        SlideSpec(
            index=2, kind="cards", headline="Что нужно для раскатки",
            blocks=[
                CardBlock(items=[
                    Card(title="Правила закупок", body="Доработка правил для закупок свыше 5 млн ₽ — 3 недели"),
                    Card(title="Кадровая система", body="Интеграция для автоподмены отсутствующих — 2 недели"),
                    Card(title="Обучение", body="340 согласующих, 4 вебинара по 40 минут"),
                    Card(title="Бюджет", body="2,4 млн ₽, из них 1,9 млн ₽ ФОТ"),
                ]),
            ],
        ),
    ],
)


def _runs(shape):
    if not shape.has_text_frame:
        return []
    return [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]


def _family(shape) -> str | None:
    runs = _runs(shape)
    return runs[0].font.name if runs else None


def _size(shape) -> float | None:
    runs = _runs(shape)
    return runs[0].font.size.pt if runs and runs[0].font.size is not None else None


def _is_logo(shape, profile: TemplateProfile) -> bool:
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or profile.assets.logo is None:
        return False
    try:
        with PptxPackage.open(TEMPLATE) as pkg:
            logo_bytes = pkg.part(profile.assets.logo.part_name)
    except Exception:
        return False
    return shape.image.blob == logo_bytes


def test_slide_is_built_on_a_layout_from_the_template():
    """Проверка аудита T04: слайд обязан сидеть на макете шаблона, а не на пустом."""
    out = build_deck(SAMPLE_SPEC, PROFILE, TEMPLATE, Variant.dense)
    prs = Presentation(str(out))
    template_layout_names = {e.name for e in PROFILE.layouts}
    assert len(prs.slides) == len(SAMPLE_SPEC.slides)
    for slide in prs.slides:
        assert slide.slide_layout.name in template_layout_names


def test_no_empty_layout_placeholder_survives_into_the_built_slide():
    """Task 13, дефект отчёта задачи №1 (критично, общий): `prs.slides.
    add_slide(layout)` копирует на слайд ВСЕ плейсхолдеры макета
    (python-pptx, `SlideShapes.clone_layout_placeholders`) — а наш код
    никогда не пишет текст В САМ плейсхолдер: `_draw_slot` всегда рисует
    отдельный `add_textbox` поверх координат слота, `_place_picture_visual`/
    `add_chart`/`add_table` — отдельные фигуры (см. докстроки `builder.py`).
    В LibreOffice пустой плейсхолдер рисуется пустотой, в PowerPoint —
    видимой надписью "Щелкните, чтобы добавить текст" на каждом слайде
    богатого макета (живой разбор ЛЦТ2026, слайд 4 макета "Содержание_1":
    19 пустых плейсхолдеров, координатор). Даже простой Title-плейсхолдер
    VK Tech (единственный, который несут его лейауты) остаётся пустым тем
    же путём — заголовок уходит в отдельный textbox, не в него.

    Автозаполняемые PowerPoint'ом плейсхолдеры (номер слайда/дата/
    колонтитул/шапка) не проверяются — python-pptx их и так не клонирует
    (`clone_layout_placeholders`: "Latent placeholders... are not cloned"),
    но даже если бы клонировал, код обязан их не трогать (см. `builder.
    _AUTO_FILLED_PLACEHOLDER_TYPES`)."""
    out = build_deck(SAMPLE_SPEC, PROFILE, TEMPLATE, Variant.dense)
    prs = Presentation(str(out))
    for i, slide in enumerate(prs.slides):
        for shape in slide.placeholders:
            text = shape.text_frame.text if shape.has_text_frame else ""
            assert text.strip(), (
                f"слайд {i}: пустой плейсхолдер {shape.name!r} "
                f"(idx={shape.placeholder_format.idx}, type={shape.placeholder_format.type}) "
                "остался в готовом слайде"
            )


def test_no_slide_is_a_single_raster_image():
    """ТЗ: слайд, выгруженный единым растровым изображением, не засчитывается."""
    prs = Presentation(str(build_deck(SAMPLE_SPEC, PROFILE, TEMPLATE, Variant.dense)))
    for slide in prs.slides:
        others = [s for s in slide.shapes if s.shape_type != MSO_SHAPE_TYPE.PICTURE]
        assert others, "слайд состоит из одних картинок"


def test_text_never_leaves_its_box():
    """Правка по итогам визуального ревью (отчёт задачи, находка №1): текст
    может видимо вылезти за рамку слота ТОЛЬКО в паре с честной находкой —
    когда усечение было бы разрушительным (пусто/до «…»), сборка
    сознательно оставляет текст целиком и вылезание ловит finding, а не
    тихо. Раньше (буквальный брифовый вариант этого теста) любое
    переполнение маскировалось усечением до «…», даже разрушительным —
    сам тест этого не видел, потому что усечённый текст ВСЕГДА укладывался
    в рамку по построению; теперь у "вылез за рамку" ровно два честных
    исхода: либо такого не произошло, либо произошло и об этом есть
    finding — тихого разрыва между "визуально ок" и "метрика говорит не
    влезает" быть не должно."""
    prs = Presentation(str(build_deck(SAMPLE_SPEC, PROFILE, TEMPLATE, Variant.dense)))
    assert len(prs.slides) == len(SAMPLE_SPEC.slides)
    for slide, slide_spec in zip(prs.slides, SAMPLE_SPEC.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            family = _family(shape) or "Arial"
            size = _size(shape) or 12.0
            metrics = measure(shape.text_frame.text, family, size, shape.width / 914400)
            if metrics.height_in > shape.height / 914400 + 0.05:
                assert slide_spec.findings, (
                    f"слайд {slide_spec.index}: текст вылез за рамку слота без единой находки"
                )


def test_only_template_colors_and_fonts_are_used():
    allowed_colors = set(PROFILE.palette_roles.values())
    allowed_fonts = set(PROFILE.type_scale.families) | set(PROFILE.type_scale.mono)
    for slide in Presentation(str(build_deck(SAMPLE_SPEC, PROFILE, TEMPLATE, Variant.dense))).slides:
        for shape in slide.shapes:
            for run in _runs(shape):
                assert run.font.name in allowed_fonts
                if run.font.color and run.font.color.type is not None:
                    assert f"#{run.font.color.rgb}" in allowed_colors


def test_decor_of_the_pattern_is_carried_over():
    """Плашки и линии паттерна переносятся — иначе слайд теряет язык шаблона."""
    prs = Presentation(str(build_deck(CARDS_SPEC, PROFILE, TEMPLATE, Variant.visual)))
    cards_slide = prs.slides[2]
    assert len([s for s in cards_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]) >= 3


def test_template_master_shapes_are_not_duplicated():
    prs = Presentation(str(build_deck(SAMPLE_SPEC, PROFILE, TEMPLATE, Variant.dense)))
    for slide in prs.slides:
        logos = [s for s in slide.shapes if _is_logo(s, PROFILE)]
        assert len(logos) <= 1


def test_cards_expand_to_the_actual_number_of_items():
    """Бриф: "повтор разворачивай под фактическое число элементов" — четыре
    карточки контент-пакета, а не шесть/сколько было на слайде-примере.

    Проверяем по ТЕЛУ карточки, не по заголовку: на VK Tech лучшая по
    вместимости "cards"-раскладка для абзацев такой длины (см. отчёт
    задачи, находка №1 и её правку) физически не несёт отдельного слота
    `card_title` — подбор паттерна теперь честно предпочитает раскладку, где
    ТЕЛО помещается целиком, а не ту, где заголовок на месте, а тело
    переполнено на 251% (было бы хуже: либо задвоенный текст с усечением до
    «…», либо тело, вываливающееся за карточку в 2.5 раза). Тело карточки —
    обязательная часть контента (`CardBlock.items[i].body` не бывает
    пустым), заголовок — необязательная (`Card.title` по умолчанию "")."""
    prs = Presentation(str(build_deck(CARDS_SPEC, PROFILE, TEMPLATE, Variant.visual)))
    cards_slide = prs.slides[2]
    texts = [s.text_frame.text for s in cards_slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert any("340 согласующих" in t for t in texts), "тело третьей карточки (Обучение) не найдено"
    assert any("1,9 млн ₽ ФОТ" in t for t in texts), "тело четвёртой карточки (Бюджет) не найдено"


def test_expanded_repeat_stays_within_the_canvas():
    """Находка ручной проверки на VK Education: паттерн "cards" там несёт
    составную единицу повтора (узкий `kpi_value` + широкий `card_body` на
    одной x-координате, друг над другом) — размер единицы обязан браться по
    САМОМУ ШИРОКОМУ слоту группы, иначе шаг пересчитывается с завышенным
    "свободным" пространством и последняя карточка уезжает за правое поле
    холста (см. `expand_repeat`)."""
    template = Path("dataset/templates/Шаблон презентации VK Education.pptx")
    profile = TemplateProfile.from_file(template)
    spec = DeckSpec(
        title="Т", language="ru",
        slides=[SlideSpec(
            index=0, kind="cards", headline="Что нужно для раскатки",
            blocks=[CardBlock(items=[
                Card(body="Доработка правил — 3 недели"),
                Card(body="Интеграция кадровой системы — 2 недели"),
                Card(body="Обучение 340 согласующих"),
                Card(body="Бюджет 2,4 млн ₽"),
            ])],
        )],
    )
    out = build_deck(spec, profile, template, Variant.dense)
    prs = Presentation(str(out))
    canvas_w_in = prs.slide_width / 914400
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left / 914400 >= -0.01
            assert (shape.left + shape.width) / 914400 <= canvas_w_in + 0.01


# ---------------------------------------------------------------------------
# Правка по итогам визуального ревью (см. .superpowers/sdd/task-9-report.md,
# раздел, дописанный после первого прохода): пустая/усечённая-до-«…» карточка
# хуже переполненной, подбор раскладки обязан использовать проверку
# пригодности (а не брать первый попавшийся паттерн вида kind), и учитывать
# «не слишком ли пусто», не только «влезает ли».
# ---------------------------------------------------------------------------


def _grid_stub() -> Grid:
    return Grid(
        margin_left=0.05, margin_right=0.05, margin_top=0.1, margin_bottom=0.1,
        columns=[], gutter=0.02, anchors={}, confidence={}, skipped_no_box=0, native_guides_used=False,
    )


def _fake_cards_pattern(
    pattern_id: str, card_box: Box, score: float = 1.0, extra_slots=(), capacity_max_items: int = 6,
) -> Pattern:
    body_slot = PatternSlot(
        role="card_body", box=card_box, size_pt=16.0 * 1.333, color_hex=None,
        align="l", max_chars=999, wraps=True,
    )
    return Pattern(
        pattern_id=pattern_id, source_slide_index=[0], layout_id="slideLayout1", kind="cards",
        slots=[body_slot, *extra_slots], repeat=RepeatSpec(axis="x", count=1, step=0.0, slot_roles=["card_body"]),
        decor=[],
        capacity=Capacity(
            max_items=capacity_max_items, max_chars_per_item=999, max_bullets=0, max_series=0, max_rows=0, max_cols=0,
        ),
        score=score, is_dark=False,
    )


def test_fits_accounts_for_shrinking_not_just_the_native_size():
    """Наивная проверка "влезает на СВОЁМ кегле" отбраковала бы паттерн,
    хотя ужатый по шкале `TypeScale` (до caption) текст реально помещается —
    `_draw_slot` потом кладёт его без всякого усечения. `fits()` обязана
    пробовать всю шкалу ужимания, иначе подборщик отбраковывает раскладки,
    которые сборка успешно уложит."""
    huge_native_size = _fake_cards_pattern("huge-size", Box(0.1, 0.1, 0.3, 0.12))
    # размер слота (16pt * норм.коэфф.) заведомо велик для узкой рамки высотой
    # 0.12 холста — на СВОЁМ кегле короткая строка не влезет, но на caption влезет.
    slide_spec = SlideSpec(
        index=0, kind="cards", headline="",
        blocks=[CardBlock(items=[Card(body="Короткая подпись карточки")])],
    )
    fit = fits(slide_spec, huge_native_size, PROFILE)
    assert fit.ok, fit.reason


def test_pattern_picker_prefers_a_layout_the_content_fits_in():
    """Бриф: подбор обязан перебрать кандидатов вида kind и взять тот, в
    который содержание влезает, а не первый попавшийся."""
    narrow = _fake_cards_pattern("narrow", Box(0.1, 0.3, 0.1, 0.05))
    wide = _fake_cards_pattern("wide", Box(0.1, 0.3, 0.4, 0.4))
    slide_spec = SlideSpec(
        index=0, kind="cards", headline="",
        blocks=[CardBlock(items=[Card(
            body="Доработка правил для закупок свыше пяти миллионов рублей потребует три недели работы команды",
        )])],
    )
    picked = _pick_pattern(slide_spec, [narrow, wide], PROFILE, Variant.dense)
    assert picked.pattern_id == "wide"


def test_pattern_picker_avoids_a_layout_that_is_mostly_empty_for_this_content():
    """ТЗ D05 ("слайд заполнен меньше четверти — брак"): раскладка на
    много карточек под мало короткого содержания технически "влезает", но
    оставляет холст почти пустым — подбор обязан предпочесть раскладку,
    где то же содержание занимает разумную долю холста."""
    sparse = _fake_cards_pattern("sparse", Box(0.1, 0.5, 0.06, 0.04))
    dense = _fake_cards_pattern("dense", Box(0.1, 0.25, 0.35, 0.45))
    slide_spec = SlideSpec(
        index=0, kind="cards", headline="",
        blocks=[CardBlock(items=[Card(body="Коротко")])],
    )
    picked = _pick_pattern(slide_spec, [sparse, dense], PROFILE, Variant.dense)
    assert picked.pattern_id == "dense"


def test_pattern_picker_prefers_the_layout_whose_capacity_matches_the_content_volume():
    """Task 9 повторное ревью, находка №1 ("главная находка"): раскладка на
    шесть карточек и раскладка на две карточки формально ОБЕ влезают под
    два элемента короткого содержания и обе не выглядят пустыми по
    `fill_ratio` (короткий текст занимает ту же долю ОДИНАКОВОГО по
    размеру слота в обоих кандидатах, см. одинаковый `card_box`) — подбор
    обязан предпочесть ту, чья `Capacity.max_items` ближе к фактическому
    числу элементов, а не первую формально подходящую (иначе раскладка,
    рассчитанная на шесть элементов, под два — плохой выбор: четыре
    пустые декоративные рамки, ровно то, что увидел постановщик на VK
    Tech, "Риски раскатки")."""
    box = Box(0.1, 0.3, 0.3, 0.3)
    roomy = _fake_cards_pattern("six-slot", box, capacity_max_items=6)
    snug = _fake_cards_pattern("two-slot", box, capacity_max_items=2)
    slide_spec = SlideSpec(
        index=0, kind="cards", headline="",
        blocks=[CardBlock(items=[Card(body="Коротко"), Card(body="И ещё короче")])],
    )
    picked = _pick_pattern(slide_spec, [roomy, snug], PROFILE, Variant.dense)
    assert picked.pattern_id == "two-slot"


# ---------------------------------------------------------------------------
# Task 18, находка №3 брифа: штраф за повтор раскладки (`_diversity_
# penalty`/`_SelectionHistory`) — "ранжирование... по двум признаками:
# влезает ли текст и не слишком ли пусто... признака «на прошлом слайде уже
# была такая же» нет, поэтому берётся самая безопасная раз за разом".
# ---------------------------------------------------------------------------


def _twin_cards_slide_spec() -> SlideSpec:
    return SlideSpec(
        index=0, kind="cards", headline="",
        blocks=[CardBlock(items=[Card(body="Коротко"), Card(body="И ещё короче")])],
    )


def test_without_history_the_same_best_candidate_wins_every_time():
    """Контроль: два кандидата, равные по score/fit/capacity/fill,
    выбираются детерминированно по порядку без истории — тот самый эффект
    "самая безопасная раз за разом", который правка ниже обязана сломать
    ТОЛЬКО когда история передана."""
    box = Box(0.1, 0.3, 0.3, 0.3)
    a = _fake_cards_pattern("twin-a", box, capacity_max_items=2)
    b = _fake_cards_pattern("twin-b", box, capacity_max_items=2)
    slide_spec = _twin_cards_slide_spec()
    assert _pick_pattern(slide_spec, [a, b], PROFILE, Variant.dense).pattern_id == "twin-a"
    assert _pick_pattern(slide_spec, [b, a], PROFILE, Variant.dense).pattern_id == "twin-b"


def test_diversity_penalty_avoids_repeating_the_pattern_used_on_the_previous_slide():
    box = Box(0.1, 0.3, 0.3, 0.3)
    a = _fake_cards_pattern("twin-a", box, capacity_max_items=2)
    b = _fake_cards_pattern("twin-b", box, capacity_max_items=2)
    slide_spec = _twin_cards_slide_spec()

    history = _SelectionHistory(last_pattern_id="twin-a")
    picked = _pick_pattern(slide_spec, [a, b], PROFILE, Variant.dense, history)
    assert picked.pattern_id == "twin-b"


def test_diversity_penalty_also_reacts_to_overall_frequency_not_only_the_previous_slide():
    box = Box(0.1, 0.3, 0.3, 0.3)
    a = _fake_cards_pattern("twin-a", box, capacity_max_items=2)
    b = _fake_cards_pattern("twin-b", box, capacity_max_items=2)
    slide_spec = _twin_cards_slide_spec()

    # "twin-a" не была на прошлом слайде (last_pattern_id=None), но уже
    # использована дважды где-то раньше в этой же колоде — штраф частоты
    # (не только немедленного повтора) обязан всё равно сместить выбор.
    history = _SelectionHistory(counts={"twin-a": 2}, last_pattern_id=None)
    picked = _pick_pattern(slide_spec, [a, b], PROFILE, Variant.dense, history)
    assert picked.pattern_id == "twin-b"


def test_diversity_penalty_never_overrides_a_real_capacity_mismatch():
    """Task 18 брифом: штраф "не должен перебивать пригодность" —
    структурно (не весом) гарантировано порядком позиций кортежа
    `_pattern_rank_key`: `_capacity_badness` идёт РАНЬШЕ штрафа за повтор,
    поэтому раскладка, чья вместимость реально хуже подходит содержанию,
    не может победить только за счёт разнообразия, даже если штраф
    максимален (использована уже несколько раз ПОДРЯД)."""
    box = Box(0.1, 0.3, 0.3, 0.3)
    roomy = _fake_cards_pattern("roomy-again", box, capacity_max_items=2)
    mismatched = _fake_cards_pattern("never-used", box, capacity_max_items=20)
    slide_spec = _twin_cards_slide_spec()

    history = _SelectionHistory(counts={"roomy-again": 4}, last_pattern_id="roomy-again")
    picked = _pick_pattern(slide_spec, [roomy, mismatched], PROFILE, Variant.dense, history)
    assert picked.pattern_id == "roomy-again", (
        "штраф за повтор пересилил реальное расхождение вместимости — не должен был"
    )


def test_text_block_goes_to_the_biggest_slot_of_its_role_not_the_first_one():
    """Task 9 повторное ревью, находка №2 ("важное"): на контрольном
    ЛЦТ2026 (слайд "two_col") паттерн несёт ДВА слота роли "body" ВНЕ
    группы повтора — один нормального размера, второй крошечный (место
    под короткую подпись, судя по геометрии, не под абзац). Раньше
    `_slots_by_role`/`_take_one` отдавали содержание ПЕРВОМУ по порядку
    мининга слоту этой роли, независимо от размера — абзац утекал в
    маленький слот у угла холста и резался краем. Содержание обязано лечь
    в тот слот роли, что вмещает (самый ёмкий по площади), а не в первый
    попавшийся — тот же приём, что `_pick_body_slot` уже применяет внутри
    группы повтора карточек."""
    from deckforge.compose.blocks import assign_content

    tiny = PatternSlot(
        role="body", box=Box(0.7, 0.85, 0.1, 0.06), size_pt=12.0, color_hex=None,
        align="l", max_chars=20, wraps=False, sample_text="01",
    )
    roomy = PatternSlot(
        role="body", box=Box(0.1, 0.2, 0.35, 0.5), size_pt=14.0, color_hex=None,
        align="l", max_chars=400, wraps=True, sample_text="Текст колонки",
    )
    pattern = Pattern(
        pattern_id="p", source_slide_index=[0], layout_id="L", kind="two_col",
        # Порядок в списке — намеренно "неудобный" (маленький слот первым):
        # старый код (pop(0) без сортировки) брал бы именно его.
        slots=[tiny, roomy], repeat=None, decor=[],
        capacity=Capacity(max_items=1, max_chars_per_item=400, max_bullets=0, max_series=0, max_rows=0, max_cols=0),
        score=1.0, is_dark=False,
    )
    slide_spec = SlideSpec(
        index=0, kind="two_col", headline="",
        blocks=[TextBlock(text="Кадровая система отдаёт данные с задержкой")],
    )
    result = assign_content(slide_spec, pattern, _grid_stub())
    body_contents = [c for c in result if c.role_hint == "body"]
    assert len(body_contents) == 1
    assert body_contents[0].slot.sample_text == "Текст колонки"


def test_decor_shape_fills_use_only_template_colors():
    """Task 9 повторное ревью, находка №3 ("мелочь"): `test_only_template_
    colors_and_fonts_are_used` выше проверяет только цвет ТЕКСТА — заливки
    декоративных фигур формально не покрыты проверкой аудита "цвет не из
    палитры". Цвет декора приходит из того же снятого `DecorShape.fill_hex`
    (см. `template/patterns.py::_to_decor`, `resolve_color` против той же
    `theme.scheme`), что и остальная палитра.

    "Разрешённые" цвета для декора — не только `palette_roles.values()`
    (найдено проверкой на реальной сборке VK Tech: три декоративные плашки
    несут `#FAFCFF`/`#FEFFFF`/`#D6ECFF` — это буквально `theme.lt1`/
    `theme.lt2`/`theme.accent6`, см. вывод `TemplateProfile.theme.scheme`
    — РЕАЛЬНЫЕ цвета темы шаблона, просто не вошедшие в узкий именованный
    набор `palette_roles` (~8 ролей на 12 цветов схемы темы). `palette_
    roles` — курируемое ПОДМНОЖЕСТВО схемы темы для роли ТЕКСТА
    (`_color_for_role` в builder.py намеренно снаряжает текст только
    именованными ролями); декор переносится "как есть" (докстрока
    `decor.py`) и законно использует ЛЮБОЙ цвет схемы темы, не только
    именованный — это не изобретённый цвет, а буквально то, что было в
    файле шаблона. Проверка аудита "цвет не из палитры" по духу — про
    "не выдуман ли цвет", а не "назван ли он ролью"; для декора это
    `palette_roles.values() | theme.scheme.values()`."""
    allowed_colors = set(PROFILE.palette_roles.values()) | set(PROFILE.theme.scheme.values())
    prs = Presentation(str(build_deck(CARDS_SPEC, PROFILE, TEMPLATE, Variant.visual)))
    checked = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
                continue
            if shape.fill.type is None:
                continue
            try:
                fore = shape.fill.fore_color
            except (TypeError, AttributeError):
                continue
            if fore.type is not None:
                checked += 1
                assert f"#{fore.rgb}" in allowed_colors
    assert checked > 0, "не нашлось ни одной закрашенной декоративной фигуры для проверки"


def test_cosmetic_truncation_keeps_most_of_the_original_text():
    original = "одно два три четыре пять шесть семь восемь девять десять"
    assert _is_cosmetic_truncation(original, "одно два три четыре пять шесть семь восемь…")


def test_destructive_truncation_is_rejected():
    original = "одно два три четыре пять шесть семь восемь девять десять"
    assert not _is_cosmetic_truncation(original, "одно…")
    assert not _is_cosmetic_truncation(original, "…")


def test_overflowing_text_is_never_truncated_to_ellipsis_or_emptied():
    """Ревью глазами (отчёт, находка №1): пустая/усечённая-до-«…» карточка —
    брак хуже переполненной. Если даже на минимальном кегле содержание не
    влезает и усечение вырезало бы больше нужного порога текста — слот
    остаётся с полным текстом (видимое переполнение + finding), а не с
    многоточием или пустотой."""
    huge_card = SlideSpec(
        index=0, kind="cards", headline="Слишком длинная карточка",
        blocks=[CardBlock(items=[Card(
            title="Заголовок",
            body="слово " * 60,
        )])],
    )
    spec = DeckSpec(title="Т", language="ru", slides=[huge_card])
    out = build_deck(spec, PROFILE, TEMPLATE, Variant.dense)
    prs = Presentation(str(out))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if text:
                assert text not in ("…",), f"слот усечён до голого многоточия: {text!r}"
    assert huge_card.findings, "переполнение обязано оставить находку"


def test_picking_the_least_bad_pattern_still_leaves_a_finding():
    """Бриф: "если ни одна раскладка не подходит, выбирай ту, где
    переполнение наименьшее, и оставляй явный след о проблеме" — проверяем
    на уровне выбора паттерна напрямую (без обхода через build_deck):
    единственный кандидат заведомо тесен, `_pick_pattern` обязан его всё
    равно вернуть (лучше плохой выбор, чем никакого — слайд не пропадает),
    а `fits()` на нём — честно сказать "не влезает"."""
    too_small = _fake_cards_pattern("too-small", Box(0.1, 0.4, 0.08, 0.03))
    slide_spec = SlideSpec(
        index=0, kind="cards", headline="",
        blocks=[CardBlock(items=[Card(body="слово " * 30)])],
    )
    picked = _pick_pattern(slide_spec, [too_small], PROFILE, Variant.dense)
    assert picked is not None
    fit = fits(slide_spec, picked, PROFILE)
    assert not fit.ok
    assert fit.overflow_ratio > 0


def test_card_body_is_not_duplicated_into_every_same_role_slot_of_a_group():
    """Корень находки №1 отчёта на VK Tech: паттерн с двумя слотами роли
    `card_body` в ОДНОЙ группе повтора (узкий "под номер" 0.03 высоты холста
    + широкий "под абзац") раньше получал ОДИН И ТОТ ЖЕ текст карточки в
    ОБА слота (`compose/blocks.py::_assign_cards`, старый код: `elif
    slot.role in ("card_body", ...): result.append(...)` внутри `for slot
    in group` — без разбора, сколько таких слотов в группе). Узкий слот
    физически не мог вместить абзац ни на каком кегле шкалы и усекался до
    «…». Текст карточки обязан лечь только в САМЫЙ ЁМКИЙ (по площади) слот
    группы этой роли — остальные (обычно узкие подписи-номера в сэмпле
    мининга) остаются пустыми, а не задвоенными."""
    from deckforge.compose.blocks import assign_content

    small = PatternSlot(
        role="card_body", box=Box(0.1, 0.1, 0.15, 0.03), size_pt=16.0, color_hex=None,
        align="l", max_chars=4, wraps=False, sample_text="01",
    )
    big = PatternSlot(
        role="card_body", box=Box(0.1, 0.15, 0.15, 0.3), size_pt=12.0, color_hex=None,
        align="l", max_chars=200, wraps=True, sample_text="Текст карточки",
    )
    pattern = Pattern(
        pattern_id="p", source_slide_index=[0], layout_id="L", kind="cards",
        slots=[small, big], repeat=RepeatSpec(axis="x", count=1, step=0.0, slot_roles=["card_body"], group_size=2),
        decor=[], capacity=Capacity(max_items=4, max_chars_per_item=200, max_bullets=0, max_series=0, max_rows=0, max_cols=0),
        score=1.0, is_dark=False,
    )
    slide_spec = SlideSpec(
        index=0, kind="cards", headline="",
        blocks=[CardBlock(items=[Card(body="Полный текст карточки про доработку правил закупок")])],
    )
    result = assign_content(slide_spec, pattern, _grid_stub())
    card_body_contents = [c for c in result if c.role_hint == "card_body"]
    assert len(card_body_contents) == 1, "текст карточки задвоен по обоим слотам роли card_body"
    # `expand_repeat` пересобирает слот через `dataclasses.replace` (новая
    # позиция по оси повтора) — сравниваем не identity, а исходный слот
    # (сэмпл-текст/размер площади), который однозначно опознаёт "широкий".
    assert card_body_contents[0].slot.sample_text == "Текст карточки"


def test_text_box_has_no_internal_margins():
    """Находка при повторном визуальном ревью (после установки Play):
    заголовок "Где уходит время" (VK Education) стал переноситься на 3
    строки и наезжать на буллеты НИЖЕ, хотя `fits()`/`_draw_slot` не
    зафиксировали переполнения. Причина — python-pptx даёт текстовым
    рамкам НЕНУЛЕВЫЕ внутренние поля по умолчанию (0.1" слева/справа,
    0.05" сверху/снизу, см. OOXML `a:bodyPr` `lIns`/`rIns`/`tIns`/`bIns`),
    а `measure()` меряет по ПОЛНОЙ ширине/высоте фигуры (`box_width_in`,
    `box_height_in` — без вычета полей) — при непереопределённых полях
    реально доступная под текст ширина в PowerPoint/LibreOffice меньше,
    чем предполагает замер, и перенос строк расходится с прогнозом.
    Поля обязаны быть обнулены, чтобы бокс, который меряет `measure()`, и
    бокс, в который реально рисует текст рендерер, совпадали."""
    prs = Presentation(str(build_deck(SAMPLE_SPEC, PROFILE, TEMPLATE, Variant.dense)))
    checked = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            tf = shape.text_frame
            assert tf.margin_left == 0
            assert tf.margin_right == 0
            assert tf.margin_top == 0
            assert tf.margin_bottom == 0
            checked += 1
    assert checked > 0


def test_text_color_follows_the_local_plaque_not_the_whole_pattern():
    """Находка ручной проверки v2 (ЛЦТ2026, слайд "cards"): паттерн в целом
    тёмный, но карточки — БЕЛЫЕ плашки декора (`DecorShape` с
    `fill_hex="#FFFFFF"`) поверх этого тёмного фона. Фон под слотом обязан
    браться от цвета плашки НЕПОСРЕДСТВЕННО под слотом, если она есть, а не
    от фона всего паттерна/макета."""
    slot_box = Box(0.15, 0.2, 0.15, 0.3)
    white_plaque = DecorShape(
        kind="shape", box=Box(0.1, 0.15, 0.25, 0.4), rotation=0.0, flip_h=False, flip_v=False,
        fill_hex="#FFFFFF", has_fill=True, fill_kind="solid",
    )
    pattern = Pattern(
        pattern_id="p", source_slide_index=[0], layout_id="L", kind="cards",
        slots=[], repeat=None, decor=[white_plaque],
        capacity=Capacity(max_items=4, max_chars_per_item=200, max_bullets=0, max_series=0, max_rows=0, max_cols=0),
        score=1.0, is_dark=True,
    )
    # Тёмный фон макета (layout_bg_luminance=0.0) должен уступить светлой
    # плашке НЕПОСРЕДСТВЕННО под слотом — итог обязан быть светлым.
    assert _local_background_luminance(slot_box, pattern, 0.0) == _relative_luminance("#FFFFFF")

    # Без плашки под слотом (карточка "висит в воздухе" на самом фоне
    # макета) решает фон макета, как и должен.
    pattern_no_plaque = Pattern(
        pattern_id="p2", source_slide_index=[0], layout_id="L", kind="cards",
        slots=[], repeat=None, decor=[],
        capacity=Capacity(max_items=4, max_chars_per_item=200, max_bullets=0, max_series=0, max_rows=0, max_cols=0),
        score=1.0, is_dark=True,
    )
    assert _local_background_luminance(slot_box, pattern_no_plaque, 0.05) == 0.05


def test_layout_background_overrides_the_pattern_own_darkness_flag():
    """Task 10 отчёт, находка №1 — воспроизводит найденную причину дословно:
    контрольный файл, слайд "Риски раскатки (детали)" — раскладка снята со
    СВЕТЛОГО слайда-примера (`pattern.is_dark=False`), но кладётся на
    ТЁМНЫЙ макет шаблона. Источник истины обязан быть фон МАКЕТА
    (`layout_bg_luminance`, из `profile.layouts`), а не `pattern.is_dark`
    примера, с которого раскладка снята — иначе текст остаётся в цвете
    "для светлого фона" (тёмный/чёрный) на фактически тёмном фоне."""
    slot_box = Box(0.1, 0.1, 0.3, 0.1)  # без плашки декора под слотом
    pattern = Pattern(
        pattern_id="p", source_slide_index=[0], layout_id="L", kind="bullets",
        slots=[], repeat=None, decor=[],
        capacity=Capacity(max_items=1, max_chars_per_item=999, max_bullets=0, max_series=0, max_rows=0, max_cols=0),
        score=1.0, is_dark=False,  # раскладка "думает", что она светлая
    )
    dark_purple_luminance = _relative_luminance("#26123F")
    luminance = _local_background_luminance(slot_box, pattern, dark_purple_luminance)
    assert luminance == dark_purple_luminance, "фон макета обязан победить pattern.is_dark примера"

    # Цвет, выбранный по этой яркости, обязан реально контрастировать с
    # фактическим тёмно-фиолетовым фоном — не просто "не быть чёрным".
    chosen = _best_contrast_color("#000000", luminance, PROFILE)
    assert contrast_ratio("#26123F", chosen) >= MIN_CONTRAST


def test_best_contrast_color_is_taken_from_the_template_palette_and_reaches_wcag_aa():
    """Бриф: "не ограничивайся признаком тёмный/светлый — считай контраст
    пары и выбирай из палитры шаблона тот цвет, который даёт контраст не
    ниже 4.5 к фактическому фону". Прямая проверка числа, не косвенного
    признака: фон — знакомый нечитаемый тёмно-фиолетовый из контрольного
    файла, кандидат — чёрный (как было в найденном браке); итог обязан
    прийти ИЗ ПАЛИТРЫ шаблона и реально пройти WCAG AA (4.5:1)."""
    bg_hex = "#2B1147"
    bg_luminance = _relative_luminance(bg_hex)
    chosen = _best_contrast_color("#000000", bg_luminance, PROFILE)
    assert chosen in set(PROFILE.palette_roles.values())
    assert contrast_ratio(bg_hex, chosen) >= MIN_CONTRAST


def test_headline_slot_is_moved_off_a_partially_overlapping_decor_plaque():
    """Task 10 отчёт, находка №2 (воспроизводит контрольный файл дословно):
    заголовок «Риски раскатки (детали)» стоял поверх розовой плашки декора
    в шапке слайда — координаты сняты с реального разбора ЛЦТ2026
    ("PARTIAL headline/decor overlap"). Проверяется НАЛОЖЕНИЕ напрямую
    (площадь пересечения), не косвенный признак."""
    headline_box = Box(0.0431, 0.0666, 0.8090, 0.0548)
    badge = DecorShape(
        kind="shape", box=Box(0.0284, 0.0473, 0.3077, 0.0905), rotation=0.0, flip_h=False, flip_v=False,
        fill_hex="#FF0053", has_fill=True, fill_kind="solid",
    )
    assert _overlap_ratio(headline_box, badge.box) > 0.5, "фикстура обязана воспроизводить реальное наложение"

    grid = _grid_stub()
    moved = _avoid_decor_overlap(headline_box, [badge], grid)
    assert _overlap_ratio(moved, badge.box) <= 0.1, "после починки слот не должен значимо наезжать на декор"
    assert (moved.width, moved.height) == (headline_box.width, headline_box.height), "размер слота не меняется"


def test_slot_fully_nested_in_its_own_plaque_is_not_treated_as_a_collision():
    """Карточная плашка-фон под своим же текстом (слот ПОЛНОСТЬЮ внутри
    декора) — легитимное намеренное расположение, не брак наложения; сдвиг
    здесь не нужен и был бы ошибкой (увёл бы текст с его собственной
    карточки)."""
    slot_box = Box(0.15, 0.2, 0.15, 0.3)
    own_plaque = DecorShape(
        kind="shape", box=Box(0.1, 0.15, 0.25, 0.4), rotation=0.0, flip_h=False, flip_v=False,
        fill_hex="#FFFFFF", has_fill=True, fill_kind="solid",
    )
    grid = _grid_stub()
    assert _avoid_decor_overlap(slot_box, [own_plaque], grid) == slot_box


def test_findings_recorded_when_content_does_not_fit():
    """Бриф: "молча обрезать нельзя, это увидит аудит" — искусственно
    огромный блок, который не влезет ни на одном кегле шкалы вплоть до
    подписи, обязан оставить finding, а не тихо потеряться."""
    huge = SlideSpec(
        index=0, kind="bullets", headline="Заголовок",
        blocks=[BulletBlock(items=["слово " * 400])],
    )
    spec = DeckSpec(title="Т", language="ru", slides=[huge])
    build_deck(spec, PROFILE, TEMPLATE, Variant.dense)
    assert huge.findings, "переполнение слота должно оставить находку"


# ---------------------------------------------------------------------------
# Аудит внутри цикла сборки (Task 13, продолжение брифа): "собрали слайд —
# проверили — не понравилось — взяли другую раскладку и пересобрали".
# `_place_best_candidate` — единственное место, знающее И про кандидатов
# паттерна, И про то, что скажет о них аудит (`compose.builder` уже
# зависит от `audit.deterministic` этим путём, не наоборот).
# ---------------------------------------------------------------------------


def _real_layout_id() -> str:
    return PROFILE.patterns[0].layout_id


def _overlapping_section_pattern(pattern_id: str) -> Pattern:
    """`headline` и `subhead` — ОДИН и тот же бокс: полное наложение (L02)
    гарантировано независимо от длины реального текста (`_check_L02.
    _effective_box` усаживает оба текстовых блока по факту нарисованных
    чернил, но оба делят один left/top/width — пересечение всегда равно
    площади МЕНЬШЕГО эффективного бокса, ratio=1.0)."""
    box = Box(0.05, 0.1, 0.5, 0.1)
    headline = PatternSlot(role="headline", box=box, size_pt=24.0, color_hex=None, align="l", max_chars=200, wraps=True)
    subhead = PatternSlot(role="subhead", box=box, size_pt=16.0, color_hex=None, align="l", max_chars=200, wraps=True)
    return Pattern(
        pattern_id=pattern_id, source_slide_index=[0], layout_id=_real_layout_id(), kind="section",
        slots=[headline, subhead], repeat=None, decor=[],
        capacity=Capacity(max_items=1, max_chars_per_item=999, max_bullets=0, max_series=0, max_rows=0, max_cols=0),
        score=1.0, is_dark=False,
    )


def _clean_section_pattern(pattern_id: str) -> Pattern:
    headline = PatternSlot(
        role="headline", box=Box(0.05, 0.1, 0.5, 0.1), size_pt=24.0, color_hex=None, align="l", max_chars=200, wraps=True,
    )
    subhead = PatternSlot(
        role="subhead", box=Box(0.05, 0.3, 0.5, 0.1), size_pt=16.0, color_hex=None, align="l", max_chars=200, wraps=True,
    )
    return Pattern(
        pattern_id=pattern_id, source_slide_index=[0], layout_id=_real_layout_id(), kind="section",
        slots=[headline, subhead], repeat=None, decor=[],
        capacity=Capacity(max_items=1, max_chars_per_item=999, max_bullets=0, max_series=0, max_rows=0, max_cols=0),
        score=1.0, is_dark=False,
    )


def _new_deck_in_progress():
    prs = Presentation(str(TEMPLATE))
    _clear_sample_slides(prs)
    canvas = Canvas(width_emu=PROFILE.canvas_width_emu, height_emu=PROFILE.canvas_height_emu)
    audit_config = AuditConfig.load()
    return prs, canvas, audit_config


def test_slide_audit_rejects_an_overlapping_layout_and_retries_the_next_candidate():
    """Тест брифа: "слайд, который на первой раскладке даёт наложение, а на
    второй не даёт, собирается на второй". Первый кандидат кладёт заголовок
    буквально поверх подзаголовка (L02) — цикл обязан отклонить его и
    уложить слайд на второй, чистый, кандидат."""
    bad = _overlapping_section_pattern("overlap-bad")
    good = _clean_section_pattern("overlap-ok")
    slide_spec = SlideSpec(index=0, kind="section", headline="Ожидание согласующих", subhead="съедает почти всё время")
    prs, canvas, audit_config = _new_deck_in_progress()

    chosen, notes = _place_best_candidate(prs, slide_spec, [bad, good], PROFILE, canvas, audit_config, bullet_char="•")

    assert chosen.pattern_id == "overlap-ok"
    assert len(prs.slides) == 1, "в колоде обязан остаться РОВНО один (финальный) слайд, не оба кандидата"
    assert any("overlap-bad" in n and "отклонена" in n for n in notes)


def test_slide_audit_retry_budget_is_capped():
    """Тест брифа: "число попыток ограничь". У ЛЦТ2026/WorkSpace на
    некоторые `kind` приходится больше десятка паттернов, а колода из
    12-15 слайдов обязана укладываться в 5 минут (ТЗ) — перебор НЕ может
    быть исчерпывающим. Четвёртый кандидат идеален (без единой находки
    уровня ошибки), но раскладка не обязана его достать, если первые три
    уже исчерпали бюджет попыток."""
    bad1 = _overlapping_section_pattern("bad-1")
    bad2 = _overlapping_section_pattern("bad-2")
    bad3 = _overlapping_section_pattern("bad-3")
    perfect = _clean_section_pattern("perfect-4th")
    # `headline`/`subhead` — реальные слова, не «Заголовок»/«Подзаголовок»
    # (см. докстроку `test_slide_is_never_left_empty_when_no_candidate_
    # passes_the_audit` про `_placeholder_text_hit`, Task 22): с текстом-
    # заглушкой ни headline, ни subhead не легли бы на слайд вовсе, и L02
    # (наложение) — за отсутствием ДВУХ нарисованных блоков — не сработал
    # бы на первом же кандидате, а весь смысл теста в том, что он ДОЛЖЕН
    # сработать и быть отклонён.
    slide_spec = SlideSpec(index=0, kind="section", headline="Итоги квартала", subhead="Что изменилось")
    prs, canvas, audit_config = _new_deck_in_progress()

    chosen, notes = _place_best_candidate(
        prs, slide_spec, [bad1, bad2, bad3, perfect], PROFILE, canvas, audit_config, bullet_char="•",
    )

    assert chosen.pattern_id != "perfect-4th", "бюджет обязан остановить перебор до четвёртого кандидата"
    assert not any("perfect-4th" in n for n in notes), "четвёртый кандидат не должен был даже пробоваться"


def test_slide_audit_picks_the_least_bad_candidate_when_none_pass():
    """Тест брифа: "потом берётся кандидат с наименьшим числом ошибок".
    `worse` наваливает ТРИ взаимных наложения (headline/subhead/source на
    одном боксе), `better` — только одно (headline/subhead); ни один не
    проходит чисто, но `better` обязан победить как менее плохой."""
    worse = _overlapping_section_pattern("worse")
    extra = PatternSlot(
        role="source", box=Box(0.05, 0.1, 0.5, 0.1), size_pt=10.0, color_hex=None, align="l", max_chars=200, wraps=True,
    )
    worse = replace(worse, slots=[*worse.slots, extra])
    better = _overlapping_section_pattern("better")
    # `headline`/`subhead` — не текст-заглушка, той же причиной, что у
    # `test_slide_audit_retry_budget_is_capped` выше: без реально
    # нарисованного текста наложению (L02) нечего накладывать.
    slide_spec = SlideSpec(
        index=0, kind="section", headline="Итоги квартала", subhead="Что изменилось", source_note="Источник данных",
    )
    prs, canvas, audit_config = _new_deck_in_progress()

    chosen, notes = _place_best_candidate(prs, slide_spec, [worse, better], PROFILE, canvas, audit_config, bullet_char="•")

    assert chosen.pattern_id == "better"
    assert any("ни один" in n for n in notes)


def test_slide_is_never_left_empty_when_no_candidate_passes_the_audit():
    """Тест брифа: "если ни один кандидат не прошёл, слайд всё равно
    собирается лучшим из возможных ... пустого слайда быть не должно
    никогда" — плюс "результат перебора логируй".

    `headline`/`subhead` — реальные слова, не «Заголовок»/«Подзаголовок»
    (Task 22, отчёт задачи: ровно эти два слова с этой задачи — маркеры
    текста-заглушки в `config/audit.yaml::integrity.placeholder_patterns`,
    `place_slide` теперь не кладёт такой текст на слайд вовсе, см.
    `builder._placeholder_text_hit` — со старым текстом фикстуры оба слота
    остались бы пустыми и тест ложно проверял бы уже не то, что заявлено в
    докстроке)."""
    bad = _overlapping_section_pattern("only-bad")
    slide_spec = SlideSpec(index=0, kind="section", headline="Итоги квартала", subhead="Что изменилось")
    prs, canvas, audit_config = _new_deck_in_progress()

    chosen, notes = _place_best_candidate(prs, slide_spec, [bad], PROFILE, canvas, audit_config, bullet_char="•")

    assert chosen is not None
    assert len(prs.slides) == 1
    texts = [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert texts, "слайд не должен оставаться пустым, даже если ни один кандидат не прошёл аудит"
    assert notes, "результат перебора обязан быть залогирован"


# ---------------------------------------------------------------------------
# Task 22, отчёт задачи: текст-заглушка не должен попасть на слайд вообще
# ---------------------------------------------------------------------------


def test_placeholder_looking_text_is_not_drawn_on_the_slide():
    """Слайд, у которого `headline`/`subhead` содержат маркер из
    `config/audit.yaml::integrity.placeholder_patterns` (регистронезависимое
    вхождение подстроки — тот же критерий, что и у постфактум-аудита I02),
    не должен получить этот текст на холсте вовсе: `place_slide` обязана
    пропустить такой слот и оставить честную находку (`_placeholder_text_
    hit`), а не полагаться на то, что I02 поймает уже собранный файл
    постфактум (находка отчёта задачи — «Титульный слайд презентации» на
    первом слайде контрольного ЛЦТ2026 ушло на защиту, живьём)."""
    pattern = _clean_section_pattern("clean")
    slide_spec = SlideSpec(index=0, kind="section", headline="Заголовок", subhead="Подзаголовок")
    prs, canvas, audit_config = _new_deck_in_progress()

    _, notes = _place_best_candidate(prs, slide_spec, [pattern], PROFILE, canvas, audit_config, bullet_char="•")

    texts = [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert texts == [], f"текст-заглушка не должен был попасть на слайд, а попал: {texts!r}"
    # `_place_best_candidate` копит находки на своей ВНУТРЕННЕЙ копии спека
    # (`trial_spec = replace(slide_spec, findings=[])`) и возвращает их
    # ВЫЗЫВАЮЩЕЙ СТОРОНЕ через `notes`, не пишет назад в переданный
    # `slide_spec` — `build_deck` сам делает `slide_spec.findings.extend
    # (notes)`, тест проверяет тот же контракт напрямую.
    assert any("похож" in n and "заглушк" in n for n in notes), (
        f"расхождение обязано остаться честной находкой: {notes!r}"
    )


def test_count_embedded_photos_matches_by_content_not_by_name(tmp_path):
    """Task 22, отчёт задачи, находка "пайплайн рапортует не то, что в
    файле": `count_embedded_photos` — единственный источник правды о том,
    что физически легло на слайды (по байтам `.pptx`), не о том, что
    РАСПРЕДЕЛИЛ планировщик (`plan.photos.PhotoAssignmentReport.placed_
    count`) — эти два числа теперь намеренно печатаются раздельно, и в
    `cli.py`, и в `scripts/build_submission.py` (см. их докстроки у места
    вызова). Сравнение — по байтам (sha256), не по имени файла: python-pptx
    переименовывает media при вставке (`ppt/media/imageN.ext`), имя файла
    контент-пакета внутри архива не сохраняется — тест вставляет фото под
    ДРУГИМ путём/именем, чем то, с которым его ищет вызывающий код, ровно
    как это происходит в реальной сборке."""
    from PIL import Image

    embedded_src = tmp_path / "source-team.jpg"
    Image.new("RGB", (400, 300), color=(10, 20, 30)).save(embedded_src)
    not_embedded_src = tmp_path / "source-approver.jpg"
    Image.new("RGB", (400, 300), color=(200, 30, 60)).save(not_embedded_src)

    prs = Presentation(str(TEMPLATE))
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.add_picture(str(embedded_src), 0, 0, width=100, height=100)
    out_path = tmp_path / "manual.pptx"
    prs.save(str(out_path))

    # `team.jpg` — контент-пакетное имя, разное с `source-team.jpg` на
    # диске (как в реальной сборке, где имя файла контент-пакета и имя
    # внутри архива никогда не совпадают) — счётчик обязан найти его по
    # содержимому.
    user_photos = {"team.jpg": embedded_src, "approver.jpg": not_embedded_src}
    assert count_embedded_photos(out_path, user_photos) == 1
    assert count_embedded_photos(out_path, {}) == 0, "без фотографий контент-пакета — честный ноль, не падение"


# ---------------------------------------------------------------------------
# Task 23, находки ручной проверки контрольного ЛЦТ2026: (1) содержание,
# которому не нашлось слота, исчезало бесследно; (2) плашки карточек
# рисовались пустыми на слайде, у которого карточек нет вовсе.
# ---------------------------------------------------------------------------

_PLAQUE_WIDTH = 0.31


def _three_plaque_cards_pattern(pattern_id: str) -> Pattern:
    """Слепок раскладки `slide24` контрольного ЛЦТ2026: заголовок, три
    слота `card_body` и три белые плашки-подложки под ними
    (`repeat_group=True`). Слота под подзаголовок раскладка НЕ несёт —
    ровно поэтому подзаголовок первого слайда и терялся."""
    headline = PatternSlot(
        role="headline", box=Box(0.05, 0.06, 0.9, 0.08), size_pt=28.0, color_hex=None,
        align="l", max_chars=120, wraps=True,
    )
    bodies = [
        PatternSlot(
            role="card_body", box=Box(0.05 + i * 0.32, 0.48, 0.27, 0.34), size_pt=14.0, color_hex=None,
            align="l", max_chars=200, wraps=True,
        )
        for i in range(3)
    ]
    plaques = [
        DecorShape(
            kind="shape", box=Box(0.035 + i * 0.32, 0.44, _PLAQUE_WIDTH, 0.5), rotation=0.0,
            flip_h=False, flip_v=False, fill_hex="#FFFFFF", has_fill=True, fill_kind="solid",
            repeat_group=True, repeat_index=i,
        )
        for i in range(3)
    ]
    return Pattern(
        pattern_id=pattern_id, source_slide_index=[0], layout_id=_real_layout_id(), kind="cards",
        slots=[headline, *bodies], repeat=RepeatSpec(axis="x", count=3, step=0.32, slot_roles=["card_body"]),
        decor=plaques,
        capacity=Capacity(max_items=3, max_chars_per_item=200, max_bullets=0, max_series=0, max_rows=0, max_cols=0),
        score=1.0, is_dark=False,
    )


def _plaque_widths_on(slide) -> list[int]:
    expected = round(_PLAQUE_WIDTH * PROFILE.canvas_width_emu)
    return [sh.width for sh in slide.shapes if abs(sh.width - expected) <= 1]


def test_content_without_a_slot_leaves_an_honest_finding():
    """Главная находка Task 23: подзаголовок с цифрами из брифа был в
    плане, слота под него в раскладке нет — на слайд он не попал и НИГДЕ
    не был упомянут (ни находки, ни строки в отчёте). Слайд по-прежнему
    собирается (это не ошибка сборки), но расхождение обязано быть
    названо."""
    pattern = _three_plaque_cards_pattern("plaques-drop")
    slide_spec = SlideSpec(
        index=0, kind="two_col", headline="Итоги 2026: учебная платформа",
        subhead="Рост доведения до 57%, стоимость выпускника −44%",
    )
    prs, canvas, audit_config = _new_deck_in_progress()

    _, notes = _place_best_candidate(prs, slide_spec, [pattern], PROFILE, canvas, audit_config, bullet_char="•")

    assert len(prs.slides) == 1
    assert any("подзаголовок" in n and "Рост доведения" in n for n in notes), (
        f"потерянный подзаголовок обязан остаться находкой: {notes!r}"
    )


def test_empty_repeat_plaques_are_not_drawn():
    """Находка №2: слайду с одним заголовком досталась раскладка на три
    карточки — три белые плашки заняли больше половины слайда и не несли
    ни буквы. Плашка, в чей слот ничего не легло, не рисуется."""
    pattern = _three_plaque_cards_pattern("plaques-empty")
    slide_spec = SlideSpec(index=0, kind="two_col", headline="Итоги 2026: учебная платформа")
    prs, canvas, audit_config = _new_deck_in_progress()

    _place_best_candidate(prs, slide_spec, [pattern], PROFILE, canvas, audit_config, bullet_char="•")

    assert _plaque_widths_on(prs.slides[0]) == [], "пустые плашки карточек не должны рисоваться"


def test_plaques_of_filled_cards_are_still_drawn():
    """Обратная сторона той же правки: карточки с текстом свои плашки
    сохраняют — убирать нужное так же плохо, как рисовать лишнее."""
    pattern = _three_plaque_cards_pattern("plaques-filled")
    slide_spec = SlideSpec(
        index=0, kind="cards", headline="Что изменилось",
        blocks=[CardBlock(items=[
            Card(title="", body="Доведение выросло с 41% до 57%."),
            Card(title="", body="Стоимость выпускника снизилась на 44%."),
            Card(title="", body="Время проверки работ сократилось вдвое."),
        ])],
    )
    prs, canvas, audit_config = _new_deck_in_progress()

    _place_best_candidate(prs, slide_spec, [pattern], PROFILE, canvas, audit_config, bullet_char="•")

    assert len(_plaque_widths_on(prs.slides[0])) == 3


def test_a_plaque_under_a_slot_taken_by_plain_text_is_kept():
    """Слоты повтора не зарезервированы, когда на слайде нет карточек —
    обычный текстовый блок вправе занять слот карточки, и плашка под ним
    обязана остаться (пустыми уходят только соседние)."""
    pattern = _three_plaque_cards_pattern("plaques-text")
    slide_spec = SlideSpec(
        index=0, kind="two_col", headline="Что изменилось",
        blocks=[TextBlock(text="Доведение выросло с 41% до 57%.")],
    )
    prs, canvas, audit_config = _new_deck_in_progress()

    _place_best_candidate(prs, slide_spec, [pattern], PROFILE, canvas, audit_config, bullet_char="•")

    assert len(_plaque_widths_on(prs.slides[0])) == 1


# ---------------------------------------------------------------------------
# Заметки докладчика: уточнение заказчика от 23 сентября 2026 — на выходе
# ждут «готовые слайды и текст к каждому слайду», потому что на защите по
# слайдам ещё и рассказывают
# ---------------------------------------------------------------------------


def test_speaker_notes_reach_the_pptx_notes_page():
    """`SlideSpec.speaker_notes` обязан доехать до страницы заметок готового
    файла.

    До этой правки поле заполнялось моделью (`agents/slide-writer/AGENT.md`,
    схема ответа), ехало через `plan.spec` и молча терялось на сборке: ни
    `build_deck`, ни выгрузка его не читали. Модель писала текст, который
    никто никогда не видел."""
    spec = replace(
        SAMPLE_SPEC,
        slides=[
            replace(SAMPLE_SPEC.slides[0], speaker_notes="Начать с цифры 31,5 часа — она держит весь рассказ."),
            replace(SAMPLE_SPEC.slides[1], speaker_notes=None),
        ],
    )

    out = build_deck(spec, PROFILE, TEMPLATE, Variant.dense)
    prs = Presentation(str(out))

    assert "31,5 часа" in prs.slides[0].notes_slide.notes_text_frame.text


def test_a_slide_without_speaker_notes_gets_no_notes_page():
    """Пустое поле не должно порождать пустую страницу заметок — иначе у
    каждого слайда колоды появляется пустой лист заметок, которого в
    исходном шаблоне не было."""
    spec = replace(
        SAMPLE_SPEC,
        slides=[replace(SAMPLE_SPEC.slides[0], speaker_notes=None)],
    )

    out = build_deck(spec, PROFILE, TEMPLATE, Variant.dense)
    prs = Presentation(str(out))

    assert not prs.slides[0].has_notes_slide
