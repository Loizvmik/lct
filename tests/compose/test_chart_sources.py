"""Задача V1: график в сборке. Родной график примера получает наши данные
со стилем дизайнера; без родного графика наш встаёт в рамку образца
оформления по его палитре и правилам шаблона."""
from __future__ import annotations
import functools
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from deckforge.compose.builder import Variant, build_deck
from deckforge.compose.charts import ChartSpec, Series, add_chart
from deckforge.ooxml.geometry import Box
from deckforge.plan.spec import BulletBlock, ChartSeriesData, ChartVisual, DeckSpec, SlideSpec, Visual
from deckforge.template.profile import TemplateProfile

TEMPLATES = Path("dataset/templates")
LCT = TEMPLATES / "ЛЦТ2026 Шаблон презентации.pptx"
VKE = TEMPLATES / "Шаблон презентации VK Education.pptx"

COVERAGE = ChartVisual(
    kind="bar", categories=["I", "II", "III", "IV"],
    series=[
        ChartSeriesData(name="Регистраций", values=[4100, 6300, 5800, 9400]),
        ChartSeriesData(name="Приступили", values=[2870, 4220, 4060, 6110]),
        ChartSeriesData(name="Завершили курс", values=[1190, 1940, 2030, 3480]),
    ],
    axis_titles=("Квартал", "Человек"),
)


@functools.lru_cache(maxsize=None)
def _profile(path: Path) -> TemplateProfile:
    return TemplateProfile.from_file(path, cache_dir=None)


def _chart_slide(profile, index: int, pattern_id: str, chart: ChartVisual = COVERAGE, blocks=()) -> SlideSpec:
    kind = next(p.kind for p in profile.patterns if p.pattern_id == pattern_id)
    return SlideSpec(
        # Заголовок в рамку примера (у ЛЦТ2026 slide21 это 14 знаков при
        # 28 pt): предел ужимания кегля (задача V3) длинный не пропустит.
        index=index, kind=kind, headline="Охват растёт",
        blocks=list(blocks), visual=Visual(kind="chart", chart=chart),
        source_note="Итоги учебной платформы 2026", pattern_id=pattern_id,
    )


def _charts(slide):
    return [sh for sh in slide.shapes if getattr(sh, "has_chart", False)]


def test_native_chart_of_the_example_takes_our_data_and_keeps_its_style():
    """ЛЦТ2026, slide21: родной столбчатый график примера. Тип, легенда и
    оформление остаются, данные и встроенная таблица Excel наши."""
    spec = DeckSpec(title="Итоги", language="ru", slides=[
        _chart_slide(_profile(LCT), 0, "slide21", blocks=[BulletBlock(items=["Регистраций в IV квартале 9 400"])]),
    ])
    out = build_deck(spec, _profile(LCT), LCT, Variant.visual)
    prs = Presentation(str(out))
    charts = _charts(prs.slides[0])
    assert len(charts) == 1
    chart = charts[0].chart
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    plot = chart.plots[0]
    assert list(plot.categories) == ["I", "II", "III", "IV"]
    assert [s.name for s in plot.series] == ["Регистраций", "Приступили", "Завершили курс"]
    assert list(plot.series[0].values) == [4100, 6300, 5800, 9400]
    assert chart.has_legend and chart.value_axis.has_title
    assert any("родной график примера" in f for f in spec.slides[0].findings)
    xlsx = chart.part.chart_workbook.xlsx_part.blob
    with zipfile.ZipFile(__import__("io").BytesIO(xlsx)) as zf:
        text = " ".join(zf.read(n).decode("utf-8", "ignore") for n in zf.namelist() if n.endswith(".xml"))
    assert "Регистраций" in text, "встроенная таблица Excel обновлена"
    fills = [s._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}schemeClr").get("val")
             for s in plot.series]
    assert len(set(fills)) == 3, "третий ряд не повторяет цвет второго"


def test_two_slides_on_one_native_chart_do_not_share_data():
    other = ChartVisual(kind="bar", categories=["А", "Б", "В"], series=[ChartSeriesData(name="Ряд", values=[1, 2, 3])],
                        axis_titles=("Категория", "Значение"))
    spec = DeckSpec(title="Итоги", language="ru", slides=[
        _chart_slide(_profile(LCT), 0, "slide21", blocks=[BulletBlock(items=["Первый"])]),
        _chart_slide(_profile(LCT), 1, "slide21", chart=other, blocks=[BulletBlock(items=["Второй"])]),
    ])
    prs = Presentation(str(build_deck(spec, _profile(LCT), LCT, Variant.visual)))
    first, second = (_charts(s)[0].chart for s in prs.slides)
    assert list(first.plots[0].categories) == ["I", "II", "III", "IV"]
    assert list(second.plots[0].categories) == ["А", "Б", "В"]
    assert first.part is not second.part


def test_chart_goes_into_the_frame_of_the_chart_sample_with_template_rules():
    """VK Education, слайд 49 «Пример оформления графика»: картинка-образец
    уходит, наш график встаёт в её рамку, в цветах образца, без сетки, с
    подписями значений и зазором не уже перекрытия (правила слайда 51)."""
    profile = _profile(VKE)
    spec = DeckSpec(title="Итоги", language="ru", slides=[_chart_slide(profile, 0, "slide49")])
    prs = Presentation(str(build_deck(spec, profile, VKE, Variant.visual)))
    slide = prs.slides[0]
    charts = _charts(slide)
    assert len(charts) == 1
    big_pictures = [
        sh for sh in slide.shapes
        if sh.shape_type == 13 and sh.width * sh.height > 0.15 * prs.slide_width * prs.slide_height
    ]
    assert big_pictures == [], "картинка-образец графика удалена"
    proto = next(p for p in profile.chart_prototypes if p.pattern_id == "slide49")
    frame = charts[0]
    assert abs(frame.left / prs.slide_width - proto.frame_box.left) < 0.02
    assert abs(frame.width / prs.slide_width - proto.frame_box.width) < 0.02
    chart = frame.chart
    assert not chart.value_axis.has_major_gridlines
    plot = chart.plots[0]
    assert plot.has_data_labels
    assert plot.overlap == -50 and plot.gap_width >= 50
    first_fill = plot.series[0].format.fill.fore_color.rgb
    assert str(first_fill) == proto.palette[0].lstrip("#")


def test_rules_are_not_applied_to_a_single_series_gap():
    """Нулевой боковой зазор у одного ряда склеил бы столбики: у одного
    ряда зазор и перекрытие шаблона не применяются, сетка убирается."""
    profile = _profile(VKE)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = add_chart(slide, Box(0.1, 0.1, 0.8, 0.8), ChartSpec(
        kind="bar", categories=["А", "Б", "В"], series=[Series(name="Ряд", values=[1, 2, 3])],
        axis_titles=("Категория", "Значение"),
    ), profile)
    plot = frame.chart.plots[0]
    assert plot.gap_width != 0
    assert not frame.chart.value_axis.has_major_gridlines


def test_three_series_on_a_two_colour_sample_get_three_colours():
    """Образец слайда 47 VK Education даёт два цвета, рядов три: третий
    берётся из палитры профиля и не совпадает на глаз с первым."""
    profile = _profile(VKE)
    spec = DeckSpec(title="Итоги", language="ru", slides=[_chart_slide(profile, 0, "slide47")])
    prs = Presentation(str(build_deck(spec, profile, VKE, Variant.visual)))
    plot = _charts(prs.slides[0])[0].chart.plots[0]
    fills = [str(s.format.fill.fore_color.rgb) for s in plot.series]
    assert len(set(fills)) == 3


def test_a_plate_filled_by_the_theme_style_counts_as_background():
    """Плашка ЛЦТ2026 (слайд 12) залита только `p:style/a:fillRef`: под ней
    светло, и заголовок и подписи графика на ней не белые."""
    from lxml import etree

    from deckforge.compose.builder import _shape_solid_fill

    xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:spPr><a:ln><a:noFill/></a:ln></p:spPr>'
        '<p:style><a:fillRef idx="1"><a:schemeClr val="lt1"/></a:fillRef></p:style></p:sp>'
    )
    color = _shape_solid_fill(etree.fromstring(xml), {"lt1": "#FFFFFF"}, {"bg1": "lt1"})
    assert color is not None and color.hex == "#FFFFFF"
