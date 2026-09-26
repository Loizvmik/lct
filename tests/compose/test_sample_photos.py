"""Фото-образцы примера не уезжают в чужую колоду (задача V2).

У VK Education «Паттерн + фото» (`slide35`) рисует фотографию девушки в
кресле на лейауте, а не на слайде-примере: майнинг её не видел, клон и
сборка с нуля выводили её на слайд с нашим текстом (живой прогон 27
сентября 2026, airy, слайд 3)."""
from __future__ import annotations
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.compose import builder
from deckforge.compose.clone import sample_slides_by_number
from deckforge.ooxml.geometry import Canvas
from deckforge.ooxml.ns import qn
from deckforge.pattern.candidates import sample_photo_fits
from deckforge.pattern.intent import SlideIntent
from deckforge.plan.spec import BulletBlock, SlideSpec

TEMPLATE = Path("dataset/templates/Шаблон презентации VK Education.pptx")


@pytest.fixture(scope="module")
def profile(profile_fixture):
    return profile_fixture(TEMPLATE.name)


@pytest.fixture
def deck():
    prs = Presentation(str(TEMPLATE))
    sources = sample_slides_by_number(prs)
    builder._clear_sample_slides(prs)
    return prs, sources


def _model(profile, pattern_id: str):
    return next(m for m in profile.patterns if m.pattern_id == pattern_id)


def _spec(pattern_id: str) -> SlideSpec:
    return SlideSpec(
        index=2, kind="bullets", headline="Ожидание сократится на 80%",
        blocks=[BulletBlock(items=["Автоподмена согласующего в отпуске", "Правила маршрута для закупок"])],
        pattern_id=pattern_id,
    )


def _layout_photo_visible(slide, photo_ids) -> bool:
    """Фото лейаута видно на слайде: графика лейаута не скрыта, либо
    картинка с тем же id скопирована на сам слайд."""
    if slide._element.get("showMasterSp") != "0":
        return True
    ids = {el.get("id") for el in slide._element.iter(qn("p:cNvPr"))}
    return bool(ids & set(photo_ids))


def test_layout_photo_of_the_example_is_a_sample_photo(profile):
    model = _model(profile, "slide35")

    assert model.photo_frames == 1 and model.layout_photo_ids
    assert model.photo_area > 0.3
    # Фирменная графика обложки (цветные полукруги) не фото.
    cover = _model(profile, "slide1")
    assert cover.photo_frames == 0 and not cover.layout_photo_ids


def test_planner_does_not_take_a_layout_whose_photo_would_leave_half_the_slide_empty(profile):
    intent = SlideIntent(index=2, outline_kind="context", intent="Контекст", items=3)

    assert not sample_photo_fits(intent, _model(profile, "slide35"))
    assert sample_photo_fits(intent, _model(profile, "slide21"))


def test_clone_hides_the_layout_photo_and_keeps_the_text(profile, deck):
    prs, sources = deck
    model = _model(profile, "slide35")
    pattern = builder._pattern_from_model(model)

    outcome = builder.place_slide_by_clone(prs, _spec("slide35"), pattern, profile, AuditConfig.load(), sources[35])

    assert outcome.reason is None
    slide = prs.slides[-1]
    assert not _layout_photo_visible(slide, model.layout_photo_ids)
    assert any("Автоподмена" in (sp.text_frame.text if sp.has_text_frame else "") for sp in slide.shapes)


def test_scratch_build_hides_the_layout_photo_too(profile, deck):
    prs, _sources = deck
    model = _model(profile, "slide35")

    builder.place_slide(prs, _spec("slide35"), builder._pattern_from_model(model), profile, AuditConfig.load())

    assert not _layout_photo_visible(prs.slides[-1], model.layout_photo_ids)


def test_hiding_keeps_the_rest_of_the_layout_graphics(profile, deck):
    """Скрыть графику лейаута целиком нельзя: вместе с фото ушёл бы логотип
    и фирменный узор. Всё, кроме фото, копируется на слайд."""
    from deckforge.compose.clone import hide_layout_photos

    prs, _sources = deck
    layout = next(l for l in prs.slide_layouts if l.name == "1_Титульный слайд")
    slide = prs.slides.add_slide(layout)
    pictures = [el for el in layout.shapes._spTree.iter(qn("p:pic")) if el.find(".//" + qn("p:ph")) is None]
    fake_photo = pictures[0].find(qn("p:nvPicPr") + "/" + qn("p:cNvPr")).get("id")
    canvas = Canvas(width_emu=prs.slide_width, height_emu=prs.slide_height)

    hidden = hide_layout_photos(slide, [fake_photo], canvas)

    assert hidden == 1 and slide._element.get("showMasterSp") == "0"
    copied = [el for el in slide.shapes._spTree.iter(qn("p:pic"))]
    assert len(copied) == len(pictures) - 1
    for pic in copied:
        rid = pic.find(".//" + qn("a:blip")).get(qn("r:embed"))
        assert slide.part.rels[rid].reltype.endswith("/image")
