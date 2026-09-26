"""Сборка слайда клоном слайда-примера (`compose.clone`, `builder.place_slide_by_clone`).

Шаблон: VK Education из `dataset/templates/` (тот же файл, что лежит в
корне репозитория, байт в байт): на нём клон проверялся глазами, и у него
есть всё, что клон обязан сохранить: картинки-иконки, значки с номерами,
пустые плейсхолдеры титула. Модель не вызывается: профиль строится
разбором без сети (`cache_dir=None`, как в `conftest.py`).
"""
from __future__ import annotations
import copy
from dataclasses import replace
import zipfile
from collections import Counter
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.compose import builder
from deckforge.compose.blocks import Paragraph
from deckforge.compose.clone import (
    bind_text, box_iou, clone_example_slide, clone_map, fill_native_table, match_slots, native_table,
    prune_unfilled, sample_slides_by_number, shape_text, slide_refs,
)
from deckforge.ooxml.geometry import Box, Canvas
from deckforge.ooxml.ns import qn
from deckforge.ooxml.walk import walk_shapes
from deckforge.plan.spec import BulletBlock, Card, CardBlock, DeckSpec, SlideSpec, TableVisual, Visual
from deckforge.plan.variants import Variant

TEMPLATE = Path("dataset/templates/Шаблон презентации VK Education.pptx")


@pytest.fixture(scope="module")
def profile(profile_fixture):
    return profile_fixture(TEMPLATE.name)


@pytest.fixture
def deck():
    """Колода шаблона без примеров и сами примеры по номеру: ровно то, с
    чем работает `build_deck`."""
    prs = Presentation(str(TEMPLATE))
    sources = sample_slides_by_number(prs)
    builder._clear_sample_slides(prs)
    return prs, sources


def _pattern(profile, pattern_id: str):
    return next(builder._pattern_from_model(m) for m in profile.patterns if m.pattern_id == pattern_id)


def _canvas(profile) -> Canvas:
    return Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)


def _texts(slide) -> list[str]:
    return [shape_text(el).strip() for el in slide._element.iter(qn("p:sp")) if shape_text(el).strip()]


def _cards_spec(n: int) -> SlideSpec:
    return SlideSpec(
        index=0, kind="two_col", headline="Что нужно для раскатки",
        blocks=[CardBlock(items=[Card(body=f"Пункт номер {i + 1} плана раскатки") for i in range(n)])],
    )


def test_clone_keeps_shapes_and_image_links(profile, deck):
    prs, sources = deck
    source = sources[26]  # три иконки-картинки над тремя тезисами
    layout = builder._find_layout(prs, _pattern(profile, "slide26").layout_id)

    slide = clone_example_slide(prs, source, layout)

    canvas = _canvas(profile)
    assert len(list(walk_shapes(slide._element, canvas))) == len(list(walk_shapes(source._element, canvas)))
    blips = list(slide._element.iter(qn("a:blip")))
    assert blips, "у примера есть картинки: у клона они обязаны остаться"
    source_blobs = {rel.target_part.blob for rel in source.part.rels.values() if rel.reltype.endswith("/image")}
    for blip in blips:
        rel = slide.part.rels[blip.get(qn("r:embed"))]
        assert rel.reltype.endswith("/image")
        assert rel.target_part.blob in source_blobs
    assert not any(r.reltype.endswith("/notesSlide") for r in slide.part.rels.values())


def test_bind_text_keeps_size_and_color_of_first_run(deck):
    _prs, sources = deck
    element = run_pr = None
    for slide in sources.values():
        for sp in slide._element.iter(qn("p:sp")):
            run = sp.find(".//" + qn("a:r"))
            r_pr = run.find(qn("a:rPr")) if run is not None else None
            if r_pr is not None and r_pr.get("sz") and r_pr.find(qn("a:solidFill")) is not None:
                element, run_pr = sp, r_pr
                break
        if element is not None:
            break
    assert element is not None, "в шаблоне нет run с явным кеглем и цветом: тест не о чем"
    size, color = run_pr.get("sz"), _color(run_pr)
    body_pr_before = _xml(element.find(qn("p:txBody")).find(qn("a:bodyPr")))

    bind_text(element, [Paragraph("Первый пункт", bullet=True), Paragraph("Второй пункт", bullet=True)], bullet_char="•")

    paragraphs = element.find(qn("p:txBody")).findall(qn("a:p"))
    assert [_paragraph_text(p) for p in paragraphs] == ["Первый пункт", "Второй пункт"]
    for p in paragraphs:
        runs = p.findall(qn("a:r"))
        assert len(runs) == 1
        assert runs[0].find(qn("a:rPr")).get("sz") == size
        assert _color(runs[0].find(qn("a:rPr"))) == color
    assert _xml(element.find(qn("p:txBody")).find(qn("a:bodyPr"))) == body_pr_before


def test_unfilled_repeat_units_are_removed(profile, deck):
    prs, sources = deck
    pattern = _pattern(profile, "slide21")  # четыре кружка с номерами над четырьмя описаниями
    spec = _cards_spec(2)

    outcome = builder.place_slide_by_clone(prs, spec, pattern, profile, AuditConfig.load(), sources[21])

    assert outcome.reason is None
    texts = _texts(prs.slides[-1])
    badges = sorted(t for t in texts if t.isdigit())
    assert badges == ["1", "2"], texts
    assert sum("плана раскатки" in t for t in texts) == 2


def test_missing_headline_shape_falls_back_to_building_from_scratch(profile, deck):
    prs, sources = deck
    pattern = _pattern(profile, "slide21")
    # Коробка заголовка уведена туда, где в примере пусто, и id фигуры
    # стёрт: клону некуда положить заголовок, и слайд обязан собраться
    # старым путём.
    slots = [
        replace(s, box=Box(0.05, 0.9, 0.3, 0.05), source_shape_id=None) if s.role == "headline" else s
        for s in pattern.slots
    ]
    broken = replace(pattern, slots=slots)
    spec = _cards_spec(2)

    outcome = builder.place_slide_by_clone(prs, spec, broken, profile, AuditConfig.load(), sources[21])
    assert outcome.reason is not None and "headline" in outcome.reason
    assert len(prs.slides) == 0

    chosen, notes = builder._place_best_candidate(
        prs, spec, [broken], profile, _canvas(profile), AuditConfig.load(), source_slides=sources,
    )
    assert chosen is broken
    assert len(prs.slides) == 1
    assert any("не принят" in n for n in notes)
    name = prs.slides[0]._element.find(qn("p:cSld")).get("name") or ""
    assert not name.startswith(builder.CLONE_MARK_PREFIX)


def test_sample_text_does_not_survive_in_unfilled_slots(profile, deck):
    prs, sources = deck
    pattern = _pattern(profile, "slide26")  # два тезиса под иконками, третий тезис и сноска
    spec = _cards_spec(2)

    outcome = builder.place_slide_by_clone(prs, spec, pattern, profile, AuditConfig.load(), sources[26])

    assert outcome.reason is None
    texts = " ".join(_texts(prs.slides[-1]))
    for slot in pattern.slots:
        for line in (slot.sample_text or "").split("\n"):
            if len(line.strip()) > 3:
                assert line.strip() not in texts, f"остался текст примера: {line!r}"


def test_build_deck_marks_cloned_slides(profile):
    spec = DeckSpec(title="Клон", language="ru", slides=[_cards_spec(2)])
    spec.slides[0].pattern_id = "slide21"
    cloned = Presentation(str(builder.build_deck(spec, profile, TEMPLATE, Variant.dense, clone_examples=True)))
    names = [s._element.find(qn("p:cSld")).get("name") or "" for s in cloned.slides]
    assert names and names[0].startswith(builder.CLONE_MARK_PREFIX)

    spec_old = DeckSpec(title="Клон", language="ru", slides=[_cards_spec(2)])
    spec_old.slides[0].pattern_id = "slide21"
    rebuilt = Presentation(str(builder.build_deck(spec_old, profile, TEMPLATE, Variant.dense, clone_examples=False)))
    assert not any(
        (s._element.find(qn("p:cSld")).get("name") or "").startswith(builder.CLONE_MARK_PREFIX)
        for s in rebuilt.slides
    )


def test_user_photo_replaces_example_picture_without_duplicate_parts(profile):
    """Фото пользователя встаёт в фигуру-картинку примера (рамка телефона,
    круглая маска), а архив не получает двух частей с одним именем: новая
    картинка берёт имя, которое python-pptx считает свободным после
    очистки примеров, а следующий клон снова связывается со старой
    картинкой примера под тем же именем (см. `fix_duplicate_partnames`)."""
    photo = Path("fixtures/content-packs/edu-platform/photos/students-project-review.jpg")
    slides = [
        SlideSpec(
            index=i, kind=_pattern(profile, pid).kind, headline="Студенты защищают проекты", pattern_id=pid,
            blocks=[BulletBlock(items=["Разбор проекта с ментором", "Демо на планшете"])],
            visual=Visual(kind="photo", photo_name="p.jpg"),
        )
        for i, pid in enumerate(["slide30", "slide3", "slide33"])
    ]
    spec = DeckSpec(title="Фото", language="ru", slides=slides)

    out = builder.build_deck(
        spec, profile, TEMPLATE, Variant.visual, user_photos={"p.jpg": photo}, clone_examples=True,
    )

    names = zipfile.ZipFile(out).namelist()
    assert [n for n, k in Counter(names).items() if k > 1] == []
    assert builder.count_embedded_photos(out, {"p.jpg": photo}) == 1
    assert any("собран клоном" in f for s in spec.slides for f in s.findings)


def _paragraph_text(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


def _color(r_pr) -> str:
    fill = r_pr.find(qn("a:solidFill"))
    return _xml(fill) if fill is not None else ""


def _xml(el) -> bytes:
    return etree.tostring(el) if el is not None else b""


def test_bind_text_marks_a_bold_paragraph_and_keeps_the_rest_regular(deck):
    """Заголовок карточки, склеенный с телом (`blocks._assign_cards`),
    приходит абзацем с `bold=True`: у его run появляется `b="1"`, у
    остальных начертание примера не меняется."""
    _prs, sources = deck
    element = next(
        sp for slide in sources.values() for sp in slide._element.iter(qn("p:sp"))
        if sp.find(".//" + qn("a:r")) is not None
    )
    bind_text(element, [Paragraph("Бюджет", bold=True), Paragraph("2,4 млн ₽")])

    runs = element.findall(".//" + qn("a:r"))
    assert [r.find(qn("a:t")).text for r in runs] == ["Бюджет", "2,4 млн ₽"]
    assert runs[0].find(qn("a:rPr")).get("b") == "1"
    assert runs[1].find(qn("a:rPr")) is None or runs[1].find(qn("a:rPr")).get("b") != "1"


def test_bind_text_into_an_empty_placeholder_leaves_the_size_to_the_layout(deck):
    """Пустой плейсхолдер примера несёт в `endParaRPr` кегль-заглушку
    экспорта (VK Education, обложка раздела: 16 pt при заголовке макета
    48 pt). Записанный явно, он перекрывал унаследованный: обложка выходила
    кеглем текста (прогон 26 сентября 2026)."""
    _prs, sources = deck
    element = next(
        (sp for slide in sources.values() for sp in slide._element.iter(qn("p:sp"))
         if sp.find(".//" + qn("p:ph")) is not None and sp.find(".//" + qn("a:r")) is None
         and sp.find(".//" + qn("a:endParaRPr")) is not None),
        None,
    )
    if element is None:
        pytest.skip("в шаблоне нет пустого плейсхолдера с endParaRPr")
    bind_text(element, [Paragraph("Итоги пилота")])

    run_pr = element.find(".//" + qn("a:r")).find(qn("a:rPr"))
    assert run_pr is None or run_pr.get("sz") is None


def test_cloned_headline_frame_is_narrowed_to_the_graphics_on_its_right(profile, deck):
    """Обложка VK Education (пример №1): плейсхолдер заголовка во всю
    ширину, справа половина слайда под графикой. Длинный заголовок заезжал
    на неё (прогон 26 сентября 2026); рамка сужается до графики."""
    prs, sources = deck
    pattern = _pattern(profile, "slide1")
    canvas = _canvas(profile)
    slide = clone_example_slide(prs, sources[1], builder._find_layout(prs, pattern.layout_id))
    matched = match_slots(slide, pattern.slots, canvas)
    headline_index = next(i for i, s in enumerate(pattern.slots) if s.role == "headline")
    ref = matched[headline_index]
    assert ref is not None
    graphics = [
        r for r in builder._frame_obstacles(slide, ref, [ref.element], canvas)
        if r.box.intersect(ref.box) is not None and not builder._contains(r.box, ref.box)
    ]
    graphics_left = min(r.box.left for r in graphics if r.box.left > ref.box.left + 0.3 * ref.box.width)
    assert ref.box.right > graphics_left, "у примера рамка заголовка заходит под графику: тест не о чем"

    shrunk = builder._shrink_frame_away_from_decor(slide, ref, [ref.element], canvas)

    assert shrunk.box.right <= graphics_left
    xfrm = shrunk.element.find(qn("p:spPr")).find(qn("a:xfrm"))
    assert int(xfrm.find(qn("a:ext")).get("cx")) == int(round(shrunk.box.width * canvas.width_emu))


# ---------------------------------------------------------------------------
# Родная таблица примера
# ---------------------------------------------------------------------------

_TABLE_ROWS = [
    ["Показатель", "До", "После", "Изменение"],
    ["Сквозная медиана", "31,5 ч", "6,2 ч", "−80%"],
    ["Доля переназначений вручную", "39%", "4%", "−35 п.п."],
    ["Заявок с просрочкой SLA", "23%", "6%", "−17 п.п."],
    ["Оценка удобства авторами (1–5)", "2,8", "4,1", "+1,3"],
    ["Обращений в поддержку", "120", "45", "−63%"],
]


def _table_frame(slide):
    return next(f for f in slide._element.iter(qn("p:graphicFrame")) if native_table(f) is not None)


def _table_texts(frame) -> list[list[str]]:
    return [
        ["".join(t.text or "" for t in tc.iter(qn("a:t"))) for tc in tr.findall(qn("a:tc"))]
        for tr in native_table(frame).findall(qn("a:tr"))
    ]


def _dark_fill(tc) -> bool:
    """Ячейка с заливкой accent1: так VK Education выделяет «Акцент» и «Итого»."""
    clr = tc.find(qn("a:tcPr") + "/" + qn("a:solidFill") + "/" + qn("a:schemeClr"))
    return clr is not None and clr.get("val") == "accent1"


def _table_spec(rows) -> SlideSpec:
    return SlideSpec(
        index=6, kind="table", headline="Пилот: −80% к сроку согласования",
        visual=Visual(kind="table", table=TableVisual(rows=rows)),
    )


def test_fill_native_table_matches_data_and_keeps_the_example_style(deck):
    """Пример №38: три столбца, шапка, строка-акцент на заливке accent1 и
    четыре обычные строки. Наша таблица 4×6: столбец добавлен, строка-акцент
    не стала второй шапкой, недостающая строка тела склонирована с
    последней строки тела, шапка осталась шапкой примера."""
    _prs, sources = deck
    frame = copy.deepcopy(_table_frame(sources[38]))
    tbl = native_table(frame)
    style_before = tbl.find(qn("a:tblPr") + "/" + qn("a:tableStyleId")).text
    header_pr_before = _xml(tbl.find(qn("a:tr")).findall(qn("a:tc"))[1].find(qn("a:tcPr")))

    fill_native_table(frame, _TABLE_ROWS, is_highlight=_dark_fill)

    assert _table_texts(frame) == _TABLE_ROWS
    assert len(tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))) == 4
    assert tbl.find(qn("a:tblPr") + "/" + qn("a:tableStyleId")).text == style_before
    header = tbl.find(qn("a:tr")).findall(qn("a:tc"))
    assert _xml(header[1].find(qn("a:tcPr"))) == header_pr_before
    body = tbl.findall(qn("a:tr"))[1:]
    assert not any(_dark_fill(tc) for tr in body for tc in tr.findall(qn("a:tc")))
    # Пустой угол шапки примера без своего кегля получил кегль соседки.
    assert header[0].find(".//" + qn("a:rPr")).get("sz") == header[1].find(".//" + qn("a:rPr")).get("sz")


def test_fill_native_table_drops_extra_rows_columns_and_merges(deck):
    """Пример №40: расписание 4×12 с объединёнными ячейками. Лишние строки
    и столбцы уходят, объединения примера к нашим данным не относятся."""
    _prs, sources = deck
    frame = copy.deepcopy(_table_frame(sources[40]))
    rows = [["Этап", "Срок"], ["Пилот", "июнь"], ["Раскатка", "сентябрь"]]

    fill_native_table(frame, rows, is_highlight=_dark_fill)

    tbl = native_table(frame)
    assert _table_texts(frame) == rows
    assert len(tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))) == 2
    for tc in tbl.iter(qn("a:tc")):
        assert not any(tc.get(a) for a in ("gridSpan", "rowSpan", "hMerge", "vMerge"))


def test_sole_table_clone_fills_the_native_table_wide_and_passes_audit(profile, deck):
    """Прогон 26 сентября 2026: на месте таблицы примера рисовалась своя,
    на четверть слайда. Теперь заполняется родная таблица примера;
    таблица-единственный блок не уже 0.6 холста, рамка равна сумме
    столбцов и строк, кегль не ниже caption, и клон проходит тот же аудит,
    что при сборке колоды."""
    prs, sources = deck
    pattern = _pattern(profile, "slide38")
    spec = _table_spec(_TABLE_ROWS[:5])

    outcome = builder.place_slide_by_clone(prs, spec, pattern, profile, AuditConfig.load(), sources[38])

    assert outcome.reason is None
    slide = prs.slides[-1]
    frames = [f for f in slide._element.iter(qn("p:graphicFrame")) if native_table(f) is not None]
    assert len(frames) == 1, "своя таблица поверх родной: рамок две"
    frame = frames[0]
    assert _table_texts(frame) == _TABLE_ROWS[:5]
    tbl = native_table(frame)
    widths = [int(gc.get("w")) for gc in tbl.iter(qn("a:gridCol"))]
    heights = [int(tr.get("h")) for tr in tbl.findall(qn("a:tr"))]
    ext = frame.find(qn("p:xfrm") + "/" + qn("a:ext"))
    assert int(ext.get("cx")) == sum(widths) and int(ext.get("cy")) == sum(heights)
    assert sum(widths) / profile.canvas_width_emu >= 0.6 - 1e-6
    caption = profile.type_scale_pt("caption")
    assert all(int(r.get("sz")) / 100 >= caption - 0.01 for r in tbl.iter(qn("a:rPr")) if r.get("sz"))
    errors = builder._clone_errors(
        builder.audit_slide_layout(slide, _canvas(profile), profile, AuditConfig.load(), index=spec.index),
    )
    assert not errors, errors


def test_lone_table_frame_starts_at_the_left_margin_when_the_left_column_is_gone(profile, deck):
    """VK Education slide38: список слева, таблица справа. Когда список
    пуст и удалён, таблица одна на слайде и начинается от левого поля, а
    не с середины (прогон 26 сентября 2026, слайд 7). Пока список на
    месте, таблица его не задевает."""
    prs, sources = deck
    pattern = _pattern(profile, "slide38")
    canvas = _canvas(profile)
    slide = clone_example_slide(prs, sources[38], builder._find_layout(prs, pattern.layout_id))
    matched = match_slots(slide, pattern.slots, canvas)
    table_slot = next(s for s in pattern.slots if s.role == "table")
    bullet_index = next(i for i, s in enumerate(pattern.slots) if s.role == "bullet")
    bullet_ref = matched[bullet_index]
    assert bullet_ref is not None
    table_ref = matched[next(i for i, s in enumerate(pattern.slots) if s.role == "table")]
    exclude = [table_ref.element] if table_ref is not None else []

    with_column = builder._table_frame_box(slide, table_slot.box, profile, canvas, exclude=exclude)
    assert with_column.left >= bullet_ref.box.right

    bullet_ref.element.getparent().remove(bullet_ref.element)
    alone = builder._table_frame_box(slide, table_slot.box, profile, canvas, exclude=exclude)
    assert abs(alone.left - profile.grid.margin_left) < 0.01
    assert alone.width >= builder._TABLE_MIN_WIDTH



# ---------------------------------------------------------------------------
# Привязка слотов к фигурам клона по id исходной фигуры
# ---------------------------------------------------------------------------


def _fresh_clone(profile, deck, pattern_id: str, number: int):
    prs, sources = deck
    pattern = _pattern(profile, pattern_id)
    layout = builder._find_layout(prs, pattern.layout_id)
    return pattern, clone_example_slide(prs, sources[number], layout)


def test_slot_finds_its_shape_by_id_when_the_box_drifted(profile, deck):
    """Коробка слота уехала ниже порога IoU (так бывает после притяжки к
    полям), а id у клона тот же, что у примера: фигура находится по id."""
    pattern, slide = _fresh_clone(profile, deck, "slide21", 21)
    canvas = _canvas(profile)
    baseline = match_slots(slide, pattern.slots, canvas)
    h = next(i for i, s in enumerate(pattern.slots) if s.role == "headline")
    assert baseline[h] is not None and pattern.slots[h].source_shape_id
    box = pattern.slots[h].box
    drifted = Box(box.left + box.width * 0.4, box.top, box.width, box.height)
    assert box_iou(drifted, baseline[h].box) < 0.85

    slots = list(pattern.slots)
    slots[h] = replace(slots[h], box=drifted)
    assert match_slots(slide, slots, canvas)[h].element is baseline[h].element

    slots[h] = replace(slots[h], source_shape_id=None)
    assert match_slots(slide, slots, canvas)[h] is None


def test_shape_of_a_foreign_kind_by_id_is_rejected_and_the_box_decides(profile, deck):
    """Id указал на картинку там, где слоту нужна надпись: такая фигура не
    принимается, и слот находит свою фигуру по коробке."""
    pattern, slide = _fresh_clone(profile, deck, "slide26", 26)
    canvas = _canvas(profile)
    baseline = match_slots(slide, pattern.slots, canvas)
    h = next(i for i, s in enumerate(pattern.slots) if s.role == "headline")
    picture = next(r for r in slide_refs(slide, canvas) if r.kind == "picture")

    slots = list(pattern.slots)
    slots[h] = replace(slots[h], source_shape_id=picture.shape_id)
    matched = match_slots(slide, slots, canvas)

    assert matched[h] is not None and matched[h].element is baseline[h].element
    assert matched[h].kind == "shape"


def test_prune_removes_unfilled_slot_and_decor_by_id(profile, deck):
    """Незаполненный слот и декор незаполненной единицы убираются по id,
    даже если их коробки в профиле не совпадают ни с одной фигурой клона."""
    pattern, slide = _fresh_clone(profile, deck, "slide26", 26)
    canvas = _canvas(profile)
    slot = next(s for s in pattern.slots if s.role == "card_body" and s.source_shape_id)
    decor = next(d for d in pattern.decor if d.source_shape_id)
    nowhere = Box(0.97, 0.97, 0.01, 0.01)
    before = clone_map(slide_refs(slide, canvas))
    assert slot.source_shape_id in before and decor.source_shape_id in before

    removed = prune_unfilled(
        slide, [replace(decor, box=nowhere)], [replace(slot, box=nowhere)], canvas,
    )

    after = clone_map(slide_refs(slide, canvas))
    assert removed >= 2
    assert slot.source_shape_id not in after and decor.source_shape_id not in after
    assert len(after) >= len(before) - removed


def test_doubled_shape_id_is_not_trusted(profile, deck):
    """Id, который на слайде встречается дважды, фигуру не узнаёт: такие
    id в карту клона не попадают, и слот идёт по коробке."""
    _pattern_, slide = _fresh_clone(profile, deck, "slide21", 21)
    canvas = _canvas(profile)
    refs = [r for r in slide_refs(slide, canvas) if r.shape_id]
    first, second = refs[0], refs[1]
    nv = second.element.find(".//" + qn("p:cNvPr"))
    nv.set("id", first.shape_id)

    ids = clone_map(slide_refs(slide, canvas))

    assert first.shape_id not in ids


def test_pattern_from_model_keeps_source_shape_ids(profile):
    for model in profile.patterns:
        pattern = builder._pattern_from_model(model)
        assert [s.source_shape_id for s in pattern.slots] == [s.source_shape_id for s in model.slots]
        assert [d.source_shape_id for d in pattern.decor] == [d.source_shape_id for d in model.decor]


def _divider(profile, pattern_id: str):
    """Разделитель VK Education, как его записал кеш v18/v20: подпись-
    подсказка «Точки используются для навигации» помечена моделью `fixed`."""
    pattern = _pattern(profile, pattern_id)
    slots = [replace(s, fixed=True) if s.role != "headline" else s for s in pattern.slots]
    return replace(pattern, slots=slots)


@pytest.mark.parametrize(("pattern_id", "number"), [("slide10", 10), ("slide11", 11)])
def test_divider_clone_drops_designer_hint_even_if_marked_fixed(profile, deck, pattern_id, number):
    """Скриншот 8.2: на разделителе остался текст шаблона «Точки
    используются для навигации». Флаг `fixed` от модели его больше не
    удерживает: подсказка не проходит фильтр постоянного текста."""
    prs, sources = deck
    pattern = _divider(profile, pattern_id)
    spec = SlideSpec(index=1, kind="section", headline="Дальше — контекст")

    outcome = builder.place_slide_by_clone(prs, spec, pattern, profile, AuditConfig.load(), sources[number])

    assert outcome.reason is None
    texts = _texts(prs.slides[-1])
    assert "Дальше — контекст" in texts
    assert not any("Точки" in t or "Пример" in t for t in texts), texts


def test_blank_headline_rejects_the_clone(profile, deck):
    """Пустая привязка не оставляет на слайде пустой плейсхолдер: клон
    отклоняется, и слайд собирается другим путём."""
    prs, sources = deck
    spec = SlideSpec(index=1, kind="section", headline="\u200b ")

    outcome = builder.place_slide_by_clone(
        prs, spec, _pattern(profile, "slide10"), profile, AuditConfig.load(), sources[10],
    )

    assert outcome.reason is not None and "пустой текст" in outcome.reason
    assert len(prs.slides) == 0


def _boxes(slide, canvas):
    return [ref for ref in walk_shapes(slide._element, canvas) if ref.box is not None]


def test_third_icon_of_an_unrecognised_unit_goes_with_its_column(profile, deck):
    """Наблюдение 8.4: у `slide26` майнинг собрал повтор из двух тезисов,
    третий тезис стал `body`, а его иконка получила `repeat_index=2` при
    `repeat.count=2`. При двух заполненных тезисах третья колонка уходит
    целиком: и иконка-декор, и кружок-картинка, и тезис примера."""
    prs, sources = deck
    pattern = _pattern(profile, "slide26")
    assert pattern.repeat.count == 2 and any(d.repeat_index >= pattern.repeat.count for d in pattern.decor)
    canvas = _canvas(profile)

    outcome = builder.place_slide_by_clone(prs, _cards_spec(2), pattern, profile, AuditConfig.load(), sources[26])

    assert outcome.reason is None
    third = [r for r in _boxes(prs.slides[-1], canvas) if 0.6 < r.box.left < 1.0 and r.box.top > 0.3]
    assert third == [], [(r.kind, r.box) for r in third]
    icons = [r for r in _boxes(prs.slides[-1], canvas) if r.kind == "picture"]
    assert len(icons) == 3  # иконка первой единицы, кружок и иконка второй


def test_unit_columns_of_filled_cards_survive_on_slide21(profile, deck):
    prs, sources = deck
    canvas = _canvas(profile)

    outcome = builder.place_slide_by_clone(
        prs, _cards_spec(3), _pattern(profile, "slide21"), profile, AuditConfig.load(), sources[21],
    )

    assert outcome.reason is None
    lefts = sorted({round(r.box.left, 2) for r in _boxes(prs.slides[-1], canvas) if r.box.top > 0.3})
    assert lefts == [0.05, 0.28, 0.5]


def test_example_content_photo_is_removed_when_the_slide_has_no_photo(profile, deck):
    """Контентная картинка примера (самая крупная, роль `image`) без нашего
    фото удаляется: скриншот поста «IT-дайвинг» с VK Education уезжал в
    чужую презентацию (27 сентября 2026). Иконки и декор остаются."""
    prs, sources = deck
    pattern = next(
        (p for p in (builder._pattern_from_model(m) for m in profile.patterns)
         if any(s.role == "image" for s in p.slots) and p.kind not in ("section", "image")),
        None,
    )
    if pattern is None:
        pytest.skip("в шаблоне нет содержательной раскладки со слотом под фото")
    spec = SlideSpec(
        index=1, kind=pattern.kind, headline="Заголовок без фотографии",
        blocks=[BulletBlock(items=["Первый тезис слайда", "Второй тезис слайда"])],
        pattern_id=pattern.pattern_id,
    )
    canvas = _canvas(profile)
    image_slot = next(s for s in pattern.slots if s.role == "image")
    outcome = builder.place_slide_by_clone(
        prs, spec, pattern, profile, AuditConfig.load(), sources[pattern.source_slide_index[0]], user_photos=None,
    )
    if outcome.reason is not None:
        pytest.skip(f"клон не принят: {outcome.reason}")
    slide = prs.slides[-1]
    leftover = [
        r for r in slide_refs(slide, canvas)
        if r.kind == "picture" and r.box is not None and builder._contains(image_slot.box, r.box, 0.02)
    ]
    assert not leftover, "картинка примера осталась в слоте под фото"

