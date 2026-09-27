"""Задача V5: предел места в контракте писателя считается тем же замером,
что проверка клона; обложечные раскладки не идут под содержание; сборка
с нуля не теряет блоки; запасная раскладка лестницы не повторяет соседа.
Модель не зовётся."""
from __future__ import annotations
import functools
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.compose import builder
from deckforge.compose.capacity import _FILLER, _TITLE_FILLER, measured_fits
from deckforge.compose.clone import sample_slides_by_number
from deckforge.ooxml.geometry import Canvas
from deckforge.pattern.candidates import candidates_for, cover_pattern_id, is_cover_like
from deckforge.pattern.forms import forms_of, pattern_form
from deckforge.pattern.intent import SlideIntent
from deckforge.pattern.scoring import look_key, overflow
from deckforge.pattern.style import load_style
from deckforge.plan.contracts import build_contract
from deckforge.plan.spec import BulletBlock, Card, CardBlock, SlideSpec
from deckforge.template.profile import TemplateProfile

TEMPLATE = Path("dataset/templates/Шаблон презентации VK Education.pptx")


@functools.lru_cache(maxsize=None)
def _profile() -> TemplateProfile:
    return TemplateProfile.from_file(TEMPLATE, cache_dir=None)


@pytest.fixture(scope="module")
def profile():
    return _profile()


def _words(n: int, max_chars: int) -> str:
    """Текст из `n` слов образца, не длиннее `max_chars` знаков."""
    words: list[str] = []
    for i in range(n):
        candidate = " ".join([*words, _FILLER[(i + 7) % len(_FILLER)]])
        if max_chars and len(candidate) > max_chars:
            break
        words.append(_FILLER[(i + 7) % len(_FILLER)])
    return " ".join(words)


def _assignment(pattern_id: str, kind: str, items: int):
    intent = SlideIntent(index=1, outline_kind="how_it_works", intent="Как работает", items=items)
    return type("A", (), {
        "intent": intent, "pattern_id": pattern_id, "kind": kind, "position": 1, "alternatives": (),
        "gap_note": lambda self: None,
    })()


def _clone(profile, spec: SlideSpec, pattern_id: str):
    prs = Presentation(str(TEMPLATE))
    sources = sample_slides_by_number(prs)
    builder._clear_sample_slides(prs)
    pattern = builder._pattern_from_model(next(m for m in profile.patterns if m.pattern_id == pattern_id))
    return builder.place_slide_by_clone(
        prs, spec, pattern, profile, AuditConfig.load(), sources[pattern.source_slide_index[0]],
    )


# ---------------------------------------------------------------------------
# Контракт против клона
# ---------------------------------------------------------------------------


def test_timeline_unit_limit_comes_from_the_clone_measure_not_the_area_estimate(profile):
    """slide42 (таймлайн из семи событий): рамка события 2,2x0,6" под 16 pt,
    а заголовок карточки клон ставит жирной первой строкой той же рамки.
    Площадная оценка профиля обещала телу 36 знаков, клон принимал около
    десяти: писатель укладывался в контракт, клон отказывал FONT_BUDGET."""
    model = next(m for m in profile.patterns if m.pattern_id == "slide42")
    area = pattern_form(model).main.unit
    measured = forms_of(profile)["slide42"].main.unit

    assert measured.max_chars < area.max_chars
    assert measured.native_words is not None and measured.native_words < measured.max_words


@pytest.mark.parametrize("pattern_id, items", [("slide42", 3), ("slide17", 4), ("slide21", 3)])
def test_text_at_the_contract_target_is_accepted_by_the_clone(profile, pattern_id, items):
    contract = build_contract(_assignment(pattern_id, "cards", items), profile)
    cards = contract.slots[0]
    assert cards.block == "cards"
    body = _words(cards.item.target_words, cards.item.max_chars)
    headline = _words(contract.headline.target_words, contract.headline.max_chars)
    spec = SlideSpec(
        index=1, kind="cards", headline=headline, pattern_id=pattern_id,
        blocks=[CardBlock(items=[Card(title=_TITLE_FILLER, body=body) for _ in range(cards.count)])],
    )

    outcome = _clone(profile, spec, pattern_id)

    assert outcome.reason is None, (outcome.code, outcome.reason, cards.item)


def test_measured_capacity_needs_the_template_file(profile):
    """Профиль без файла шаблона (собран в тесте, файл удалён): замера нет,
    форма берёт прежнюю оценку профиля, а не падает."""
    orphan = profile.model_copy(update={"source_path": "/нет/такого/шаблона.pptx"})
    assert measured_fits(orphan) == {}
    assert forms_of(orphan)["slide42"].main.unit.native_words is None


def test_planner_sees_a_layout_too_tight_for_the_style(profile):
    """Раскладка, где кеглем примера ложится меньше 60% слов стиля на
    единицу, это переполнение для планировщика."""
    intent = SlideIntent(index=1, outline_kind="how_it_works", intent="Шаги", items=3)
    forms = forms_of(profile)
    dense = load_style("dense")

    assert overflow(intent, forms["slide42"], dense) > 0.5
    assert overflow(intent, forms["slide21"], dense) == 0.0


# ---------------------------------------------------------------------------
# Обложечные раскладки
# ---------------------------------------------------------------------------


def test_cover_layouts_are_not_candidates_for_a_content_slide(profile):
    forms = forms_of(profile)
    covers = {p.pattern_id for p in profile.patterns if is_cover_like(p, profile)}
    assert {"slide1", "slide2", "slide4"} <= covers

    content = SlideIntent(index=3, outline_kind="risks", intent="Риски", items=3)
    found = candidates_for(content, profile, forms, position=3, last=10, cover_id=cover_pattern_id(profile))
    assert found.pattern_ids and not covers & set(found.pattern_ids), found

    title = SlideIntent(index=0, outline_kind="title", intent="Тема")
    first = candidates_for(title, profile, forms, position=0, last=10, cover_id=cover_pattern_id(profile))
    assert cover_pattern_id(profile) in first.pattern_ids


def test_ladder_does_not_fall_back_to_a_cover_under_content(profile):
    patterns = [builder._pattern_from_model(m) for m in profile.patterns]
    by_id = {p.pattern_id: p for p in patterns}
    spec = SlideSpec(index=3, kind="bullets", headline="Риски", pattern_id="slide13",
                     blocks=[BulletBlock(items=["Первый риск", "Второй риск"])])
    kept = builder._content_candidates(spec, [by_id["slide13"], by_id["slide2"], by_id["slide4"], by_id["slide53"]], profile)
    assert [p.pattern_id for p in kept] == ["slide13"]

    only_covers = builder._content_candidates(spec, [by_id["slide2"], by_id["slide4"]], profile)
    assert [p.pattern_id for p in only_covers] == ["slide2", "slide4"], "других нет: слайд всё равно собирается"


# ---------------------------------------------------------------------------
# Сборка с нуля не теряет блоки
# ---------------------------------------------------------------------------


def test_scratch_keeps_the_list_instead_of_a_headline_only_divider(profile):
    """airy 28 сентября 2026, «Риски и как их снимаем»: сборка с нуля взяла
    разделитель без места под список (одна находка D05) вместо раскладки
    со списком (две находки), и на слайде остался один заголовок."""
    prs = Presentation(str(TEMPLATE))
    builder._clear_sample_slides(prs)
    patterns = {m.pattern_id: builder._pattern_from_model(m) for m in profile.patterns}
    spec = SlideSpec(
        index=11, kind="section", headline="Риски и как их снимаем", pattern_id="slide16",
        blocks=[BulletBlock(items=[
            "Задержка кадровых данных до суток: оставляем ручную подмену.",
            "Четыре подразделения с особыми регламентами: включим во вторую очередь.",
        ])],
    )
    canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)
    notes: list[str] = []
    candidates = [patterns[pid] for pid in ("slide16", "slide11", "slide10", "slide13")]

    chosen = builder._scratch_candidates(
        prs, spec, candidates, profile, canvas, AuditConfig.load(), notes,
        bullet_char="•", user_photos=None, image_bytes=None,
    )

    grid = builder._grid_from_model(profile.grid)
    assert not builder._loses_all_content(spec, chosen, grid), (chosen.pattern_id, notes)
    texts = " ".join(sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame)
    assert "Задержка кадровых данных" in texts


# ---------------------------------------------------------------------------
# Разнообразие в лестнице
# ---------------------------------------------------------------------------


def _ladder(profile, candidates, avoid):
    return builder._Ladder(
        prs=None, slide_spec=SlideSpec(index=1, kind="cards", headline="x"), candidates=candidates,
        profile=profile, canvas=None, audit_config=None, forms={}, repair=None, room=0, kw={}, look=None,
        avoid_looks=frozenset(avoid),
    )


def test_ladder_spare_avoids_the_neighbours_look(profile):
    by_id = {m.pattern_id: builder._pattern_from_model(m) for m in profile.patterns}
    chosen, same, other = by_id["slide42"], by_id["slide21"], by_id["slide24"]
    pool = [chosen, same, by_id["slide44"], by_id["slide26"], other]

    ordered = _ladder(profile, pool, {look_key(same)})._diverse(pool[1:4])
    # С задачи T1 пример 44 разбор видит таймлайном, и его облик тоже не
    # соседский: первым встаёт любой из двух, но не облик соседа.
    assert ordered[0] in (by_id["slide26"], by_id["slide44"]), "сначала облик, которого нет у соседа"
    assert look_key(ordered[0]) != look_key(same)

    ladder = _ladder(profile, pool, {look_key(same), look_key(by_id["slide26"]), look_key(by_id["slide44"])})
    assert other in ladder.spares(), "все запасные повторяют соседей: добавлен кандидат другого облика"


def test_crowded_looks_cover_neighbours_and_the_repeat_limit(profile):
    by_id = {m.pattern_id: builder._pattern_from_model(m) for m in profile.patterns}
    slides = [SlideSpec(index=i, kind="cards", headline="x", pattern_id=pid)
              for i, pid in enumerate(["slide21", "slide13", "slide24", "slide26"])]
    placed = [look_key(by_id["slide13"])] * 3 + [look_key(by_id["slide21"])]
    diversity = load_style("dense").diversity

    crowded = builder._crowded_looks(slides, 2, placed, by_id, diversity)

    assert look_key(by_id["slide21"]) in crowded, "сосед слева"
    assert look_key(by_id["slide26"]) in crowded, "сосед справа"
    assert look_key(by_id["slide13"]) in crowded, "предел повторов исчерпан"
    assert look_key(by_id["slide24"]) not in crowded


def test_a_list_gets_no_more_items_than_the_place_has_lines(profile):
    """Пример 13: главное текстовое место это полоса во всю ширину в две
    строки. Одной фразой она держит три десятка слов, а списком только два
    пункта: пункт это абзац не короче строки. Живой прогон dense 28
    сентября 2026: три коротких пункта клон отклонил, слайд разделили на
    два пустоватых."""
    main = forms_of(profile)["slide13"].main
    assert main.block in ("bullets", "text")
    assert main.unit.lines is not None and main.units <= main.unit.lines

    spec = SlideSpec(
        index=2, kind="bullets", headline="98,5% времени ожидание", pattern_id="slide13",
        blocks=[BulletBlock(items=["Сквозная медиана: 31,5 часа", "Чистая работа: 28 минут"][: main.units])],
    )
    assert _clone(profile, spec, "slide13").reason is None


def test_relaxing_the_form_does_not_bring_covers_back(profile):
    """Живой прогон visual 29 сентября 2026: пункт с одним показателем не
    нашёл раскладки нужной формы, ограничение формы ослабили, и вместе с
    ним вернулась обложка примера 2. Запрет обложки ослабляется последним."""
    forms = forms_of(profile)
    covers = {p.pattern_id for p in profile.patterns if is_cover_like(p, profile)}
    kpi = SlideIntent(index=1, outline_kind="context", intent="Главная цифра", items=1, form="kpi")

    found = candidates_for(kpi, profile, forms, position=1, last=11, cover_id=cover_pattern_id(profile))

    assert found.pattern_ids and not covers & set(found.pattern_ids), found
