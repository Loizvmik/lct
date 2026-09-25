"""Тесты `compose/tables.py` (Task 10) — brief дословно по составу
случаев: кегль ужимается циклом, но не ниже пола; высоты строк — честный
замер, а не равное деление; цвет текста шапки выбирается по контрасту;
таблица строится и без стиля таблицы в самом файле (VK Tech).

Task 10 код-ревью — дописана регрессия на находку №3 отчёта задачи (тело
таблицы — заливка того же полюса, что фактический фон слайда), параметри-
зованная по всем трём учебным шаблонам (находка №3 код-ревью)."""
from __future__ import annotations

from pathlib import Path

import pytest

from deckforge.compose.tables import FONT_FLOOR_RATIO, TableSpec, add_table, column_shares, header_fill
from deckforge.template.naming import contrast_ratio

# Дубликат `TEMPLATE_NAMES` из `tests/compose/conftest.py` — намеренно, не
# импорт (см. докстроку conftest.py про независимость тестовых модулей).
_TEMPLATE_NAMES = [
    "VK Tech шаблон.pptx",
    "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
    "Шаблон презентации VK Education.pptx",
]


def _relative_luminance(hex_color: str) -> float:
    h = hex_color.lstrip("#")

    def lin(c: float) -> float:
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def _poles(profile) -> dict:
    poles: dict = {}
    for layout in profile.layouts:
        if layout.background.luminance is None:
            continue
        pole = "dark" if layout.is_dark else "light"
        poles.setdefault(pole, layout)
    return poles


def _body_cell_fill_hex(table) -> str:
    return "#" + str(table.cell(1, 0).fill.fore_color.rgb)


def _cell_sizes(table) -> list[float]:
    sizes = []
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        sizes.append(run.font.size.pt)
    return sizes


def _header_text_color(table) -> str:
    run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    return f"#{str(run.font.color.rgb)}"


def _contrast(hex_a: str, hex_b: str) -> float:
    return contrast_ratio(hex_a, hex_b)


def test_table_shrinks_font_until_it_fits_but_not_below_floor(new_slide, PROFILE, SMALL_BOX):
    frame = add_table(
        new_slide(), SMALL_BOX,
        TableSpec(
            header=["Метрика", "Было", "Стало"],
            rows=[["Медианное время ожидания заявки в очереди", "4 ч", "1 ч"]] * 6,
        ),
        PROFILE,
    )
    sizes = _cell_sizes(frame.table)
    assert sizes
    # `type_scale.steps` нормирован к эталонному холсту 13.333″ — пол цикла
    # ужимания (`tables._fit_font_size`) денормирован к РЕАЛЬНОМУ холсту
    # `PROFILE` (`TemplateProfile.type_scale_pt`, см. находку аудита T02,
    # отчёт Task 10a), сырой `type_scale.steps["caption"]` — кегль ДРУГОГО,
    # эталонного холста, сравнивать построенную таблицу с ним напрямую
    # неверно на любом шаблоне с холстом, отличным от эталонного.
    # 0.01pt слабины — квантование `python-pptx` при записи `sz` в XML
    # (сотые доли пункта), не сама проверка: кегль, ЗАЖАТЫЙ ровно на полу
    # (`max(size * _SHRINK_FACTOR, floor)`), при обратном чтении из XML
    # может round-trip'нуться на пару сотых пункта ниже `floor` — тот же
    # источник шума, что и `config/audit.yaml: template.size_tolerance_pt`.
    assert min(sizes) >= PROFILE.type_scale_pt("caption") * FONT_FLOOR_RATIO - 0.01


def test_row_heights_are_measured_not_split_evenly(new_slide, PROFILE, BOX):
    """Высота строки в .pptx — минимум, а не размер: редактор растит строку
    под текст, но не ужимает, и таблица тихо уезжает на подвал."""
    frame = add_table(
        new_slide(), BOX,
        TableSpec(header=["A", "B"], rows=[["коротко", "коротко"], ["очень длинный текст " * 6, "x"]]),
        PROFILE,
    )
    heights = [row.height for row in frame.table.rows]
    assert heights[2] > heights[1]


def test_header_text_color_is_chosen_by_contrast(new_slide, PROFILE, BOX):
    frame = add_table(new_slide(), BOX, TableSpec(header=["A", "B"], rows=[["1", "2"]]), PROFILE)
    assert _contrast(_header_text_color(frame.table), header_fill(PROFILE)) >= 4.5


def test_table_without_template_table_style_still_builds(new_slide, profile_fixture, BOX):
    """У VK Tech tableStyles.xml отсутствует вовсе."""
    profile = profile_fixture("VK Tech шаблон.pptx")
    frame = add_table(new_slide(), BOX, TableSpec(header=["A", "B"], rows=[["1", "2"]]), profile)
    assert frame.has_table


def test_column_alignment_is_applied(new_slide, PROFILE, BOX):
    frame = add_table(
        new_slide(), BOX,
        TableSpec(header=["Метрика", "Значение"], rows=[["Строк", "10"]], align=["l", "r"]),
        PROFILE,
    )
    from pptx.enum.text import PP_ALIGN
    assert frame.table.cell(0, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT


@pytest.mark.parametrize("template_name", _TEMPLATE_NAMES)
def test_table_body_fill_matches_the_actual_slide_background_pole(
    real_slide_factory, profile_fixture, template_name, BOX,
):
    """Task 10 отчёт, находка №3: тело таблицы — заливка ТОГО ЖЕ полюса
    светлый/тёмный, что фактический фон слайда под ней, а не слепая
    `palette_roles["surface"]` (у VK Tech/WorkSpace `surface` чёрный —
    большая часть шаблона тёмная, но демо-раскладка задачи светлая; таблица
    чёрной плашкой на бледном фоне выглядела вырезанной из другого файла).
    Проверяем на КАЖДОМ полюсе фона, который реально есть у шаблона
    (WorkSpace — только тёмный, см. `_poles`): заливка тела ближе к
    фактическому фону слайда, чем противоположный полюс surface/on_surface."""
    profile = profile_fixture(template_name)
    poles = _poles(profile)
    assert poles, template_name

    surface_lum = _relative_luminance(profile.palette_roles["surface"])
    on_surface_lum = _relative_luminance(profile.palette_roles["on_surface"])

    for pole_name, layout in poles.items():
        slide = real_slide_factory(template_name, layout.part_name)
        frame = add_table(slide, BOX, TableSpec(header=["A", "B"], rows=[["1", "2"]]), profile)
        body_lum = _relative_luminance(_body_cell_fill_hex(frame.table))
        bg_lum = layout.background.luminance

        other_lum = on_surface_lum if abs(body_lum - surface_lum) < abs(body_lum - on_surface_lum) else surface_lum
        assert abs(body_lum - bg_lum) <= abs(other_lum - bg_lum), (template_name, pole_name)


def test_header_fill_is_theme_accent1_not_a_model_assigned_role(new_slide, profile_fixture, BOX):
    """Прогон 26 сентября 2026: шапка на VK Education вышла розовой, потому
    что модель назначила роль `accent` цвету accent2 темы (#FF3885). Шапка
    красится accent1 темы, цветом, которым шаблон сам заливает таблицы."""
    profile = profile_fixture("Шаблон презентации VK Education.pptx")
    pink = profile.model_copy(update={"palette_roles": {**profile.palette_roles, "accent": "#FF3885"}})
    frame = add_table(new_slide(), BOX, TableSpec(header=["A", "B"], rows=[["1", "2"]]), pink)
    assert str(frame.table.cell(0, 0).fill.fore_color.rgb) == "0077FF"


def test_column_shares_give_text_columns_more_room_than_numbers():
    rows = [
        ["Показатель", "До", "После"],
        ["Доля переназначений вручную", "39%", "4%"],
        ["Сквозная медиана", "31,5 ч", "6,2 ч"],
    ]
    shares = column_shares(rows)
    assert abs(sum(shares) - 1) < 1e-9
    assert shares[0] > 2 * shares[1]


def test_frame_height_is_the_sum_of_measured_rows(new_slide, PROFILE, BOX):
    """Рамка по строкам, а не по коробке слота: иначе аудит видит площадь,
    которой на слайде нет."""
    frame = add_table(new_slide(), BOX, TableSpec(header=["A", "B"], rows=[["1", "2"]]), PROFILE)
    assert frame.height == sum(row.height for row in frame.table.rows)
    assert frame.height < round(BOX.height * PROFILE.canvas_height_emu)


def _sole_table_spec() -> "SlideSpec":
    from deckforge.plan.spec import SlideSpec, TableVisual, Visual
    rows = [
        ["Показатель", "До", "После", "Изменение"],
        ["Сквозная медиана", "31,5 ч", "6,2 ч", "−80%"],
        ["Доля переназначений вручную", "39%", "4%", "−35 п.п."],
        ["Заявок с просрочкой SLA", "23%", "6%", "−17 п.п."],
        ["Оценка удобства авторами (1–5)", "2,8", "4,1", "+1,3"],
    ]
    return SlideSpec(
        index=6, kind="table", headline="Пилот: −80% к сроку согласования",
        visual=Visual(kind="table", table=TableVisual(rows=rows)),
    )


def test_sole_table_built_from_scratch_spans_the_content_width(profile_fixture):
    """Прогон 26 сентября 2026: таблица единственным блоком слайда занимала
    четверть ширины (слот снят с рамки примера 3 000 000 EMU), слова в
    ячейках рвались по буквам. Такая таблица растягивается до правого поля
    и не уже 0.6 холста, кегль тела не ниже caption, столбцы не срезаны."""
    from pptx import Presentation

    from deckforge.audit.config import AuditConfig
    from deckforge.compose import builder

    name = "Шаблон презентации VK Education.pptx"
    profile = profile_fixture(name)
    prs = Presentation(str(Path("dataset/templates") / name))
    builder._clear_sample_slides(prs)
    pattern = next(builder._pattern_from_model(m) for m in profile.patterns if m.pattern_id == "slide38")
    spec = _sole_table_spec()

    builder.place_slide(prs, spec, pattern, profile, AuditConfig.load())

    frames = [s for s in prs.slides[-1].shapes if s.has_table]
    assert len(frames) == 1
    frame = frames[0]
    width = frame.width / profile.canvas_width_emu
    assert width >= 0.6 - 1e-6
    assert (frame.left + frame.width) / profile.canvas_width_emu <= 1 - profile.grid.margin_right + 1e-6
    assert len(frame.table.columns) == 4
    assert min(_cell_sizes(frame.table)) >= profile.type_scale_pt("caption") - 0.01
