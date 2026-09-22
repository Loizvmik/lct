"""Общая инфраструктура тестов Task 10 (`test_charts.py`/`test_tables.py`/
`test_diagrams.py`) — профиль строится на РЕАЛЬНЫХ файлах `dataset/
templates/` (не синтетика: графики/таблицы/схемы читают
`profile.chart_series`/`shape_vocabulary`/`type_scale`/`assets.icons`,
которые имеют смысл только на настоящем разборе), кешируется на весь
прогон (`functools.lru_cache`), тот же приём, что и `tests/template/
conftest.py`.
"""
from __future__ import annotations
import functools
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.ooxml.geometry import Box
from deckforge.template.profile import TemplateProfile

TEMPLATES_DIR = Path("dataset/templates")

# ЛЦТ2026 — контрольный шаблон, в тестах не участвует (тот же принцип, что
# и во всех остальных tests/*/conftest.py проекта).
DEFAULT_TEMPLATE = "VK Tech шаблон.pptx"

# Все три учебных шаблона — Task 10 код-ревью, находка №3: графики/таблицы/
# схемы гоняются по умолчанию ТОЛЬКО на `DEFAULT_TEMPLATE` (VK Tech), на
# котором случайно нет ни коллизии цветов ряда, ни проблемы словаря форм —
# регрессия на шести остальных найденных находках (текст графика/сектора
# pie/тело таблицы/карточка схемы/подложка пиктограммы/порядок иконок)
# никем не замечалась бы. Каждый файл, дублирующий этот список локально
# (`_TEMPLATE_NAMES` в `test_charts.py`/`test_tables.py`/`test_diagrams.py`
# — намеренная копия, не импорт: тесты, в отличие от `template/`, обязаны
# оставаться независимыми друг от друга и от внутренних приватных имён
# соседних тестовых модулей), обязан прогонять свои дефекто-специфичные
# тесты по всем трём.
TEMPLATE_NAMES = [
    "VK Tech шаблон.pptx",
    "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
    "Шаблон презентации VK Education.pptx",
]


@functools.lru_cache(maxsize=None)
def _profile(name: str) -> TemplateProfile:
    return TemplateProfile.from_file(TEMPLATES_DIR / name, cache_dir=None)


@pytest.fixture(scope="session")
def profile_fixture():
    """`profile_fixture(name) -> TemplateProfile` — вызываемая фикстура,
    как в `tests/template/conftest.py`."""
    return _profile


@pytest.fixture(scope="session")
def PROFILE(profile_fixture) -> TemplateProfile:
    return profile_fixture(DEFAULT_TEMPLATE)


@pytest.fixture
def new_slide():
    """`new_slide()` — чистый слайд на ПУСТОМ шаблоне `python-pptx` (не на
    файле шаблона DeckForge): графикам/таблицам/схемам из `TemplateProfile`
    нужны только данные профиля (палитра/шкала/словарь форм), не реальные
    лейауты конкретного .pptx — развязка позволяет тестам не зависеть от
    того, что несёт лейаут №6 в каждом из четырёх файлов датасета."""
    def _make():
        prs = Presentation()
        return prs.slides.add_slide(prs.slide_layouts[6])
    return _make


@pytest.fixture
def BOX() -> Box:
    return Box(left=0.1, top=0.15, width=0.6, height=0.5)


@pytest.fixture
def SMALL_BOX() -> Box:
    return Box(left=0.1, top=0.15, width=0.2, height=0.12)


@pytest.fixture
def real_slide_factory():
    """`real_slide_factory(template_name, layout_part_name) -> Slide` —
    слайд на РЕАЛЬНОМ лейауте настоящего файла шаблона (не `new_slide()`,
    пустой `Presentation()` без единого макета шаблона) — нужен
    регрессионным тестам находок отчёта Task 10 №1/№3/№4 (цвет текста
    графика/заливка тела таблицы/заливка карточки схемы, все три — по
    ФАКТИЧЕСКОМУ фону макета), потому что этот фон резолвится
    `colorpick.slide_background_luminance` через `slide.slide_layout.part.
    partname` — атрибут, которого у слайда на синтетическом пустом
    `Presentation()` попросту нет так, чтобы он что-то значил (совпадений
    с `profile.layouts` там быть не может по построению)."""
    def _make(template_name: str, layout_part_name: str):
        prs = Presentation(str(TEMPLATES_DIR / template_name))
        for master in prs.slide_masters:
            for layout in master.slide_layouts:
                if str(layout.part.partname).lstrip("/") == layout_part_name:
                    return prs.slides.add_slide(layout)
        raise AssertionError(f"лейаут {layout_part_name!r} не найден в {template_name!r}")
    return _make
