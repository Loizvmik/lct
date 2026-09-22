"""Тесты `compose/charts.py` (Task 10) — brief дословно по составу
случаев: все восемь типов графика строятся, цвета рядов/точек берутся из
`profile.chart_series` (не из темы Office), у графика с одним рядом
красятся ТОЧКИ, а не ряд, легенда появляется только когда несёт смысл,
подписи осей+единиц обязательны, отрицательные значения уводят подписи
категорий вниз.

Task 10 код-ревью — дописаны регрессии на находки визуального ревью отчёта
задачи (№2 критично: попарная различимость цветов ряда; №1: текст графика
по фактическому фону макета; №2 отчёта: секторы pie/doughnut все разного
цвета), параметризованные по всем трём учебным шаблонам (находка №3
код-ревью: раньше всё гонялось на одном VK Tech)."""
from __future__ import annotations

import pytest
from pptx.enum.chart import XL_TICK_LABEL_POSITION

from deckforge.compose.charts import CHART_KINDS, ChartSpec, Series, add_chart
from deckforge.ooxml.color import delta_e76
from deckforge.template.chart_palette import MIN_DELTA_E
from deckforge.template.naming import MIN_CONTRAST

# Дубликат `TEMPLATE_NAMES` из `tests/compose/conftest.py` — намеренно, не
# импорт (см. докстроку conftest.py про независимость тестовых модулей).
_TEMPLATE_NAMES = [
    "VK Tech шаблон.pptx",
    "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
    "Шаблон презентации VK Education.pptx",
]


def _series_colors(chart) -> list[str]:
    return [str(series.format.fill.fore_color.rgb) for series in chart.plots[0].series]


def _point_colors(series) -> list[str]:
    return [str(point.format.fill.fore_color.rgb) for point in series.points]


def _value_axis_title(chart) -> str:
    return chart.value_axis.axis_title.text_frame.text


def _tick_label_position(axis) -> str:
    return axis.tick_label_position.name


def _relative_luminance(hex_color: str) -> float:
    h = hex_color.lstrip("#")

    def lin(c: float) -> float:
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def _contrast(l_a: float, l_b: float) -> float:
    lighter, darker = max(l_a, l_b), min(l_a, l_b)
    return (lighter + 0.05) / (darker + 0.05)


def _poles(profile) -> dict:
    poles: dict = {}
    for layout in profile.layouts:
        if layout.background.luminance is None:
            continue
        pole = "dark" if layout.is_dark else "light"
        poles.setdefault(pole, layout)
    return poles


def test_every_chart_kind_builds(new_slide, PROFILE, BOX):
    assert set(CHART_KINDS) == {
        "bar", "bar_stacked", "bar_h", "line", "area", "pie", "doughnut", "scatter",
    }
    for kind in CHART_KINDS:
        frame = add_chart(
            new_slide(), BOX,
            ChartSpec(kind=kind, categories=["A", "B", "C"], series=[Series("Выручка", [1, 2, 3])], unit="млн ₽"),
            PROFILE,
        )
        assert frame.has_chart, kind


def test_series_colors_come_from_the_template_palette(new_slide, PROFILE, BOX):
    """Без явной окраски python-pptx отдаёт раскраску теме файла, и график
    выходит в цветах Office."""
    frame = add_chart(
        new_slide(), BOX,
        ChartSpec(kind="bar", categories=["A", "B"], series=[Series("s1", [1, 2]), Series("s2", [2, 1])]),
        PROFILE,
    )
    used = _series_colors(frame.chart)
    allowed = {c.lstrip("#").upper() for c in PROFILE.chart_series}
    assert set(used) <= allowed
    assert used  # хоть что-то покрашено


def test_single_series_colors_points_not_series(new_slide, PROFILE, BOX):
    """У одного ряда окраска ряда даёт одинаковые столбцы; красить надо точки."""
    frame = add_chart(
        new_slide(), BOX,
        ChartSpec(kind="bar", categories=["A", "B", "C"], series=[Series("s", [1, 2, 3])], highlight_index=1),
        PROFILE,
    )
    colors = _point_colors(frame.chart.plots[0].series[0])
    assert len(colors) == 3
    assert len(set(colors)) == 2


def test_chart_has_axis_titles_and_units(new_slide, PROFILE, BOX):
    """Проверка аудита I05: у диаграммы нет подписей осей, единиц или легенды."""
    frame = add_chart(
        new_slide(), BOX,
        ChartSpec(
            kind="line", categories=["Янв", "Фев"], series=[Series("Заявки", [10, 12])],
            unit="штук", axis_titles=("Месяц", "Заявки"),
        ),
        PROFILE,
    )
    assert _value_axis_title(frame.chart) == "Заявки, штук"
    assert frame.chart.category_axis.axis_title.text_frame.text == "Месяц"


def test_legend_appears_only_when_it_carries_information(new_slide, PROFILE, BOX):
    one = add_chart(new_slide(), BOX, ChartSpec(kind="bar", categories=["A"], series=[Series("s", [1])]), PROFILE)
    two = add_chart(
        new_slide(), BOX,
        ChartSpec(kind="bar", categories=["A"], series=[Series("s1", [1]), Series("s2", [2])]),
        PROFILE,
    )
    assert one.chart.has_legend is False
    assert two.chart.has_legend is True


def test_pie_legend_reflects_categories_not_series(new_slide, PROFILE, BOX):
    """У круговой диаграммы ряд всегда один по природе формата — легенду
    включает число КАТЕГОРИЙ (секторов), а не рядов."""
    single_category = add_chart(
        new_slide(), BOX, ChartSpec(kind="pie", categories=["Всё"], series=[Series("s", [1])]), PROFILE,
    )
    many_categories = add_chart(
        new_slide(), BOX,
        ChartSpec(kind="pie", categories=["A", "B", "C"], series=[Series("s", [1, 2, 3])]),
        PROFILE,
    )
    assert single_category.chart.has_legend is False
    assert many_categories.chart.has_legend is True


def test_negative_values_move_category_labels_low(new_slide, PROFILE, BOX):
    frame = add_chart(
        new_slide(), BOX, ChartSpec(kind="bar", categories=["A", "B"], series=[Series("s", [-5, 7])]), PROFILE,
    )
    assert _tick_label_position(frame.chart.category_axis) == "LOW"


def test_positive_values_keep_default_tick_label_position(new_slide, PROFILE, BOX):
    frame = add_chart(
        new_slide(), BOX, ChartSpec(kind="bar", categories=["A", "B"], series=[Series("s", [5, 7])]), PROFILE,
    )
    assert frame.chart.category_axis.tick_label_position != XL_TICK_LABEL_POSITION.LOW


def test_unknown_chart_kind_raises(new_slide, PROFILE, BOX):
    with pytest.raises(ValueError):
        add_chart(new_slide(), BOX, ChartSpec(kind="bogus", categories=["A"], series=[Series("s", [1])]), PROFILE)


@pytest.mark.parametrize("template_name", _TEMPLATE_NAMES)
def test_series_colors_are_pairwise_distinguishable(profile_fixture, template_name):
    """Task 10 код-ревью, находка №2 (критично): отбор по весу/насыщенности
    не проверял попарную различимость вовсе — контрольный шаблон ЛЦТ2026
    (вне датасета тестов) даёт `#FE095F`/`#FF0053` в одной палитре, ΔE76
    ≈ 7.0, ниже уверенного порога, кольцевая диаграмма сливает секторы.
    Палитра рядов обязана быть различима ПОПАРНО в CIE Lab, не только
    валидной по составу (старый тест, `test_series_colors_come_from_the_
    template_palette`, проверял только это)."""
    profile = profile_fixture(template_name)
    series = profile.chart_series
    assert len(series) >= 2, (template_name, series)
    for i in range(len(series)):
        for j in range(i + 1, len(series)):
            distance = delta_e76(series[i], series[j])
            assert distance >= MIN_DELTA_E, (template_name, series[i], series[j], distance)


@pytest.mark.parametrize("template_name", _TEMPLATE_NAMES)
def test_chart_text_color_matches_the_actual_slide_background(
    real_slide_factory, profile_fixture, template_name, BOX,
):
    """Task 10 отчёт, находка №1 (критично): цвет текста графика раньше был
    фиксированной ролью `on_surface`, читаемой только если фактический фон
    МАКЕТА совпадает с общей парой `surface`/`on_surface` (VK Tech, лейаут
    "free" — декоративный фон не совпадал, заголовок и подписи осей вышли
    почти невидимыми). Проверяем контраст к ФАКТИЧЕСКОМУ фону КАЖДОГО
    полюса шаблона (не к роли), тем же кодом, что реально строит график."""
    profile = profile_fixture(template_name)
    poles = _poles(profile)
    assert poles, template_name

    for pole_name, layout in poles.items():
        slide = real_slide_factory(template_name, layout.part_name)
        frame = add_chart(
            slide, BOX, ChartSpec(kind="bar", categories=["A", "B"], series=[Series("s", [1, 2])]), profile,
        )
        text_hex = "#" + str(frame.chart.font.color.rgb)
        contrast = _contrast(_relative_luminance(text_hex), layout.background.luminance)
        assert contrast >= MIN_CONTRAST, (template_name, pole_name, contrast)


@pytest.mark.parametrize("kind", ["pie", "doughnut"])
@pytest.mark.parametrize("template_name", _TEMPLATE_NAMES)
def test_pie_and_doughnut_sectors_are_all_distinct_colors(new_slide, profile_fixture, BOX, template_name, kind):
    """Task 10 отчёт, находка №2: у pie/doughnut каждая ТОЧКА — своя
    категория, а не повтор ряда — красить надо циклом палитры по точкам, не
    "база+акцент" (тот приём сливал три из четырёх секторов в один цвет на
    живом рендере VK Tech). Раньше регрессия проверялась только у
    `test_single_series_colors_points_not_series` (bar, не pie/doughnut) —
    параметризовано по трём учебным шаблонам."""
    profile = profile_fixture(template_name)
    frame = add_chart(
        new_slide(), BOX,
        ChartSpec(
            kind=kind, categories=["Янв", "Фев", "Мар", "Апр"],
            series=[Series("s", [10, 20, 30, 40])], highlight_index=1,
        ),
        profile,
    )
    colors = _point_colors(frame.chart.plots[0].series[0])
    assert len(colors) == 4
    assert len(set(colors)) == 4, (template_name, kind, colors)
