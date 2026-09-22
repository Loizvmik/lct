"""Тесты вставки пользовательских фотографий контент-пакета (Task 20) —
`compose.builder._place_picture_visual`/`_contain_box` с `user_photos`.

Большинство тестов зовут `_place_picture_visual` напрямую на СИНТЕТИЧЕСКОМ
`Pattern` (тот же приём, что и остальной `test_builder.py`: `_fake_cards_
pattern`/ручные `PatternSlot`), а не через `build_deck`/выбор раскладки —
какой ИМЕННО паттерн реального шаблона достанется слайду, зависит от
ранжира (`_pattern_rank_key`) и не детерминировано с точки зрения ЭТОГО
теста (например, VK Tech внешне "не несёт" `kind=photo_text/image`, но
несколько его `bullets`/`cards`/`section`-раскладок ФИЗИЧЕСКИ содержат
слот `image` — задача этого файла проверить поведение `_place_picture_
visual` при слоте и без него, не поведение подборщика раскладки, это
покрыто отдельно в `test_builder.py`/`test_variants.py`).

Один интеграционный тест (`test_user_photo_flows_through_build_deck_end_
to_end`) идёт через настоящий `build_deck` на реальном шаблоне, чтобы
проверить, что `user_photos` действительно доходит от `cli.py`-подобного
вызова до готового `.pptx` — по одному разу достаточно, детали "contain"/
находок покрыты юнит-тестами ниже."""
from __future__ import annotations
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckforge.compose.builder import Variant, _place_picture_visual, build_deck
from deckforge.ooxml.geometry import Box
from deckforge.plan.spec import DeckSpec, SlideSpec, Visual
from deckforge.template.patterns import Capacity, Pattern, PatternSlot
from deckforge.template.profile import TemplateProfile

EDU_TEMPLATE = Path("dataset/templates/Шаблон презентации VK Education.pptx")

_EMPTY_CAPACITY = Capacity(
    max_items=0, max_chars_per_item=0, max_bullets=0, max_series=0, max_rows=0, max_cols=0,
)


def _make_photo(tmp_path: Path, name: str, size: tuple[int, int]) -> Path:
    path = tmp_path / name
    Image.new("RGB", size, color=(10, 20, 30)).save(path)
    return path


def _picture_shapes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _pattern_with_image_slot(box: Box) -> Pattern:
    slot = PatternSlot(role="image", box=box, size_pt=0.0, color_hex=None, align="l", max_chars=0, wraps=False)
    return Pattern(
        pattern_id="with-image", source_slide_index=[0], layout_id="slideLayout1",
        kind="photo_text", slots=[slot], repeat=None, decor=[], capacity=_EMPTY_CAPACITY,
        score=1.0, is_dark=False,
    )


def _pattern_without_image_slot() -> Pattern:
    slot = PatternSlot(
        role="body", box=Box(left=0.1, top=0.1, width=0.5, height=0.3), size_pt=16.0,
        color_hex=None, align="l", max_chars=200, wraps=True,
    )
    return Pattern(
        pattern_id="no-image", source_slide_index=[0], layout_id="slideLayout1",
        kind="bullets", slots=[slot], repeat=None, decor=[], capacity=_EMPTY_CAPACITY,
        score=1.0, is_dark=False,
    )


@pytest.fixture(scope="module")
def profile() -> TemplateProfile:
    return TemplateProfile.from_file(EDU_TEMPLATE, cache_dir=None)


@pytest.mark.parametrize(
    "slot_box,native_size",
    [
        (Box(left=0.1, top=0.1, width=0.6, height=0.2), (300, 900)),  # слот шире, фото уже — height-constrained
        (Box(left=0.1, top=0.1, width=0.2, height=0.6), (1600, 300)),  # слот выше, фото шире — width-constrained
    ],
    ids=["slot-wide-photo-tall", "slot-tall-photo-wide"],
)
def test_user_photo_is_contained_not_stretched(new_slide, profile, tmp_path, slot_box, native_size):
    """L07 (аудит: искажённые пропорции) — вписанная фотография несёт то
    же соотношение сторон, что исходный файл, независимо от формы слота."""
    slide = new_slide()
    pattern = _pattern_with_image_slot(slot_box)
    slide_spec = SlideSpec(
        index=0, kind="photo_text", headline="Демо",
        visual=Visual(kind="photo", photo_name="demo.jpg"),
    )
    photo_path = _make_photo(tmp_path, "demo.jpg", native_size)

    _place_picture_visual(slide, slide_spec, pattern, profile, "photo", {"demo.jpg": photo_path})

    pics = _picture_shapes(slide)
    assert len(pics) == 1
    pic = pics[0]
    native_aspect = native_size[0] / native_size[1]
    assert pic.width / pic.height == pytest.approx(native_aspect, rel=0.02)

    # "Contain": ни одна сторона картинки не может выйти за слот.
    slot_left = round(slot_box.left * profile.canvas_width_emu)
    slot_top = round(slot_box.top * profile.canvas_height_emu)
    slot_width = round(slot_box.width * profile.canvas_width_emu)
    slot_height = round(slot_box.height * profile.canvas_height_emu)
    assert pic.left >= slot_left - 1
    assert pic.top >= slot_top - 1
    assert pic.left + pic.width <= slot_left + slot_width + 1
    assert pic.top + pic.height <= slot_top + slot_height + 1
    assert not slide_spec.findings


def test_no_findings_when_photo_smaller_or_equal_native_aspect_matches_slot(new_slide, profile, tmp_path):
    """Регрессия по построению: если фото и слот одной физической пропорции
    (`Box` — доля ХОЛСТА, не квадратные единицы — канва 16:9, поэтому доля
    ширины/высоты сама по себе не равна физическому аспекту слота, отсюда
    пересчёт ниже), "contain" заполняет слот целиком (не просто "влезает")."""
    slide = new_slide()
    target_aspect = 2.0  # тот же аспект, что и у тестового фото ниже (800x400)
    width_frac = 0.4
    height_frac = (width_frac * profile.canvas_width_emu) / (target_aspect * profile.canvas_height_emu)
    box = Box(left=0.1, top=0.1, width=width_frac, height=height_frac)
    pattern = _pattern_with_image_slot(box)
    slide_spec = SlideSpec(
        index=0, kind="photo_text", headline="Демо", visual=Visual(kind="photo", photo_name="demo.jpg"),
    )
    photo_path = _make_photo(tmp_path, "demo.jpg", (800, 400))  # тот же аспект 2:1

    _place_picture_visual(slide, slide_spec, pattern, profile, "photo", {"demo.jpg": photo_path})

    pic = _picture_shapes(slide)[0]
    slot_width = round(box.width * profile.canvas_width_emu)
    slot_height = round(box.height * profile.canvas_height_emu)
    assert pic.width == pytest.approx(slot_width, abs=2)
    assert pic.height == pytest.approx(slot_height, abs=2)


def test_missing_slot_reports_a_finding_and_places_nothing(new_slide, profile, tmp_path):
    """Бриф задачи, п.3: раскладка под картинку в шаблоне может не
    найтись — тогда фотография не ставится, но об этом сказано находкой
    (не молча, в отличие от ассета каталога шаблона в этой же ситуации)."""
    slide = new_slide()
    pattern = _pattern_without_image_slot()
    slide_spec = SlideSpec(
        index=2, kind="bullets", headline="Команда",
        visual=Visual(kind="photo", photo_name="team.jpg"),
    )
    photo_path = _make_photo(tmp_path, "team.jpg", (800, 600))

    _place_picture_visual(slide, slide_spec, pattern, profile, "photo", {"team.jpg": photo_path})

    assert _picture_shapes(slide) == []
    assert any(
        "не вставлена" in f and "team.jpg" in f for f in slide_spec.findings
    ), slide_spec.findings


def test_missing_slot_without_user_photo_is_not_a_finding(new_slide, profile):
    """Тот же случай отсутствия слота, но ассет — каталога ШАБЛОНА (не
    пользовательская фотография, `photo_name=None`) — поведение ДО этой
    задачи: молча ничего не подставляем, это не находка (декоративная
    картинка на бедной раскладке — обычный случай, не то, что пользователь
    специально принёс и ждёт увидеть)."""
    slide = new_slide()
    pattern = _pattern_without_image_slot()
    slide_spec = SlideSpec(index=0, kind="bullets", headline="Обычный слайд")

    _place_picture_visual(slide, slide_spec, pattern, profile, "photo", None)

    assert _picture_shapes(slide) == []
    assert slide_spec.findings == []


def test_unreadable_user_photo_file_reports_a_finding_and_places_nothing(new_slide, profile, tmp_path):
    """Бриф: молчаливая тишина хуже отсутствия — битый/нечитаемый файл
    фотографии не подставляет молча что-то другое, а честно сообщает."""
    slide = new_slide()
    box = Box(left=0.1, top=0.1, width=0.5, height=0.3)
    pattern = _pattern_with_image_slot(box)
    slide_spec = SlideSpec(
        index=0, kind="photo_text", headline="Демо", visual=Visual(kind="photo", photo_name="demo.jpg"),
    )
    broken = tmp_path / "demo.jpg"
    broken.write_bytes(b"not an image at all")

    _place_picture_visual(slide, slide_spec, pattern, profile, "photo", {"demo.jpg": broken})

    assert _picture_shapes(slide) == []
    assert any("не читается" in f and "demo.jpg" in f for f in slide_spec.findings)


def test_missing_user_photos_mapping_entry_falls_back_to_template_catalog(new_slide, profile, tmp_path):
    """`photo_name` проставлен, но словарь `user_photos` не несёт запись
    под этим именем (например, файл потерялся между шагами) — деградация
    на ассет каталога шаблона (обычный путь до Task 20), не пустой слот и
    не падение."""
    slide = new_slide()
    box = Box(left=0.1, top=0.1, width=0.5, height=0.3)
    pattern = _pattern_with_image_slot(box)
    slide_spec = SlideSpec(
        index=0, kind="photo_text", headline="Демо", visual=Visual(kind="photo", photo_name="missing.jpg"),
    )

    _place_picture_visual(slide, slide_spec, pattern, profile, "photo", {})

    # У VK Education есть каталог фото — ассет каталога подставлен как и раньше.
    assert len(_picture_shapes(slide)) == 1


def test_user_photo_flows_through_build_deck_end_to_end(tmp_path):
    """Интеграционная проверка: `user_photos`, переданный в `build_deck`
    (как это делает `cli.py` после `plan.photos.assign_photos`), реально
    доходит до готового `.pptx` — на реальном шаблоне и реальном подборе
    раскладки, не на синтетическом `Pattern`."""
    profile = TemplateProfile.from_file(EDU_TEMPLATE, cache_dir=None)
    photo_path = _make_photo(tmp_path, "demo.jpg", (900, 600))
    deck = DeckSpec(
        title="Т", language="ru",
        slides=[SlideSpec(
            index=0, kind="image", headline="Демо продукта",
            visual=Visual(kind="photo", photo_name="demo.jpg"),
        )],
    )

    out = build_deck(deck, profile, EDU_TEMPLATE, Variant.dense, user_photos={"demo.jpg": photo_path})
    prs = Presentation(str(out))
    pics = [sh for slide in prs.slides for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1
