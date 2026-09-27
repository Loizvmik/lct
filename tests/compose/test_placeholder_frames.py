"""Задача T4: пустые рамки «Вставить фото/QR», рыба лейаута и пустые
необязательные места в клоне, ряд единиц сборки с нуля ниже заголовка,
график без образца в стиле шаблона. Шаблоны настоящие, модель не
вызывается."""
from __future__ import annotations
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.compose import builder
from deckforge.compose.blocks import assign_content_with_drops
from deckforge.compose.charts import ChartSpec, Series, add_chart
from deckforge.compose.clone import (
    hide_layout_prompts, remove_orphan_plates, remove_stray_text, shape_text, slide_refs,
)
from deckforge.ooxml.geometry import Box, Canvas
from deckforge.ooxml.ns import qn
from deckforge.plan.spec import BulletBlock, Card, CardBlock, DeckSpec, SlideSpec, TextBlock
from deckforge.plan.variants import Variant
from deckforge.template.profile import TemplateProfile

TEMPLATES = Path("dataset/templates")
TECH = TEMPLATES / "VK Tech шаблон.pptx"
WORKSPACE = TEMPLATES / "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx"
LCT = TEMPLATES / "ЛЦТ2026 Шаблон презентации.pptx"
EDU = TEMPLATES / "Шаблон презентации VK Education.pptx"


def _canvas(prs) -> Canvas:
    return Canvas(width_emu=prs.slide_width, height_emu=prs.slide_height)


def _ids(slide) -> set[str]:
    return {el.get("id") for el in slide._element.iter(qn("p:cNvPr"))}


@pytest.fixture(scope="module")
def lct_profile():
    return TemplateProfile.from_file(LCT, cache_dir=None)


@pytest.fixture(scope="module")
def ws_profile():
    return TemplateProfile.from_file(WORKSPACE, cache_dir=None)


# ---------------------------------------------------------------------------
# Рамки-заглушки
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("number,text_id,plate_id", [(5, "442", "441"), (9, "507", "506"), (4, "433", "432")])
def test_insert_photo_or_qr_hint_leaves_with_its_plate(number, text_id, plate_id):
    """VK Tech, примеры 4, 5, 9: «Вставить фото/QR» на белом квадрате. Без
    фото пользователя уходит и подсказка, и квадрат (у QR подсказка шире
    квадрата, поэтому подложка ищется по центру надписи)."""
    prs = Presentation(str(TECH))
    slide = prs.slides[number - 1]
    assert {text_id, plate_id} <= _ids(slide)
    keep = [
        ref.element for ref in slide_refs(slide, _canvas(prs))
        if ref.is_placeholder and shape_text(ref.element).strip()
    ]
    remove_stray_text(slide, _canvas(prs), keep=keep)
    assert text_id not in _ids(slide) and plate_id not in _ids(slide)


def test_plate_with_something_left_on_it_stays():
    """Подложка, на которой осталась надпись или картинка, не трогается."""
    prs = Presentation(str(TECH))
    slide = prs.slides[4]
    canvas = _canvas(prs)
    plate = next(r for r in slide_refs(slide, canvas) if r.shape_id == "441")
    # Надпись «Вставить QR» ещё на месте: подложку снимать нельзя.
    assert remove_orphan_plates(slide, canvas, [plate.box]) == 0
    assert "441" in _ids(slide)


def test_note_plate_goes_with_its_unfilled_note():
    """WorkSpace, пример 3: белая плашка «Заметка». Незаполненная заметка
    снимается вместе с плашкой."""
    prs = Presentation(str(WORKSPACE))
    slide = prs.slides[2]
    canvas = _canvas(prs)
    note = next(r for r in slide_refs(slide, canvas) if r.shape_id == "414")
    note.element.getparent().remove(note.element)
    remove_orphan_plates(slide, canvas, [note.box])
    # Плашка уходит вместе со своим уголком-значком.
    assert not {"413", "415", "416"} & _ids(slide)


def test_shape_beyond_the_canvas_is_removed():
    """VK Education, пример 26: пустая рамка текста справа за краем слайда.
    Аудит L01 отклонял из-за неё клон, и фото слайда терялось."""
    prs = Presentation(str(EDU))
    slide = prs.slides[25]
    assert "764" in _ids(slide)
    remove_stray_text(slide, _canvas(prs))
    assert "764" not in _ids(slide)


def test_layout_prompt_text_is_hidden_when_its_placeholder_is_absent():
    """ЛЦТ2026, пример 24: LibreOffice рисует подсказки лейаута «Решение» /
    «Ваше предложение...», если плейсхолдера на слайде нет. Графика
    лейаута выключается, оформление копируется на слайд."""
    prs = Presentation(str(LCT))
    slide = prs.slides[23]
    canvas = _canvas(prs)
    before = len(list(slide._element.iter(qn("p:cNvPr"))))
    assert hide_layout_prompts(slide, canvas)
    assert slide._element.get("showMasterSp") == "0"
    assert len(list(slide._element.iter(qn("p:cNvPr")))) >= before
    # Повторный вызов ничего не копирует второй раз.
    after = len(list(slide._element.iter(qn("p:cNvPr"))))
    hide_layout_prompts(slide, canvas)
    assert len(list(slide._element.iter(qn("p:cNvPr")))) == after


def test_two_cards_of_three_leave_no_layout_prompts_on_lct_clone(lct_profile, tmp_path):
    """ЛЦТ2026, клон 24 с двумя карточками из трёх: в файле слайд не
    показывает графику лейаута, и рыба «Решение» не выходит на картинку."""
    cards = CardBlock(items=[Card(title="Проблема", body="Согласование идёт девять дней"),
                             Card(title="Аналоги", body="Ручной обход и почта")])
    spec = DeckSpec(title="t", language="ru", slides=[SlideSpec(
        index=0, kind="two_col", headline="Пилот сократил согласование вдвое", blocks=[cards], pattern_id="slide24",
    )])
    out = builder.build_deck(spec, lct_profile, LCT, Variant.dense)
    slide = Presentation(str(out)).slides[0]
    texts = " ".join(shape_text(el) for el in slide._element.iter(qn("p:sp")))
    assert "Ваше предложение" not in texts
    assert any("клоном слайда-примера №24" in f for f in spec.slides[0].findings)
    assert slide._element.get("showMasterSp") == "0"


def test_layout_with_a_third_of_the_canvas_left_empty_is_not_taken(lct_profile):
    """ЛЦТ2026, пример 20: содержание легло в строки справа, левая рамка на
    треть холста пустеет. Такой клон отклоняется."""
    cards = CardBlock(items=[Card(title=f"Шаг {i}", body="Коротко по делу") for i in range(1, 5)])
    spec = DeckSpec(title="t", language="ru", slides=[SlideSpec(
        index=0, kind="photo_text", headline="Пилот сократил согласование вдвое", blocks=[cards],
        pattern_id="slide20",
    )])
    builder.build_deck(spec, lct_profile, LCT, Variant.dense)
    assert any("slide20" in f and "пустеет" in f for f in spec.slides[0].findings)


# ---------------------------------------------------------------------------
# Сборка с нуля: ряд единиц ниже заголовка
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("pattern_id", ["slide3", "slide4", "slide17"])
def test_scattered_units_stay_below_the_headline(ws_profile, pattern_id):
    grid = builder._grid_from_model(ws_profile.grid)
    pattern = next(builder._pattern_from_model(m) for m in ws_profile.patterns if m.pattern_id == pattern_id)
    spec = SlideSpec(index=1, kind=pattern.kind, headline="Пилот сократил согласование вдвое", blocks=[
        BulletBlock(items=["Первый пункт", "Второй пункт", "Третий пункт"]), TextBlock(text="Вывод абзацем"),
    ])
    contents, _drops = assign_content_with_drops(spec, pattern, grid, cards_as_text=True)
    aligned, _plaques = builder._align_scattered_units(contents, pattern, grid)
    headline = next(c.slot.box for c in aligned if c.role_hint == "headline")
    for content in aligned:
        if content.role_hint in builder._UNIT_ROLES and content.slot.box.width < builder._UNIT_MAX_WIDTH:
            assert content.slot.box.top >= headline.top + headline.height, (pattern_id, content.role_hint)


# ---------------------------------------------------------------------------
# График без образца
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("path", [TECH, WORKSPACE, LCT], ids=["tech", "workspace", "lct"])
def test_chart_without_prototype_follows_the_template_style(path):
    profile = TemplateProfile.from_file(path, cache_dir=None)
    prs = Presentation(str(path))
    slide = prs.slides[0]
    chart_spec = ChartSpec(
        kind="bar", categories=["I", "II", "III"], series=[Series(name="Заявки", values=[3.0, 5.0, 4.0])],
    )
    frame = add_chart(slide, Box(0.1, 0.3, 0.6, 0.5), chart_spec, profile)
    chart = frame.chart
    assert not chart.value_axis.has_major_gridlines
    assert not chart.category_axis.has_major_gridlines
    plot = chart.plots[0]
    assert plot.has_data_labels
    assert plot.data_labels.font.size.pt == pytest.approx(profile.type_scale_pt("caption", 12.0))
    assert chart.font.name == (profile.type_scale.families[0] if profile.type_scale.families else "Arial")
    space = chart._chartSpace
    for owner in (space, space.find(qn("c:chart") + "/" + qn("c:plotArea"))):
        sp_pr = owner.find(qn("c:spPr"))
        assert sp_pr is not None and sp_pr.find(qn("a:noFill")) is not None
        assert sp_pr.find(qn("a:ln") + "/" + qn("a:noFill")) is not None
    fill = plot.series[0].points[0].format.fill.fore_color.rgb
    assert f"#{fill}".upper() == profile.chart_series[0].upper()
