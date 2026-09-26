"""Классификация отказа и политика починки (задача V3): каждая неудачная
попытка называет причину, путь лестницы зависит от главной причины, а
кегль при подгонке не ужимается глубже предела. Модель не зовётся: там,
где лестнице нужна починка текста, она фейковая."""
from __future__ import annotations
import functools
from dataclasses import dataclass, replace
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.audit.config import AuditConfig
from deckforge.audit.deterministic import _build_context_inmemory, _check_T02
from deckforge.compose import builder
from deckforge.compose.clone import sample_slides_by_number, set_text_size
from deckforge.compose.failure import (
    CATEGORIES, DEFAULT_POLICY, FontBudget, RepairPolicy, classify, condense_to_theses, from_findings, primary,
)
from deckforge.ooxml.geometry import Canvas
from deckforge.pattern.forms import forms_of
from deckforge.plan.spec import BulletBlock, Card, CardBlock, DeckSpec, SlideSpec, TextBlock
from deckforge.settings import Settings
from deckforge.template.profile import TemplateProfile

TEMPLATE = Path("dataset/templates/Шаблон презентации VK Education.pptx")
APP_YAML = Path(__file__).resolve().parents[2] / "config" / "app.yaml"


@functools.lru_cache(maxsize=None)
def _profile() -> TemplateProfile:
    return TemplateProfile.from_file(TEMPLATE, cache_dir=None)


@pytest.fixture(scope="module")
def profile():
    return _profile()


@dataclass
class _Finding:
    check_id: str
    repair: str = "structural"
    message: str = ""


# ---------------------------------------------------------------------------
# Классификация
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("code, category", [
    ("UNITS_OVER_REPEAT", "TOO_MANY_UNITS"),
    ("SLOT_NO_SHAPE", "MISSING_SLOT"),
    ("HEADLINE_NO_SHAPE", "MISSING_SLOT"),
    ("BLOCK_NO_SLOT", "MISSING_SLOT"),
    ("FONT_BUDGET", "TEXT_OVERFLOW"),
    ("UNDERFILLED", "LOW_DENSITY"),
    ("SHORTEN_EMPTY", "MODEL_FAILURE"),
    ("NO_MODEL", "MODEL_FAILURE"),
    ("CLONE_EXCEPTION", "BUILD_ERROR"),
])
def test_clone_refusal_codes_map_to_categories(code, category):
    failure = classify(code, "причина", "slide1")
    assert failure.category == category
    assert failure.code == code and failure.pattern_id == "slide1"
    assert failure.label == f"{category}/{code}"


@pytest.mark.parametrize("check_id, category", [
    ("L01", "OVERLAP"), ("L02", "OVERLAP"), ("L03", "TEXT_OVERFLOW"), ("L04", "TEXT_OVERFLOW"),
    ("D05", "LOW_DENSITY"), ("T02", "TEXT_OVERFLOW"),
])
def test_audit_findings_map_to_categories(check_id, category):
    (failure,) = from_findings([_Finding(check_id)])
    assert (failure.category, failure.code) == (category, check_id)


def test_recoverable_by_comes_from_the_audit_repair_class():
    failures = from_findings([_Finding("L03", "local"), _Finding("L03", "structural"), _Finding("D05", "none")])
    by_code = {f.code: f.recoverable_by for f in failures}
    assert by_code == {"L03": "structural", "D05": "none"}, "одна причина на код, самая тяжёлая"


def test_primary_failure_is_the_most_structural():
    overflow, units = classify("FONT_BUDGET"), classify("UNITS_OVER_REPEAT")
    assert primary([overflow, units]) is units
    assert primary([classify("UNDERFILLED"), overflow]) is overflow
    assert primary([]) is None
    assert set(DEFAULT_POLICY) <= set(CATEGORIES)


# ---------------------------------------------------------------------------
# Политика
# ---------------------------------------------------------------------------


def test_policy_sends_overflow_and_missing_slot_different_ways():
    policy = RepairPolicy.from_mapping(None)
    overflow, missing = policy.steps("TEXT_OVERFLOW"), policy.steps("MISSING_SLOT")
    assert overflow[:3] == ("roomier", "shorten", "split")
    assert missing[0] == "alternate"
    assert "shorten" not in missing, "сокращение не создаёт места под блок"
    assert policy.steps("TOO_MANY_UNITS")[:2] == ("larger", "split")
    assert policy.steps("MODEL_FAILURE") == ("fallback_text",)


def test_policy_from_app_yaml_is_valid_and_rejects_typos():
    config = Settings.load(APP_YAML).compose
    policy = RepairPolicy.from_mapping(config.repair_policy)
    assert "shorten" not in policy.steps("MISSING_SLOT")
    assert (config.font_degradation_budget.min_ratio, config.font_degradation_budget.max_steps) == (0.8, 2)
    with pytest.raises(ValueError):
        RepairPolicy.from_mapping({"TEXT_OVERFLOW": ["shrink_more"]})
    with pytest.raises(ValueError):
        RepairPolicy.from_mapping({"OVERFLOW": ["split"]})


# ---------------------------------------------------------------------------
# Предел кегля
# ---------------------------------------------------------------------------


def test_font_floor_is_the_stricter_of_ratio_and_two_steps():
    budget = FontBudget(min_ratio=0.8, max_steps=2)
    vk = [12.0, 14.0, 16.0, 44.0, 48.0, 88.0]
    # 44 -> 40 -> 36 (промежуточные ступени), 0,8 дали бы 35,2.
    assert budget.floor_pt(44.0, vk) == pytest.approx(36.0)
    # 16 -> 14 -> 12 по шкале, но 0,8 * 16 = 12,8 строже.
    assert budget.floor_pt(16.0, vk) == pytest.approx(12.8)
    assert FontBudget(min_ratio=0.0, max_steps=0).floor_pt(20.0, vk) == 20.0


def _clone_on_section(profile, headline: str, monkeypatch, budget: FontBudget):
    monkeypatch.setattr(builder, "font_budget", lambda: budget)
    prs = Presentation(str(TEMPLATE))
    sources = sample_slides_by_number(prs)
    builder._clear_sample_slides(prs)
    pattern = builder._pattern_from_model(next(m for m in profile.patterns if m.pattern_id == "slide16"))
    spec = SlideSpec(index=0, kind="section", headline=headline, pattern_id="slide16")
    outcome = builder.place_slide_by_clone(
        prs, spec, pattern, profile, AuditConfig.load(), sources[pattern.source_slide_index[0]],
    )
    return outcome, prs


def test_font_budget_rejects_a_clone_that_needs_a_deeper_shrink(profile, monkeypatch):
    headline = "Итоги пилота по сокращению времени согласования заявок в трёх регионах и дальнейшие шаги"
    loose, prs = _clone_on_section(profile, headline, monkeypatch, FontBudget(min_ratio=0.0, max_steps=99))
    assert loose.reason is None and len(prs.slides) == 1, "без предела заголовок влезает мелким кеглем"

    tight = FontBudget(hero_min_ratio=0.8, hero_max_steps=2)
    strict, prs = _clone_on_section(profile, headline, monkeypatch, tight)
    assert strict.code == "FONT_BUDGET"
    assert "не помещается" in strict.reason
    assert len(prs.slides) == 0, "отклонённый клон убран из колоды"


def test_hero_headline_has_a_softer_budget_than_other_slots():
    budget = FontBudget()
    assert budget.for_slot("headline", "section").floor_pt(44.0, [16.0, 44.0]) == pytest.approx(28.0)
    assert budget.for_slot("headline", "cards") is budget
    assert budget.for_slot("caption", "section") is budget


def test_long_cover_title_stays_a_clone_on_vk_education(profile):
    """Название презентации даёт пользователь, сокращать его нельзя:
    длинный заголовок обложки ужимается мельче примера, но слайд остаётся
    клоном, а не собирается с нуля."""
    slide = SlideSpec(
        index=0, kind="section", pattern_id="slide16",
        headline="Сокращение времени согласования заявок на закупку в трёх регионах",
    )
    outcome, prs = _ladder(profile, slide, None, ("slide16",))
    assert outcome.rung == "clone", outcome.notes
    assert "failure_codes" not in slide.meta
    assert len(prs.slides) == 1


def test_audit_flags_text_shrunk_below_the_budget_as_structural(profile):
    prs = Presentation(str(TEMPLATE))
    sources = sample_slides_by_number(prs)
    builder._clear_sample_slides(prs)
    pattern = builder._pattern_from_model(next(m for m in profile.patterns if m.pattern_id == "slide16"))
    spec = SlideSpec(index=0, kind="section", headline="Итоги пилота", pattern_id="slide16")
    outcome = builder.place_slide_by_clone(
        prs, spec, pattern, profile, AuditConfig.load(), sources[pattern.source_slide_index[0]],
    )
    assert outcome.reason is None
    canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)
    slide = prs.slides[0]
    slot = next(s for s in pattern.slots if s.role == "headline")
    shape = next(sh for sh in slide.shapes if str(sh.shape_id) == slot.source_shape_id)

    config = AuditConfig.load()

    def t02():
        return _check_T02(_build_context_inmemory(0, slide, canvas, profile), profile, config)

    set_text_size(shape._element, 40.0)
    assert not any(f.repair == "structural" for f in t02()), "ступень в пределах: не находка"

    set_text_size(shape._element, 16.0)
    bad = t02()
    assert any(f.repair == "structural" and "глубже предела" in f.message for f in bad)


# ---------------------------------------------------------------------------
# Сокращение без модели
# ---------------------------------------------------------------------------


def test_condense_keeps_first_sentences_and_says_none_when_nothing_to_cut():
    slide = SlideSpec(
        index=0, kind="bullets", headline="Итоги",
        blocks=[
            BulletBlock(items=["Медиана 31,5 часа. Работа 28 минут.", "Ожидание 98,5%"]),
            CardBlock(items=[Card(title="Шаг", body="Первое. Второе.")]),
            TextBlock(text="Абзац один. Абзац два."),
        ],
    )
    condensed = condense_to_theses(slide)
    assert condensed.blocks[0].items == ["Медиана 31,5 часа.", "Ожидание 98,5%"]
    assert condensed.blocks[1].items[0].body == "Первое."
    assert condensed.blocks[2].text == "Абзац один."
    assert slide.blocks[0].items[0].endswith("28 минут."), "исходный слайд не тронут"
    assert condense_to_theses(condensed) is None


# ---------------------------------------------------------------------------
# Лестница на реальном шаблоне
# ---------------------------------------------------------------------------


class _FakeRepair:
    def __init__(self, result=None):
        self.result, self.calls = result, []

    def shorten(self, slide, pattern_id, problems):
        self.calls.append(pattern_id)
        return self.result(slide) if self.result else None


def _ladder(profile, slide: SlideSpec, repair, pattern_ids):
    prs = Presentation(str(TEMPLATE))
    sources = sample_slides_by_number(prs)
    builder._clear_sample_slides(prs)
    by_id = {m.pattern_id: m for m in profile.patterns}
    patterns = [builder._pattern_from_model(by_id[pid]) for pid in pattern_ids]
    canvas = Canvas(width_emu=profile.canvas_width_emu, height_emu=profile.canvas_height_emu)
    outcome = builder._place_with_ladder(
        prs, slide, patterns, profile, canvas, AuditConfig.load(), forms=forms_of(profile), repair=repair,
        room=5, bullet_char="•", user_photos=None, image_bytes=None, source_slides=sources,
    )
    return outcome, prs


def test_missing_slot_goes_to_a_layout_with_places_without_shortening(profile):
    """Карточки на раскладке-разделителе без мест под них: ищем раскладку
    с местами, а модель сокращать текст не зовём, это не поможет."""
    slide = SlideSpec(
        index=0, kind="cards", headline="Три шага раскатки", pattern_id="slide16",
        blocks=[CardBlock(items=[Card(title=f"Шаг {i}", body=f"Короткое описание шага {i}") for i in range(3)])],
    )
    repair = _FakeRepair(lambda s: s)

    outcome, prs = _ladder(profile, slide, repair, ("slide16", "slide21"))

    assert repair.calls == []
    assert outcome.rung == "adapt" and outcome.pattern.pattern_id == "slide21"
    assert slide.meta["failure_codes"].startswith("MISSING_SLOT/")
    assert slide.meta["ladder_path"] == "clone>alternate"
    assert len(prs.slides) == 1


def test_model_failure_falls_back_to_theses_without_a_model(profile):
    """Модель не ответила: текст сокращается до первых предложений и снова
    идёт клоном, а не сразу с нуля."""
    tail = " " + " ".join(["Подробности про очередь и ожидание согласующего."] * 10)
    theses = ["Медиана 31,5 часа, работа 28 минут.", "Ожидание 98,5% времени."]
    slide = SlideSpec(
        index=0, kind="bullets", headline="Заявка ждёт", pattern_id="slide13",
        blocks=[BulletBlock(items=[t + tail for t in theses])],
    )
    repair = _FakeRepair(None)

    outcome, _prs = _ladder(profile, slide, repair, ("slide13",))

    assert len(repair.calls) == 1
    assert outcome.rung == "shorten"
    assert "MODEL_FAILURE/SHORTEN_EMPTY" in slide.meta["failure_codes"]
    assert "shorten>fallback_text" in slide.meta["ladder_path"]
    assert slide.blocks[0].items == theses
    assert any("без модели" in n for n in outcome.notes)


def test_build_deck_records_rungs_and_reasons_per_slide(profile, tmp_path, monkeypatch):
    from deckforge.plan.variants import Variant

    monkeypatch.setattr(builder, "_output_path", lambda spec, tpl, variant: tmp_path / "deck.pptx")
    spec = DeckSpec(title="Проверка", language="ru", slides=[
        SlideSpec(index=0, kind="section", headline="Итоги квартала", pattern_id="slide16"),
        SlideSpec(
            index=1, kind="section", headline="Три шага", pattern_id="slide16",
            blocks=[CardBlock(items=[Card(title=f"Шаг {i}", body=f"Описание {i}") for i in range(3)])],
        ),
    ])
    builder.build_deck(spec, profile, TEMPLATE, Variant.dense)

    counts = builder.ladder_counts(spec)
    assert sum(counts.values()) == len(spec.slides)
    for slide in spec.slides:
        assert slide.meta["ladder_rung"] in builder.LADDER_RUNGS
        assert counts[slide.meta["ladder_rung"]] >= 1
    assert "failure_codes" not in spec.slides[0].meta
    assert spec.slides[1].meta["failure_codes"].startswith("MISSING_SLOT/")
    assert spec.slides[1].meta["failure_primary"].startswith("MISSING_SLOT/")


def test_html_report_says_why_a_slide_is_not_a_clone():
    from deckforge.export.html import _ladder_html, _why_not_clone_html

    slide = SlideSpec(index=0, kind="bullets", headline="Итоги")
    assert _why_not_clone_html(slide) == ""
    slide.meta.update({
        "failure_codes": "TEXT_OVERFLOW/FONT_BUDGET,MODEL_FAILURE/SHORTEN_EMPTY",
        "failure_primary": "TEXT_OVERFLOW/FONT_BUDGET",
        "ladder_path": "clone>roomier>shorten>fallback_text", "ladder_rung": "shorten",
    })
    html = _why_not_clone_html(slide)
    assert 'class="slide-ladder"' in html
    assert "текст не влез" in html and "модель не сократила текст" in html
    assert "clone&gt;roomier&gt;shorten" in html
    deck = DeckSpec(title="т", language="ru", slides=[slide], meta={"ladder_shorten": "1"})
    assert "почему не клон: текст не влез 1" in _ladder_html(deck)
