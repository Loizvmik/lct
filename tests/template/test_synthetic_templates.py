"""Незнакомый шаблон не роняет пайплайн (задача D4).

Четыре шаблона датасета разобраны и выверены вручную, на защите придёт
пятый. Три синтетических шаблона из `fixtures/templates-synthetic/`
(генератор `scripts/make_synthetic_templates.py`) бьют по допущениям,
которые на датасете никогда не нарушались: лейауты без своих коробок
плейсхолдеров, колонтитулы на слайдах-примерах, холст 4:3, карточки
группами, шрифт, которого нет в системе, пустая тема шрифтов.

Всё офлайн: ключа модели тесты не видят, разбор и сборка идут запасными
путями, как у пользователя без `.env`."""
from __future__ import annotations

import functools
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckforge.audit.config import AuditConfig
from deckforge.audit.deterministic import run_deterministic
from deckforge.compose.builder import Variant, build_deck
from deckforge.plan.spec import (
    BulletBlock, Card, CardBlock, ChartSeriesData, ChartVisual, DeckSpec, SlideSpec, TableVisual, TextBlock, Visual,
)
from deckforge.template.profile import TemplateProfile

SYNTH_DIR = Path("fixtures/templates-synthetic")
TEMPLATES = ["corporate-16x9.pptx", "grouped-4x3.pptx", "bare-3-slides.pptx"]
# Почти пустому шаблону разрешено не дать ни одной раскладки: колода тогда
# собирается на лейаутах мастера. Остальным хотя бы одна обязана найтись.
MAY_HAVE_NO_PATTERNS = {"bare-3-slides.pptx"}
# Пустой слайд, слайд-картинка, неоткрываемый файл, текст за холстом или
# обрезанный краем: то, что на защите видно с первого взгляда.
FORBIDDEN_CHECKS = {"I01", "I03", "I04", "L01", "L04"}


@functools.lru_cache(maxsize=None)
def _profile(name: str) -> TemplateProfile:
    return TemplateProfile.from_file(SYNTH_DIR / name, cache_dir=None)


def _spec() -> DeckSpec:
    """Шесть слайдов основных видов с содержанием пакета queue-latency:
    обложка, список, карточки, таблица, график, финал."""
    return DeckSpec(
        title="Сокращение времени согласования заявок",
        language="ru",
        slides=[
            SlideSpec(index=0, kind="section", headline="Автоматическая маршрутизация заявок",
                      subhead="Продуктовый комитет, июнь 2026"),
            SlideSpec(
                index=1, kind="bullets", headline="Где уходит время",
                blocks=[BulletBlock(items=[
                    "Ожидание первого согласующего: медиана 18 часов",
                    "Ожидание второго согласующего: медиана 11 часов",
                    "Чистая работа людей: 28 минут из 31,5 часа цикла",
                ])],
                source_note="Замер 1 240 заявок, март–май 2026",
            ),
            SlideSpec(
                index=2, kind="cards", headline="Что нужно для раскатки",
                blocks=[CardBlock(items=[
                    Card(title="Правила закупок", body="Доработка правил для закупок свыше 5 млн ₽"),
                    Card(title="Кадровая система", body="Интеграция для автоподмены отсутствующих"),
                    Card(title="Обучение", body="340 согласующих, 4 вебинара по 40 минут"),
                ])],
            ),
            SlideSpec(
                index=3, kind="table", headline="Пилот: до и после",
                visual=Visual(kind="table", table=TableVisual(rows=[
                    ["Показатель", "До", "После"],
                    ["Сквозная медиана", "31,5 ч", "6,2 ч"],
                    ["Переназначения вручную", "39%", "4%"],
                ])),
            ),
            SlideSpec(
                index=4, kind="bullets", headline="Медиана ожидания по этапам",
                blocks=[TextBlock(text="Дольше всего заявка ждёт первого согласующего.")],
                visual=Visual(kind="chart", chart=ChartVisual(
                    kind="bar", categories=["Первый", "Второй", "Исполнение"],
                    series=[ChartSeriesData(name="Медиана, ч", values=[18, 11, 2.5])], unit="ч",
                )),
            ),
            SlideSpec(index=5, kind="section", headline="Просим согласовать раскатку",
                      subhead="Старт в июле, полная раскатка к октябрю"),
        ],
    )


@pytest.mark.parametrize("name", TEMPLATES)
def test_offline_parse_gives_usable_profile(name):
    profile = _profile(name)
    assert profile.layouts, "у шаблона нет ни одного лейаута в профиле"
    assert profile.canvas_width_emu > 0 and profile.canvas_height_emu > 0
    if name not in MAY_HAVE_NO_PATTERNS:
        assert profile.patterns, "разбор не снял ни одной раскладки"
    # Шкала кеглей есть всегда, даже без темы шрифтов: сборке нужен кегль.
    assert profile.type_scale.steps.get("body")
    assert profile.type_scale.families and all(f.strip() for f in profile.type_scale.families), "пустая гарнитура шкалы"


@pytest.mark.parametrize("variant", [Variant.dense, Variant.visual])
@pytest.mark.parametrize("name", TEMPLATES)
def test_offline_build_has_six_slides_without_empty_or_broken_ones(name, variant):
    profile = _profile(name)
    spec = _spec()
    out = build_deck(spec, profile, SYNTH_DIR / name, variant, max_slides=len(spec.slides))
    prs = Presentation(str(out))
    assert len(prs.slides) == 6
    _assert_no_broken_slides(out, profile)
    # Раскладок таблицы и карточек у части шаблонов нет: содержание обязано
    # лечь всё равно (сборка с нуля, свободное место, карточки текстом).
    assert any(sh.has_table for sh in prs.slides[3].shapes), "таблица пропала"
    assert any(sh.has_chart for sh in prs.slides[4].shapes), "график пропал"
    cards_text = " ".join(sh.text_frame.text for sh in _leaves(prs.slides[2].shapes) if sh.has_text_frame)
    assert "Кадровая система" in cards_text, "карточки пропали"


@pytest.mark.parametrize("name", TEMPLATES)
def test_template_without_any_pattern_builds_on_master_layouts(name):
    """Майнинг не снял ни одной раскладки (примеров нет или все отброшены):
    колода всё равно собирается, на лейаутах мастера, без пустых слайдов."""
    profile = _profile(name).model_copy(update={"patterns": []})
    spec = _spec()
    out = build_deck(spec, profile, SYNTH_DIR / name, Variant.dense, max_slides=len(spec.slides))
    assert len(Presentation(str(out)).slides) == 6
    assert any("лейауте мастера" in f for s in spec.slides for f in s.findings)
    _assert_no_broken_slides(out, profile)


def _leaves(shapes):
    """Фигуры слайда с заходом в группы: карточки (b) собраны группами."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _leaves(shape.shapes)
        else:
            yield shape


def _assert_no_broken_slides(out: Path, profile: TemplateProfile) -> None:
    findings = run_deterministic(out, profile, AuditConfig.load())
    bad = [f"{f.check_id} слайд {f.slide_index}: {f.message}" for f in findings if f.check_id in FORBIDDEN_CHECKS]
    assert not bad, "\n".join(bad)
